"""生成测试：NASTRAN 工况控制语句（Case Control）。

工程背景
--------
Nastran 的工况控制段是"一行一个子工况、一列一条语句"的形态：

```
SUBCASE 1
SUBTITLE = HEAD SEA
LOAD = 1001
SPC = 1
SUBCOM 4
SUBSEQ = 1, 1.2, 1.3
```

这类数据在实践里往往是从表格 / 文本 / 另一个软件里整块拿到的，所以本测试用
``excel.local_direction: vertical``（指南 §19）：一个子工况一行，可以直接粘贴。

**两个模板都是 ``engine: excel``（公式模式）** —— 改 Excel 里的参数，Code 表里的语句
自己就变了，不需要跑命令。快照模式做不到这一点，所以这里刻意用公式模式。

要覆盖的四件事
--------------
1. **语句是可选的吗** —— 每条语句留空时，输出里**没有这条语句的文本**，
   而它在表里占的那一行是**空的**（语句按行对齐，方便整列复制）；
2. **SUBCASE 与 SUBCOM 两种块**都对，``SUBSEQ`` 的内容原样输出（工具不解释它）；
3. **公式模式的一致性**：Code 表里**公式算出来的文本** == Python 渲染的文本
   （这就是"改 Excel 就改输出"的保证 —— 与 ``excel-codegen check --values`` 同一件事）；
4. **整块拼起来**是一段合法的 case control（按 ``seq`` 排序、删掉空行）。

另外还钉住两个曾经会出问题的地方：``{% if %}`` 的条件在公式模式下**必须写 ``.value``**，
以及 ``filename`` 里用到只做排序的变量 ``seq``。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen import create_template, write_results
from excel_codegen.formula_eval import evaluate_template_values
from excel_codegen.models import ProjectConfig, load_config
from excel_codegen.renderer import collect_variables, export_files, render_all
from excel_codegen.utils import ExcelError

CASE_CONTROL_YAML = """\
version: 1

excel:
  output: "cc.xlsx"
  template_sheet: null
  howto_sheet: null
  local_direction: "vertical"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Case Control", "Summary"]

variables:
  global:
    - name: analysis
      default: ""
      type: string
  local:
    - name: seq
      default: "01"
      type: string
    - name: kind
      default: "SUBCASE"
      type: string
      choices: ["SUBCASE", "SUBCOM"]
    - name: subcase_id
      default: 1
      type: int
      min: 1
    - name: label
      default: ""
      type: string
    - name: subtitle
      default: ""
      type: string
    - name: method
      default: ""
      type: string
    - name: load
      default: ""
      type: string
    - name: spc
      default: ""
      type: string
    - name: subseq
      default: ""
      type: string

asserts:
  - "kind.value != 'SUBCOM' or subseq.value != ''"
  - "kind.value != 'SUBCASE' or subseq.value == ''"

templates:
  - name: case_control
    output_sheet: "Case Control"
    start_cell: "B2"
    direction: "horizontal"
    write_case_headers: true
    engine: "excel"
    filename: "cc_{{ seq }}_{{ case_name }}.inc"
    code: |
      {{ kind }} {{ subcase_id }}
      {% if label.value %}LABEL = {{ label }}{% endif %}
      {% if subtitle.value %}SUBTITLE = {{ subtitle }}{% endif %}
      {% if analysis.value %}ANALYSIS = {{ analysis }}{% endif %}
      {% if method.value %}METHOD = {{ method }}{% endif %}
      {% if load.value %}LOAD = {{ load }}{% endif %}
      {% if spc.value %}SPC = {{ spc }}{% endif %}
      {% if subseq.value %}SUBSEQ = {{ subseq }}{% endif %}
  - name: cc_summary
    output_sheet: "Summary"
    start_cell: "B2"
    direction: "vertical"
    write_case_headers: true
    engine: "excel"
    filename: "cc_summary_{{ seq }}_{{ case_name }}.md"
    code: |
      {{ kind }} {{ subcase_id }} ｜ LOAD={{ load }} ｜ SPC={{ spc }} ｜ SUBSEQ={{ subseq }}
"""

#: 模板里语句的**行序**（空语句也占一行，所以各工况之间是对齐的）
STATEMENT_ROWS = ["KIND/ID", "LABEL", "SUBTITLE", "ANALYSIS", "METHOD", "LOAD", "SPC", "SUBSEQ"]

#: 样例子工况：第 4 行是"组合"块，最后一行**除编号外全空**
ROWS: dict[str, dict[str, object]] = {
    "LC1": {"seq": "01", "kind": "SUBCASE", "subcase_id": 1, "subtitle": "HEAD SEA", "load": "101", "spc": "1"},
    "LC2": {"seq": "02", "kind": "SUBCASE", "subcase_id": 2, "subtitle": "BEAM SEA", "load": "102", "spc": "1"},
    "LC3": {
        "seq": "03",
        "kind": "SUBCASE",
        "subcase_id": 3,
        "label": "QTR",
        "subtitle": "QUARTERING SEA",
        "method": "1",
        "load": "103",
        "spc": "2",
    },
    "COMB4": {
        "seq": "04",
        "kind": "SUBCOM",
        "subcase_id": 4,
        "subtitle": "COMBINED 1.0 / 1.2 / 1.3",
        "subseq": "1, 1.2, 1.3",
    },
    "LC5": {"seq": "05", "kind": "SUBCASE", "subcase_id": 5},
}

#: 期望的渲染结果：**空字符串代表这一格是空的**（语句被省略，但行还在）
EXPECTED: dict[str, list[str]] = {
    "LC1": ["SUBCASE 1", "", "SUBTITLE = HEAD SEA", "ANALYSIS = STATICS", "", "LOAD = 101", "SPC = 1", ""],
    "LC2": ["SUBCASE 2", "", "SUBTITLE = BEAM SEA", "ANALYSIS = STATICS", "", "LOAD = 102", "SPC = 1", ""],
    "LC3": [
        "SUBCASE 3",
        "LABEL = QTR",
        "SUBTITLE = QUARTERING SEA",
        "ANALYSIS = STATICS",
        "METHOD = 1",
        "LOAD = 103",
        "SPC = 2",
        "",
    ],
    "COMB4": [
        "SUBCOM 4",
        "",
        "SUBTITLE = COMBINED 1.0 / 1.2 / 1.3",
        "ANALYSIS = STATICS",
        "",
        "",
        "",
        "SUBSEQ = 1, 1.2, 1.3",
    ],
    # 只填了编号：别的语句一条都不输出
    "LC5": ["SUBCASE 5", "", "", "ANALYSIS = STATICS", "", "", "", ""],
}

#: 删掉空行后的语句（这才是"最后的输出"里该有的东西）
EXPECTED_DECK = """\
SUBCASE 1
SUBTITLE = HEAD SEA
ANALYSIS = STATICS
LOAD = 101
SPC = 1
SUBCASE 2
SUBTITLE = BEAM SEA
ANALYSIS = STATICS
LOAD = 102
SPC = 1
SUBCASE 3
LABEL = QTR
SUBTITLE = QUARTERING SEA
ANALYSIS = STATICS
METHOD = 1
LOAD = 103
SPC = 2
SUBCOM 4
SUBTITLE = COMBINED 1.0 / 1.2 / 1.3
ANALYSIS = STATICS
SUBSEQ = 1, 1.2, 1.3
SUBCASE 5
ANALYSIS = STATICS"""


def _statements(lines: list[str]) -> list[str]:
    """去掉空语句，只留下真正会写进 deck 的行。"""
    return [line for line in lines if line]


@pytest.fixture()
def project(tmp_path: Path) -> ProjectConfig:
    path = tmp_path / "case_control.yaml"
    path.write_text(CASE_CONTROL_YAML, encoding="utf-8")
    return load_config(path)


@pytest.fixture()
def workbook_path(tmp_path: Path, project: ProjectConfig) -> Path:
    """建竖排 Local 表并按 ROWS 填值 —— 这一步就是"整块粘贴"的自动化版本。"""
    path = create_template(project, tmp_path / "cc.xlsx", cases=list(ROWS))
    workbook = load_workbook(path)
    try:
        workbook["Global Parameter"]["B2"] = "STATICS"
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        rows = {sheet.cell(row=r, column=1).value: r for r in range(2, sheet.max_row + 1)}
        for case, values in ROWS.items():
            for name, value in values.items():
                sheet.cell(row=rows[case], column=columns[name], value=value)
        workbook.save(path)
    finally:
        workbook.close()
    return path


# --------------------------------------------------------------------------- #
# 逐个 Case：语句可选
# --------------------------------------------------------------------------- #
def test_each_subcase_block(workbook_path: Path, project: ProjectConfig) -> None:
    out = render_all(project, workbook_path)
    results = out.results["case_control"]
    assert [result.case_name for result in results] == list(ROWS)
    for result in results:
        assert result.lines == EXPECTED[result.case_name], result.case_name


def test_blank_statements_are_not_in_the_output(workbook_path: Path, project: ProjectConfig) -> None:
    """留空 = 这条语句不出现在输出里（它在表里占的那一行是空的）。"""
    results = {r.case_name: r.lines for r in render_all(project, workbook_path).results["case_control"]}
    assert not any("LABEL" in line for line in results["LC1"])

    # 语句按行对齐：LABEL=第 2 行、ANALYSIS=第 4 行、LOAD=第 6 行
    assert results["LC3"][1] == "LABEL = QTR"
    assert results["LC1"][1] == ""
    assert results["LC1"][3] == "ANALYSIS = STATICS"
    assert results["LC1"][5] == "LOAD = 101"

    # LC5 只有 SUBCASE 与全局的 ANALYSIS 两条内容
    assert _statements(results["LC5"]) == ["SUBCASE 5", "ANALYSIS = STATICS"]


def test_subcom_uses_subseq_and_keeps_it_verbatim(workbook_path: Path, project: ProjectConfig) -> None:
    """SUBSEQ 的内容原样输出（工具不解释"子工况号 / 比例因子"的写法）。"""
    lines = next(r for r in render_all(project, workbook_path).results["case_control"] if r.case_name == "COMB4").lines
    assert lines[0] == "SUBCOM 4"
    assert "SUBSEQ = 1, 1.2, 1.3" in lines
    assert not any(line.startswith("LOAD") for line in lines)
    assert not any(line.startswith("SPC") for line in lines)


def test_every_row_starts_with_kind_and_id(workbook_path: Path, project: ProjectConfig) -> None:
    for result in render_all(project, workbook_path).results["case_control"]:
        kind, _, cid = result.lines[0].partition(" ")
        assert kind in {"SUBCASE", "SUBCOM"}, result.case_name
        assert cid.isdigit(), result.case_name


def test_statement_rows_are_aligned(workbook_path: Path, project: ProjectConfig) -> None:
    """每一行是哪条语句是固定的 —— 整列复制出去才能按行对齐。"""
    results = {r.case_name: r.lines for r in render_all(project, workbook_path).results["case_control"]}
    for case, lines in results.items():
        assert len(lines) == len(STATEMENT_ROWS), case
        assert lines[1] == "" or lines[1].startswith("LABEL = "), case
        assert lines[3] == "" or lines[3].startswith("ANALYSIS = "), case
        assert lines[5] == "" or lines[5].startswith("LOAD = "), case
        assert lines[7] == "" or lines[7].startswith("SUBSEQ = "), case


# --------------------------------------------------------------------------- #
# 公式模式：Excel 里的公式算出来 == Python 渲染
# --------------------------------------------------------------------------- #
def test_sheet_holds_formulas_not_snapshots(workbook_path: Path, project: ProjectConfig) -> None:
    """公式模式：Code 表里是**公式**。快照模式下改 Excel 不会改输出。"""
    write_results(workbook_path, project, render_all(project, workbook_path).results)
    workbook = load_workbook(workbook_path)
    try:
        value = workbook["Case Control"].cell(row=2, column=2).value
        assert isinstance(value, str) and value.startswith("="), value
    finally:
        workbook.close()


def test_formula_values_equal_python_render(workbook_path: Path, project: ProjectConfig) -> None:
    """核心保证：Code 表里**公式算出来的文本** == Python 渲染的文本。

    "改 Excel 就改输出"成立的前提就是这条 —— 与 `excel-codegen check --values` 同一件事。
    """
    expected = render_all(project, workbook_path)
    write_results(workbook_path, project, expected.results)

    workbook = load_workbook(workbook_path)
    try:
        for template in project.templates:
            results = expected.results[template.name]
            got = evaluate_template_values(workbook, project, template, [r.case_name for r in results])
            for result in results:
                assert got[result.case_name] == result.lines, (template.name, result.case_name)
    finally:
        workbook.close()


def test_editing_the_sheet_changes_the_output(workbook_path: Path, project: ProjectConfig) -> None:
    """改一格参数、**不重跑渲染**，公式算出来的语句就跟着变。"""
    write_results(workbook_path, project, render_all(project, workbook_path).results)

    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        rows = {sheet.cell(row=r, column=1).value: r for r in range(2, sheet.max_row + 1)}
        sheet.cell(row=rows["LC1"], column=columns["load"], value="9001")
        sheet.cell(row=rows["LC1"], column=columns["spc"]).value = None  # 顺手清掉 SPC
        workbook.save(workbook_path)
    finally:
        workbook.close()

    workbook = load_workbook(workbook_path)
    try:
        got = evaluate_template_values(workbook, project, project.templates[0], ["LC1"])
    finally:
        workbook.close()
    assert _statements(got["LC1"]) == ["SUBCASE 1", "SUBTITLE = HEAD SEA", "ANALYSIS = STATICS", "LOAD = 9001"]

    # 与"用新参数重新渲染"逐行一致
    fresh = render_all(project, workbook_path)
    assert got["LC1"] == fresh.results["case_control"][0].lines


# --------------------------------------------------------------------------- #
# 拼成整段 case control
# --------------------------------------------------------------------------- #
def test_exported_files_concatenate_into_one_deck(workbook_path: Path, project: ProjectConfig, tmp_path: Path) -> None:
    out = render_all(project, workbook_path)
    files = export_files(project, out.results, tmp_path / "deck")
    names = sorted(path.name for path in files if path.suffix == ".inc")
    assert names == ["cc_01_LC1.inc", "cc_02_LC2.inc", "cc_03_LC3.inc", "cc_04_COMB4.inc", "cc_05_LC5.inc"]

    # 空语句在文件里是空行 —— 删掉就是可以直接用的 case control 段
    deck = "".join((tmp_path / "deck" / name).read_text(encoding="utf-8") for name in names)
    assert "\n".join(line for line in deck.splitlines() if line) == EXPECTED_DECK


def test_filename_can_use_any_parameter(workbook_path: Path, project: ProjectConfig, tmp_path: Path) -> None:
    """``seq`` 只出现在 filename 里 —— 不算"定义了没人用"，也要真的能渲染出来。"""
    template = project.templates[0]
    assert "seq" in collect_variables(template, base_dir=project.source_dir)

    out = render_all(project, workbook_path)
    files = export_files(project, out.results, tmp_path / "deck")
    assert (tmp_path / "deck" / "cc_04_COMB4.inc").exists()
    assert len(files) == 2 * len(ROWS)  # 两个模板 × 5 个 Case


# --------------------------------------------------------------------------- #
# Output 表写回
# --------------------------------------------------------------------------- #
def test_blocks_are_written_to_the_sheet(workbook_path: Path, project: ProjectConfig) -> None:
    write_results(workbook_path, project, render_all(project, workbook_path).results)
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Case Control"]
        # 横向：一个 Case 一列；表头在 B1 起
        assert [sheet.cell(row=1, column=c).value for c in range(2, 7)] == list(ROWS)
        # 每个 Case 的列里都是公式（第 1 行 = KIND/ID，第 2 行 = LABEL）
        assert sheet.cell(row=2, column=2).value.startswith("=")
        assert sheet.cell(row=3, column=2).value.startswith("=")
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 填错地方要被拦下来（asserts / choices）
# --------------------------------------------------------------------------- #
def test_subcom_without_subseq_is_rejected(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        # openpyxl 的 cell(..., value=None) 不会真的清空，要直接赋值
        sheet.cell(row=5, column=columns["subseq"]).value = None  # COMB4 的 SUBSEQ 抠掉
        workbook.save(workbook_path)
    finally:
        workbook.close()
    with pytest.raises(ExcelError, match=r"(?s)COMB4.*subseq"):
        render_all(project, workbook_path)


def test_subseq_on_a_plain_subcase_is_rejected(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        sheet.cell(row=2, column=columns["subseq"], value="1, 2")  # LC1 是 SUBCASE，不该有 SUBSEQ
        workbook.save(workbook_path)
    finally:
        workbook.close()
    with pytest.raises(ExcelError, match=r"asserts"):
        render_all(project, workbook_path)


def test_unknown_kind_is_rejected(workbook_path: Path, project: ProjectConfig) -> None:
    """kind 声明了 choices：写成别的词要在读表时就被拦住。"""
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        sheet.cell(row=2, column=columns["kind"], value="SUBCASE 1")  # 把整条语句塞进一格里
        workbook.save(workbook_path)
    finally:
        workbook.close()
    with pytest.raises(ExcelError, match=r"choices|允许列表"):
        render_all(project, workbook_path)


# --------------------------------------------------------------------------- #
# 仓库里那个示例本身也要是好的
# --------------------------------------------------------------------------- #
def _example_paths() -> tuple[Path, Path, Path]:
    here = Path(__file__).resolve().parents[1]
    return (
        here / "examples" / "nastran_case_control.yaml",
        here / "examples" / "nastran_case_control.xlsx",
        here / "examples" / "generated_nastran" / "case_control.deck",
    )


def test_shipped_example_is_formula_mode_and_vertical() -> None:
    """示例的**形态**要钉住（工作表是给大家改的，所以不钉它的取值）。"""
    yaml_path, _, _ = _example_paths()
    if not yaml_path.exists():  # pragma: no cover - sdist 里一定带着
        pytest.skip("示例 YAML 不在")

    config = load_config(yaml_path)
    assert config.excel.local_direction == "vertical", "一个子工况一行，才能整块粘贴"
    assert [template.name for template in config.templates] == ["case_control", "cc_summary"]
    for template in config.templates:
        assert template.engine == "excel", f"{template.name} 必须是公式模式（改 Excel 就改输出）"


def test_shipped_example_workbook_is_consistent() -> None:
    """示例工作簿是**给人改的**，所以这里只查"自洽"，不查具体数值。

    三件事：能渲染、每一块的形状对、Code 表里的公式算出来与 Python 渲染一致
    （也就是"改 Excel 就改输出"在示例上成立）。
    """
    yaml_path, book, _ = _example_paths()
    if not (yaml_path.exists() and book.exists()):  # pragma: no cover - sdist 里一定带着
        pytest.skip("示例不在")
    config = load_config(yaml_path)

    expected = render_all(config, book)
    template = config.templates[0]
    results = expected.results[template.name]
    assert results, "示例工作簿里一个工况都没有？"

    # 形状：每块第一行是 SUBCASE/SUBCOM + 编号；语句行对齐到固定位置
    for result in results:
        kind, _, cid = result.lines[0].partition(" ")
        assert kind in {"SUBCASE", "SUBCOM"}, result.case_name
        assert cid.isdigit(), result.case_name
        assert len(result.lines) == len(STATEMENT_ROWS), result.case_name

    # 公式一致性（工作簿可能被人改过参数 —— 那也不该破坏这条不变式）
    workbook = load_workbook(book)
    try:
        got = evaluate_template_values(workbook, config, template, [r.case_name for r in results])
    finally:
        workbook.close()
    for result in results:
        assert got[result.case_name] == result.lines, result.case_name


def test_shipped_deck_matches_the_workbook() -> None:
    """仓库里那个 `.deck` 应当就是当前示例工作簿渲染出来、删掉空行的结果。"""
    yaml_path, book, deck_path = _example_paths()
    if not (yaml_path.exists() and book.exists() and deck_path.exists()):  # pragma: no cover
        pytest.skip("示例产物不在")

    config = load_config(yaml_path)
    out = render_all(config, book)
    lines = {r.case_name: r.lines for r in out.results["case_control"]}

    workbook = load_workbook(book)
    try:
        sheet = workbook["Local Parameter"]
        columns = {sheet.cell(row=1, column=c).value: c for c in range(2, sheet.max_column + 1)}
        rows = {sheet.cell(row=r, column=1).value: r for r in range(2, sheet.max_row + 1)}
        order = sorted(lines, key=lambda case: str(sheet.cell(row=rows[case], column=columns["seq"]).value))
    finally:
        workbook.close()

    rebuilt = [line for case in order for line in _statements(lines[case])]
    assert deck_path.read_text(encoding="utf-8").splitlines() == rebuilt
