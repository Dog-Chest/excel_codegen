"""公式求值器测试：把 engine: excel 写出来的公式在 Python 里算一遍。

这是"公式算出来的值对不对"的唯一证据链：解析**工作簿里真实的公式文本**，
与 Python 渲染逐行比对。
"""

from __future__ import annotations

from pathlib import Path
from typing import ClassVar

import pytest
from openpyxl import load_workbook

from spreadsheet_codegen.excel_io import create_template, write_results
from spreadsheet_codegen.formula_eval import (
    FormulaEvalError,
    evaluate_formula,
    evaluate_template_values,
)
from spreadsheet_codegen.models import ProjectConfig, load_config
from spreadsheet_codegen.renderer import render_all

FORMULA_YAML = """\
version: 1

excel:
  output: "eval.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output", "Output Vertical"]

variables:
  global:
    - name: baud
      default: 115200
      type: int
    - name: mcu
      default: "STM32F103"
  local:
    - name: port
      default: "A"
      prefix: "GPIO"
      suffix: "_PORT"
    - name: mode
      default: "TX_RX"
      prefix: "MODE_"

templates:
  - name: uart_init
    output_sheet: "Output"
    start_cell: "B2"
    direction: "horizontal"
    engine: "excel"
    code: |
      // Case: {{ case_name }}
      UART_Init({{ baud }}, {{ port }}, {{ mode }});
      // 纯值 {{ port.value }} / 前缀 {{ port.prefix }} / 后缀 {{ port.suffix }}
      // 常量行 with "quotes"
  - name: pins
    output_sheet: "Output Vertical"
    start_cell: "B2"
    direction: "vertical"
    engine: "excel"
    code: |
      {{ case_name }},{{ port }},{{ baud }}
"""


class FakeReader:
    """给求值器喂一个内存表：``{表名: {变量名: (值, 空前缀/后缀)}}`` 太绕，
    这里直接按单元格地址存。"""

    def __init__(self, cells: dict[str, dict[str, object]], names: dict[str, dict[str, int]]):
        self.cells = cells
        self.names = names

    def cell_value(self, sheet: str, column: int, row: int):
        from spreadsheet_codegen.utils import column_index_to_letter

        address = f"{column_index_to_letter(column)}{row}"
        if sheet not in self.cells:
            raise FormulaEvalError(f"没有工作表 {sheet}")
        value = self.cells[sheet].get(address)
        if isinstance(value, str) and value.startswith("="):
            raise FormulaEvalError(f"{sheet}!{address} 本身是公式")
        return value, value is None or value == ""

    def name_row(self, sheet: str, variable: str) -> int:
        try:
            return self.names[sheet][variable]
        except KeyError as exc:
            raise FormulaEvalError(f"{sheet} 表里没有变量 {variable!r}") from exc

    def name_column(self, sheet: str, variable: str) -> int:
        raise FormulaEvalError(f"{sheet} 表第 1 行里没有变量 {variable!r}")


class _Cfg:
    """只给 sheet_names_of 用。"""

    class excel:
        class sheets:
            global_: ClassVar[str] = "Global Parameter"
            local: ClassVar[str] = "Local Parameter"
            outputs: ClassVar[list[str]] = ["Output"]

        template_sheet: ClassVar[str | None] = None
        howto_sheet: ClassVar[str | None] = None


@pytest.fixture()
def reader() -> FakeReader:
    return FakeReader(
        cells={
            "Global Parameter": {"B2": 115200, "D2": None, "E2": None},
            "Local Parameter": {"C2": "GPIO", "D2": "_PORT", "E2": "A"},
        },
        names={"Global Parameter": {"baud": 2}, "Local Parameter": {"port": 2}},
    )


# --------------------------------------------------------------------------- #
# 求值器
# --------------------------------------------------------------------------- #
def test_literals_and_concat(reader: FakeReader) -> None:
    assert evaluate_formula('="a"&"b"', reader=reader, config=_Cfg) == "ab"
    assert evaluate_formula('="他说""你好"""', reader=reader, config=_Cfg) == '他说"你好"'
    assert evaluate_formula('="x"&1&"y"&2.5', reader=reader, config=_Cfg) == "x1y2.5"


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ('=IF(3>2,"y","n")', "y"),
        ('=IF(3<2,"y","n")', "n"),
        ('=IF(2>=2,"y","n")', "y"),
        ('=IF(2<=1,"y","n")', "n"),
        ('=IF(2<>3,"y","n")', "y"),
        ('=IF("B">"A","y","n")', "y"),  # 文本按字典序比
        ('=IF("A"<"B","y","n")', "y"),
        ('=AND((3>2),("A"="A"))', "TRUE"),  # 行内 {% if %} 的 and
        ('=OR((1>2),("A"="A"))', "TRUE"),
        ("=NOT(1>2)", "TRUE"),
        ("=AND((1>2),(2>3))", "FALSE"),
        ('=IF(AND((1>2),(2>3)),"y","n")', "n"),
    ],
)
def test_comparison_and_logical_operators(reader: FakeReader, formula: str, expected: str) -> None:
    """行内 {% if %} 会用 < > <= >= AND OR NOT —— 求值器必须跟编译器同步支持。"""
    assert evaluate_formula(formula, reader=reader, config=_Cfg) == expected


def test_absolute_row_number_is_a_scalar(reader: FakeReader) -> None:
    """绝对行号 `$，纵向输入 + 纵向输出时 INDEX 的行号就是这样。

    （曾经这里直接抛"看不懂的记号 punct='$'" —— 只有"输入纵向 + 输出纵向"才会生成它。）
    """
    assert evaluate_formula("=INDEX('Local Parameter'!$A:$ZZ,$2,3)", reader=reader, config=_Cfg) == "GPIO"


def test_isblank_and_if_fallback(reader: FakeReader) -> None:
    # 空单元格：ISBLANK 为真 → 回落默认值
    assert (
        evaluate_formula(
            "=IF(ISBLANK(INDEX('Global Parameter'!D:D,MATCH(\"baud\",'Global Parameter'!A:A,0))),\"\","
            "INDEX('Global Parameter'!D:D,MATCH(\"baud\",'Global Parameter'!A:A,0)))",
            reader=reader,
            config=_Cfg,
        )
        == ""
    )
    # 非空单元格：取原值
    assert (
        evaluate_formula(
            "=IF(ISBLANK(INDEX('Local Parameter'!C:C,MATCH(\"port\",'Local Parameter'!A:A,0))),\"X\","
            "INDEX('Local Parameter'!C:C,MATCH(\"port\",'Local Parameter'!A:A,0)))",
            reader=reader,
            config=_Cfg,
        )
        == "GPIO"
    )
    # 默认值分支
    assert (
        evaluate_formula(
            "=IF(ISBLANK(INDEX('Global Parameter'!Z:Z,MATCH(\"baud\",'Global Parameter'!A:A,0))),115200,1)",
            reader=reader,
            config=_Cfg,
        )
        == "115200"
    )


def test_text_formats(reader: FakeReader) -> None:
    assert evaluate_formula('=TEXT(115200,"0")', reader=reader, config=_Cfg) == "115200"
    assert evaluate_formula('=TEXT(340.0,"0")', reader=reader, config=_Cfg) == "340"
    assert evaluate_formula('=TEXT(20.559,"0.############")', reader=reader, config=_Cfg) == "20.559"
    assert evaluate_formula('=TEXT(8.0,"0.############")', reader=reader, config=_Cfg) == "8"


def test_direct_references_and_bool_literals(reader: FakeReader) -> None:
    # 相对列 / 绝对列的写法都要能解析（值本身按地址取）
    assert evaluate_formula("='Local Parameter'!E$2&'Local Parameter'!$E$2", reader=reader, config=_Cfg) == "AA"
    assert evaluate_formula('=IF(TRUE,"y","n")', reader=reader, config=_Cfg) == "y"
    assert evaluate_formula('=IF(FALSE,"y","n")', reader=reader, config=_Cfg) == "n"


def test_unsupported_things_raise(reader: FakeReader) -> None:
    with pytest.raises(FormulaEvalError, match="不支持的函数"):
        evaluate_formula("=SUM(1,2)", reader=reader, config=_Cfg)
    with pytest.raises(FormulaEvalError, match="本身是公式"):
        FakeReader({"S": {"A1": "=1+1"}}, {"S": {}}).cell_value("S", 1, 1)
    with pytest.raises(FormulaEvalError, match="没有变量"):
        evaluate_formula(
            "=INDEX('Local Parameter'!E:E,MATCH(\"nope\",'Local Parameter'!$A:$A,0))", reader=reader, config=_Cfg
        )


# --------------------------------------------------------------------------- #
# 与真实工作簿对拍
# --------------------------------------------------------------------------- #
@pytest.fixture()
def formula_workbook(tmp_path: Path) -> tuple[ProjectConfig, Path]:
    config_path = tmp_path / "eval.yaml"
    config_path.write_text(FORMULA_YAML, encoding="utf-8")
    config = load_config(config_path)
    path = create_template(config, tmp_path / "eval.xlsx", cases=2)
    output = render_all(config, path)
    write_results(path, config, output.results)
    return config, path


def test_formula_values_equal_python_render(formula_workbook) -> None:
    """核心断言：Output 表里公式算出来的文本 == render_all() 渲染的文本。"""
    config, path = formula_workbook
    expected = render_all(config, path)
    workbook = load_workbook(path)
    try:
        for template in config.templates:
            results = expected.results[template.name]
            got = evaluate_template_values(workbook, config, template, [result.case_name for result in results])
            for result in results:
                assert got[result.case_name] == result.lines, (
                    template.name,
                    result.case_name,
                )
    finally:
        workbook.close()


def test_formula_values_follow_parameter_change(formula_workbook) -> None:
    """改参数不重跑：公式求值要跟着变，并且与"用新参数再渲染一次"一致。"""
    config, path = formula_workbook
    workbook = load_workbook(path)
    try:
        workbook["Global Parameter"]["B2"] = 9600
        workbook["Local Parameter"]["E2"] = "C"
        workbook.save(path)
    finally:
        workbook.close()

    workbook = load_workbook(path)
    try:
        template = config.templates[0]
        got = evaluate_template_values(workbook, config, template, ["Case1"])
    finally:
        workbook.close()
    assert got["Case1"][1] == "UART_Init(9600, GPIOC_PORT, MODE_TX_RX);"

    # 与"用新参数重新渲染"逐行一致
    fresh = render_all(config, path)
    assert got["Case1"] == fresh.results["uart_init"][0].lines


def test_structural_break_is_detected(formula_workbook) -> None:
    """在 Local 表中间插一个 Case 列：公式文本没变，但算出来的值会指错工况。"""
    config, path = formula_workbook
    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        sheet.insert_cols(6)  # 在 E 之后插一列 → F 变成新列，原 F 挪到 G
        sheet.cell(row=1, column=6, value="Case1b")
        workbook.save(path)
    finally:
        workbook.close()

    template = config.templates[0]
    cases = [case.name for case in render_all(config, path).cases]
    assert cases == ["Case1", "Case1b", "Case2"]

    workbook = load_workbook(path)
    try:
        # 输出表里只有 Case1 / Case2 两列，第 3 个位置的表头写着 Case2 但我们要读 Case1b
        # —— 求值器会当场报"对不上"，而不是把 Case2 的值当成 Case1b 的（那才是真的坏）
        with pytest.raises(FormulaEvalError, match="对不上"):
            evaluate_template_values(workbook, config, template, cases)

        # 表头还对得上的那几个仍然能算：报错是因为"错位"，不是因为整张表废了
        got = evaluate_template_values(workbook, config, template, ["Case1"])
    finally:
        workbook.close()

    fresh = render_all(config, path)
    assert got["Case1"] == fresh.results["uart_init"][0].lines
    # Case1b 那一列的公式还指着 F 列（Case2 的参数）—— 重跑 --write-excel 前用不了
    assert fresh.results["uart_init"][1].lines != fresh.results["uart_init"][2].lines
