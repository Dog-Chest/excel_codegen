"""行列风格（``excel.local_direction``）测试。

输入侧（Local Parameter 表）的工况排布有**两种**，由 ``excel.local_direction`` 控制：

* ``horizontal``（默认）：一个工况**一列**（E1 起写 Case 名）；
* ``vertical``：一个工况**一行**（A2 起写 Case 名）。

它和每个模板自己的 ``direction``（**输出**排布）互相独立 —— 输入竖着填、输出横着写
是最常见的组合（工况控制语句按行给、生成结果按列排）。

这里同时钉住三件事：读回来的工况一致、公式模式的引用指向正确的格、
以及"插空行 / 插空列就多一个工况"这个约定在两种布局下都成立。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen.excel_io import (
    check_required_sheets,
    check_value_constraints,
    create_template,
    read_cases,
    write_results,
)
from excel_codegen.formula import compile_formulas, local_cell
from excel_codegen.formula_eval import evaluate_template_values
from excel_codegen.models import ProjectConfig, load_config
from excel_codegen.renderer import render_all
from excel_codegen.utils import ExcelError

VERTICAL_YAML = """\
version: 1

excel:
  output: "vertical.xlsx"
  template_sheet: "Template"
  howto_sheet: null
  local_direction: "vertical"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output", "Output Vertical"]

variables:
  global:
    - name: baud
      default: 115200
      type: int
  local:
    - name: port
      description: "端口"
      default: "A"
      prefix: "GPIO"
      suffix: "_PORT"
      choices: ["A", "B", "C"]
    - name: mode
      default: "TX_RX"
      prefix: "MODE_"
    - name: width
      type: int
      derived: "baud / 14400"

templates:
  - name: uart_init
    output_sheet: "Output"
    start_cell: "B2"
    direction: "horizontal"
    engine: "excel"
    filename: "uart_init_{{ case_name }}.c"
    code: |
      UART_Init({{ baud }}, {{ port }}, {{ width }});
      // {{ port.value }} | {{ port.prefix }} | {{ port.suffix }}
  - name: uart_rows
    output_sheet: "Output Vertical"
    start_cell: "B2"
    direction: "vertical"
    filename: "uart_rows_{{ case_name }}.md"
    code: |
      {{ case_name }},{{ port }},{{ mode }}
"""


def _write(tmp_path: Path, text: str, name: str = "vertical.yaml") -> ProjectConfig:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return load_config(path)


@pytest.fixture()
def project(tmp_path: Path) -> ProjectConfig:
    return _write(tmp_path, VERTICAL_YAML)


@pytest.fixture()
def workbook_path(tmp_path: Path, project: ProjectConfig) -> Path:
    return create_template(project, tmp_path / "vertical.xlsx", cases=3, overwrite=True)


def _fill(path: Path, values: dict[str, dict[str, object]]) -> None:
    """按 ``{Case 名: {变量: 值}}`` 往纵向 Local 表里填值。"""
    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        rows = {str(sheet.cell(row=row, column=1).value): row for row in range(2, sheet.max_row + 1)}
        columns = {str(sheet.cell(row=1, column=column).value): column for column in range(2, sheet.max_column + 1)}
        for case, items in values.items():
            for variable, value in items.items():
                sheet.cell(row=rows[case], column=columns[variable], value=value)
        workbook.save(path)
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 建表
# --------------------------------------------------------------------------- #
def test_vertical_local_sheet_layout(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        assert [sheet.cell(row=1, column=column).value for column in range(1, 5)] == [
            "Case",
            "port",
            "mode",
            "width",
        ]
        assert [sheet.cell(row=row, column=1).value for row in range(2, 5)] == ["Case1", "Case2", "Case3"]
        assert sheet.freeze_panes == "B2"
        # 纵向布局没有 Prefix / Suffix 列（第 1 行整行都是变量名）
        assert sheet.cell(row=1, column=5).value is None
    finally:
        workbook.close()


def test_vertical_sheet_has_no_prefix_suffix_columns(workbook_path: Path) -> None:
    """前后缀只来自 YAML，写进表头格的批注里。"""
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        comment = sheet.cell(row=1, column=2).comment
        assert comment is not None
        assert "GPIO" in comment.text and "_PORT" in comment.text
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 读工况
# --------------------------------------------------------------------------- #
def test_read_cases_vertical(workbook_path: Path, project: ProjectConfig) -> None:
    _fill(
        workbook_path,
        {
            "Case1": {"port": "A", "mode": "TX"},
            "Case2": {"port": "C", "mode": "RX"},
        },
    )
    workbook = load_workbook(workbook_path)
    try:
        cases = read_cases(workbook, project)
    finally:
        workbook.close()

    assert [case.name for case in cases] == ["Case1", "Case2", "Case3"]
    assert [case.row for case in cases] == [2, 3, 4]
    assert all(case.column is None for case in cases)
    assert [case.where for case in cases] == ["第 2 行", "第 3 行", "第 4 行"]
    assert [case.values["port"].text for case in cases] == ["A", "C", "A"]
    assert [case.values["mode"].text for case in cases] == ["TX", "RX", "TX_RX"]
    # 派生参数：115200 / 14400 = 8
    assert [case.values["width"].text for case in cases] == ["8", "8", "8"]
    # create_template 会把 YAML 默认值填进去，所以初始三行都算"填过"
    assert [case.explicit_values for case in cases] == [True, True, True]


def test_vertical_explicit_values_follow_the_sheet(workbook_path: Path, project: ProjectConfig) -> None:
    """整行都是空的工况要能被认出来（doctor 的"空 Case 行"检查靠它）。"""
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        for column in range(2, 5):
            sheet.cell(row=4, column=column).value = None  # Case3 整行清空
        cases = read_cases(workbook, project)
    finally:
        workbook.close()
    assert [case.explicit_values for case in cases] == [True, True, False]


def test_vertical_case_discovery_stops_at_blank_row(workbook_path: Path, project: ProjectConfig) -> None:
    """空行 = 工况结束（下拉就能加工况，不用改 YAML）。"""
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        sheet.cell(row=3, column=1).value = None  # 抠掉 Case2 → 后面的 Case3 也不再看
        cases = read_cases(workbook, project)
    finally:
        workbook.close()
    assert [case.name for case in cases] == ["Case1"]


def test_vertical_without_any_case_raises(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=1).value = None
        with pytest.raises(ExcelError, match="没有工况"):
            read_cases(workbook, project)
    finally:
        workbook.close()


def test_vertical_case_names_must_be_unique(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        workbook["Local Parameter"].cell(row=3, column=1).value = "Case1"
        with pytest.raises(ExcelError, match="Case 名重复"):
            read_cases(workbook, project)
    finally:
        workbook.close()


def test_vertical_variable_names_must_be_unique(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        workbook["Local Parameter"].cell(row=1, column=3).value = "port"
        with pytest.raises(ExcelError, match="变量名 'port' 重复"):
            read_cases(workbook, project)
    finally:
        workbook.close()


def test_horizontal_and_vertical_agree(tmp_path: Path) -> None:
    """同一份参数换个排布，读出来的 Case 必须一模一样。"""
    horizontal = _write(tmp_path, VERTICAL_YAML.replace('local_direction: "vertical"', 'local_direction: "horizontal"'))
    path = create_template(horizontal, tmp_path / "horizontal.xlsx", cases=2, overwrite=True)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        sheet["E2"] = "A"  # port 那一行
        sheet["F2"] = "C"
        sheet["E3"] = "TX"
        sheet["F3"] = "RX"
        workbook.save(path)
    finally:
        workbook.close()

    workbook = load_workbook(path)
    try:
        flat = read_cases(workbook, horizontal)
    finally:
        workbook.close()

    vertical = _write(tmp_path, VERTICAL_YAML, name="vertical2.yaml")
    vpath = create_template(vertical, tmp_path / "vertical2.xlsx", cases=2, overwrite=True)
    _fill(vpath, {"Case1": {"port": "A", "mode": "TX"}, "Case2": {"port": "C", "mode": "RX"}})
    workbook = load_workbook(vpath)
    try:
        tall = read_cases(workbook, vertical)
    finally:
        workbook.close()

    assert [case.name for case in flat] == [case.name for case in tall]
    for left, right in zip(flat, tall, strict=True):
        assert {key: value.text for key, value in left.values.items()} == {
            key: value.text for key, value in right.values.items()
        }


# --------------------------------------------------------------------------- #
# 约束 / 表头检查
# --------------------------------------------------------------------------- #
def test_constraint_message_points_at_the_row(workbook_path: Path, project: ProjectConfig) -> None:
    _fill(workbook_path, {"Case2": {"port": "Z"}})
    workbook = load_workbook(workbook_path)
    try:
        cases = read_cases(workbook, project)
    finally:
        workbook.close()
    with pytest.raises(ExcelError, match=r"第 3 行.*'port'"):
        check_value_constraints(project, {}, cases)


def test_required_sheets_accepts_vertical_header(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        check_required_sheets(workbook, project)
        workbook["Local Parameter"]["A1"] = "Variable"  # 横向布局的表头，纵向布局不认
        with pytest.raises(ExcelError, match="表头应为 'Case'"):
            check_required_sheets(workbook, project)
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 公式模式
# --------------------------------------------------------------------------- #
def test_vertical_formula_references(project: ProjectConfig) -> None:
    """纵向布局的引用形态：二维 INDEX + 按变量名在第 1 行 MATCH。"""
    template = project.templates[0]
    lines = compile_formulas(template, project, case_axes=[2])[0]
    joined = "\n".join(lines)
    assert "INDEX('Local Parameter'!$A:$ZZ,2,MATCH(\"port\",'Local Parameter'!$1:$1,0))" in joined
    # Local 表按列定位的老形态不该再出现（Global 表仍然是变量做行，会有 $A:$A）
    assert "'Local Parameter'!$A:$A" not in joined
    assert "'Local Parameter'!$E:$E" not in joined
    assert "GPIO" in joined  # prefix 直接来自 YAML
    assert "_PORT" in joined


def test_local_cell_shapes() -> None:
    assert (
        local_cell("Local Parameter", "port", column=5)
        == "INDEX('Local Parameter'!$E:$E,MATCH(\"port\",'Local Parameter'!$A:$A,0))"
    )
    assert (
        local_cell("Local Parameter", "port", row=2)
        == "INDEX('Local Parameter'!$A:$ZZ,$2,MATCH(\"port\",'Local Parameter'!$1:$1,0))"
    )
    assert (
        local_cell("Local Parameter", "port", row=2, absolute=False)
        == "INDEX('Local Parameter'!$A:$ZZ,2,MATCH(\"port\",'Local Parameter'!$1:$1,0))"
    )
    with pytest.raises(Exception, match="要给 column 或 row"):
        local_cell("Local Parameter", "port")


def test_vertical_formula_values_equal_python_render(workbook_path: Path, project: ProjectConfig) -> None:
    """输出侧公式算出来的文本 == render_all() 渲染的文本（输入纵向 + 输出横向 / 纵向）。"""
    _fill(
        workbook_path,
        {
            "Case1": {"port": "A", "mode": "TX"},
            "Case2": {"port": "C", "mode": "RX"},
        },
    )
    expected = render_all(project, workbook_path)
    write_results(workbook_path, project, expected.results)

    workbook = load_workbook(workbook_path)
    try:
        for template in project.templates:
            results = expected.results[template.name]
            got = evaluate_template_values(workbook, project, template, [result.case_name for result in results])
            for result in results:
                assert got[result.case_name] == result.lines, (template.name, result.case_name)
    finally:
        workbook.close()


def test_vertical_formula_values_follow_parameter_change(workbook_path: Path, project: ProjectConfig) -> None:
    """改某个 Case 那一行的参数（不重跑）时，公式算出来的值要跟着变。"""
    _fill(workbook_path, {"Case1": {"port": "A"}, "Case3": {"port": "C"}})
    write_results(workbook_path, project, render_all(project, workbook_path).results)

    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        sheet.cell(row=2, column=2, value="B")
        sheet.cell(row=4, column=2, value="A")
        workbook["Global Parameter"]["B2"] = 28800
        workbook.save(workbook_path)
    finally:
        workbook.close()

    workbook = load_workbook(workbook_path)
    try:
        template = project.templates[0]
        got = evaluate_template_values(workbook, project, template, ["Case1", "Case3"])
    finally:
        workbook.close()

    assert got["Case1"][0] == "UART_Init(28800, GPIOB_PORT, 2);"
    assert got["Case3"][1] == "// A | GPIO | _PORT"

    fresh = render_all(project, workbook_path)
    assert got["Case1"] == fresh.results["uart_init"][0].lines
    assert got["Case3"] == fresh.results["uart_init"][2].lines


def test_vertical_inserted_row_is_picked_up(workbook_path: Path, project: ProjectConfig) -> None:
    """在末尾下拉一行 = 多一个工况（纵向布局的"右拉复制"）。"""
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        sheet.cell(row=5, column=1, value="Case4")
        sheet.cell(row=5, column=2, value="B")
        workbook.save(workbook_path)
    finally:
        workbook.close()

    workbook = load_workbook(workbook_path)
    try:
        cases = read_cases(workbook, project)
    finally:
        workbook.close()
    assert [case.name for case in cases] == ["Case1", "Case2", "Case3", "Case4"]
    assert cases[-1].values["port"].text == "B"
