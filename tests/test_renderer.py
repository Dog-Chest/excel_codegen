"""renderer 模块测试：VarValue 语义、过滤器、上下文组装、渲染与导出。"""

from __future__ import annotations

from pathlib import Path

import pytest

from openpyxl import load_workbook

from excel_codegen.excel_io import create_template, write_results
from excel_codegen.models import ProjectConfig, RenderResult, TemplateDef, load_config
from excel_codegen.renderer import (
    build_context,
    build_environment,
    case_matches,
    collect_variables,
    compile_case_filter,
    export_files,
    pvs,
    render_all,
    render_template,
    validate_template,
    wrap,
)
from excel_codegen.utils import ExcelError, RenderError, VarValue

EXAMPLES_DIR = Path(__file__).resolve().parents[1] / "examples"


def make_template(code: str, **overrides: object) -> TemplateDef:
    data = {"name": "tpl", "code": code}
    data.update(overrides)
    return TemplateDef.model_validate(data)


# --------------------------------------------------------------------------- #
# VarValue：Prefix + Value + Suffix
# --------------------------------------------------------------------------- #
def test_var_value_composition() -> None:
    value = VarValue("A", "GPIO", "_PORT")
    assert str(value) == "GPIOA_PORT"
    assert value.value == "A"
    assert value.prefix == "GPIO"
    assert value.suffix == "_PORT"
    assert value.text == "A"
    assert value.as_dict() == {"value": "A", "prefix": "GPIO", "suffix": "_PORT"}


def test_var_value_number_normalisation() -> None:
    assert str(VarValue(115200.0)) == "115200"
    assert str(VarValue(115200)) == "115200"
    assert str(VarValue(1.5)) == "1.5"
    assert str(VarValue(True)) == "true"
    assert str(VarValue(None)) == ""


def test_var_value_empty_and_truthiness() -> None:
    assert VarValue("").is_empty is True
    assert bool(VarValue("")) is False
    assert bool(VarValue("", "GPIO", "")) is True  # 有前缀即视为有内容
    assert bool(VarValue(0)) is True  # "0" 是有效值
    assert VarValue("x", None, None).prefix == ""


def test_pvs_and_wrap_filters() -> None:
    assert pvs("A", "GPIO", "_PORT") == "GPIOA_PORT"
    assert pvs(115200, "BAUD_", "U") == "BAUD_115200U"
    assert pvs("A") == "A"
    assert wrap("A", "GPIO", "_PORT") == "GPIOA_PORT"
    assert wrap("A") == pvs("A")


def test_filters_registered_in_environment() -> None:
    environment = build_environment()
    template = environment.from_string('{{ "A" | pvs("GPIO", "_PORT") }}/{{ "A" | wrap("X", "Y") }}')
    assert template.render() == "GPIOA_PORT/XAY"


# --------------------------------------------------------------------------- #
# 上下文
# --------------------------------------------------------------------------- #
def test_build_context_merges_global_local_and_case_name() -> None:
    globals_ = {"baud": VarValue(115200)}
    locals_ = {"port": VarValue("A", "GPIO", "_PORT")}
    context = build_context(globals_, locals_, "Case1")

    assert context["case_name"] == "Case1"
    assert str(context["baud"]) == "115200"
    assert str(context["port"]) == "GPIOA_PORT"


def test_build_context_local_overrides_global() -> None:
    context = build_context({"x": VarValue("global")}, {"x": VarValue("local")}, "Case1")
    assert str(context["x"]) == "local"


# --------------------------------------------------------------------------- #
# 渲染
# --------------------------------------------------------------------------- #
def test_render_template_prefix_value_suffix_variants() -> None:
    template = make_template(
        "{{ port }}|{{ port.value }}|{{ port.prefix }}|{{ port.suffix }}|{{ baud }}"
    )
    context = {
        "port": VarValue("A", "GPIO", "_PORT"),
        "baud": VarValue(115200.0),
    }
    assert render_template(template, context) == "GPIOA_PORT|A|GPIO|_PORT|115200"


def test_render_template_missing_variable_raises_render_error() -> None:
    template = make_template("{{ nope }}")
    with pytest.raises(RenderError, match="变量缺失"):
        render_template(template, {"baud": VarValue(1)})


def test_render_template_syntax_error_raises_render_error() -> None:
    template = make_template("line1\n{% if x %}\nline3")
    with pytest.raises(RenderError, match="语法错误"):
        render_template(template, {})


def test_render_template_from_file(tmp_path: Path) -> None:
    tpl_file = tmp_path / "tpl.c.j2"
    tpl_file.write_text("value={{ v }}", encoding="utf-8")
    template = TemplateDef.model_validate({"name": "file_tpl", "template_file": tpl_file.name})

    assert render_template(template, {"v": VarValue("x")}, base_dir=tmp_path) == "value=x"
    validate_template(template, base_dir=tmp_path)

    missing = TemplateDef.model_validate({"name": "missing", "template_file": "nope.j2"})
    with pytest.raises(RenderError, match="不存在"):
        validate_template(missing, base_dir=tmp_path)


def test_validate_and_collect_variables() -> None:
    template = make_template("{{ baud }} {{ port }}{{ case_name }}")
    validate_template(template)
    used = collect_variables(template)
    assert {"baud", "port", "case_name"} <= used

    with pytest.raises(RenderError, match="语法错误"):
        validate_template(make_template("{% for %}{% endfor %}"))


# --------------------------------------------------------------------------- #
# 端到端
# --------------------------------------------------------------------------- #
def test_render_all_end_to_end(workbook_path: Path, project: ProjectConfig) -> None:
    output = render_all(project, workbook_path)

    assert len(output.cases) == 2
    assert output.total() == 4  # 2 模板 × 2 Case

    first = output.results["uart_init"][0]
    assert first.case_name == "Case1"
    assert first.text.splitlines() == [
        "// Case: Case1",
        "UART_Init(115200, GPIOA_PORT, MODE_TX_RX);",
        "// 纯值: A, 前缀: GPIO, 后缀: _PORT",
        "// 过滤器: GPIOA_PORT",
        "// MCU: STM32F103",
    ]
    assert output.results["uart_init"][1].case_name == "Case2"
    assert "| Case1 | 115200 | GPIOA_PORT | MODE_TX_RX |" in output.results["uart_summary"][0].text
    assert "### uart_init · Case1" in output.preview()


def test_render_all_only_cases(workbook_path: Path, project: ProjectConfig) -> None:
    output = render_all(project, workbook_path, only_cases=["Case2"])
    assert [case.name for case in output.cases] == ["Case2"]
    assert output.results["uart_init"][0].case_name == "Case2"

    with pytest.raises(ExcelError, match="不存在这些 Case"):
        render_all(project, workbook_path, only_cases=["Case9"])


def test_render_all_uses_user_filled_values(
    workbook_path: Path, project: ProjectConfig
) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path)
    try:
        workbook["Global Parameter"]["B2"] = 9600
        sheet = workbook["Local Parameter"]
        sheet.cell(row=2, column=5, value="B")
        sheet.cell(row=3, column=5, value="MODE_RX")
        workbook.save(workbook_path)
    finally:
        workbook.close()

    output = render_all(project, workbook_path)
    text = output.results["uart_init"][0].text
    assert "UART_Init(9600, GPIOB_PORT, MODE_RX);" in text


def test_render_all_reports_missing_excel(tmp_path: Path, project: ProjectConfig) -> None:
    with pytest.raises(ExcelError, match="不存在"):
        render_all(project, tmp_path / "missing.xlsx")


# --------------------------------------------------------------------------- #
# 导出
# --------------------------------------------------------------------------- #
def test_export_files_uses_filename_pattern(
    tmp_path: Path, project: ProjectConfig, workbook_path: Path
) -> None:
    output = render_all(project, workbook_path, only_cases=["Case1"])
    written = export_files(project, output.results, tmp_path / "generated")

    names = sorted(path.name for path in written)
    assert names == ["uart_init_Case1.c", "uart_summary_Case1.md"]
    assert (tmp_path / "generated" / "uart_init_Case1.c").read_text(encoding="utf-8").startswith(
        "// Case: Case1"
    )


def test_export_files_default_naming_and_overwrite_guard(
    tmp_path: Path, config_text: str
) -> None:
    config_path = tmp_path / "c.yaml"
    config_path.write_text(config_text, encoding="utf-8")
    config = load_config(config_path)
    config.templates[0].filename = None
    config.templates[0].extension = ".h"

    results = {"uart_init": [RenderResult("uart_init", "Case1", "hello")]}

    written = export_files(config, results, tmp_path / "gen")
    assert written[0].name == "uart_init_Case1.h"
    assert written[0].read_text(encoding="utf-8") == "hello\n"

    with pytest.raises(RenderError, match="已存在"):
        export_files(config, results, tmp_path / "gen", overwrite=False)

    # overwrite=True 时允许覆盖
    assert export_files(config, results, tmp_path / "gen")[0].exists()


# --------------------------------------------------------------------------- #
# 仓库自带的示例配置
# --------------------------------------------------------------------------- #
def test_shipped_example_config_renders(tmp_path: Path) -> None:
    config_path = EXAMPLES_DIR / "example.yaml"
    assert config_path.exists(), "examples/example.yaml 应该随仓库提供"

    config = load_config(config_path)
    assert len(config.templates) == 2

    environment = build_environment()
    for template in config.templates:
        validate_template(template, env=environment, base_dir=config.source_dir)

    excel_path = create_template(config, tmp_path / "template.xlsx", cases=2, overwrite=True)
    output = render_all(config, excel_path)

    text = output.results["uart_init"][0].text
    assert "UART_Init(115200, GPIOA_PORT, MODE_TX_RX);" in text
    assert "纯值: A, 前缀: GPIO, 后缀: _PORT" in text
    assert "过滤器: GPIOA_PORT" in text
    assert "MCU: STM32F103" in text


# --------------------------------------------------------------------------- #
# 回归：FINDINGS.md 里逐条修复的问题
# --------------------------------------------------------------------------- #
def test_value_attribute_is_normalised_like_combined(tmp_path: Path, config_text: str) -> None:
    """FINDINGS #4：`{{ x.value }}` 必须和 `{{ x }}` 一样，整数浮点不带 .0。"""
    config_path = tmp_path / "value.yaml"
    config_path.write_text(
        config_text.replace(
            "      // MCU: {{ mcu }}",
            "      // value: {{ baud.value }} / {{ baud }} / {{ mcu.value }}",
        ),
        encoding="utf-8",
    )
    project = load_config(config_path)
    excel_path = create_template(project, tmp_path / "t.xlsx", cases=1)
    workbook = load_workbook(excel_path)
    try:
        workbook["Global Parameter"]["B2"] = 340.0  # 浮点写进 Excel
        workbook.save(excel_path)
    finally:
        workbook.close()

    output = render_all(project, excel_path)
    line = [item for item in output.results["uart_init"][0].lines if item.startswith("// value:")][0]
    assert line == "// value: 340 / 340 / STM32F103"

    # 直接测环境：bool / None / 大数 的行为也是确定的
    environment = build_environment()
    assert environment.from_string("{{ 340.0 }}|{{ true }}|{{ none }}").render() == "340|true|"
    assert environment.from_string("{{ 1e20 }}").render() == "1e+20"


TWOSETS_YAML = """\
version: 1

excel:
  output: "twosets.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Code EXT", "Code INT"]

variables:
  global:
    - name: L
      description: "船长 —— 两套规则共用"
      default: 340
  local:
    - name: kind
      description: "这个 Case 属于哪一套规则"
      default: "EXT"
      type: string

templates:
  - name: ext_code
    output_sheet: "Code EXT"
    start_cell: "B2"
    direction: "horizontal"
    case_filter: "kind == 'EXT'"
    code: |
      // EXT {{ case_name }} (L={{ L }})
  - name: int_code
    output_sheet: "Code INT"
    start_cell: "B2"
    direction: "horizontal"
    case_filter: "kind == 'INT'"
    code: |
      // INT {{ case_name }} (L={{ L }})
"""


def test_case_filter_limits_template_to_matching_cases(tmp_path: Path) -> None:
    """FINDINGS #3(b)：per-template 的 Case 过滤，一本工作簿放两套规则。"""
    config_path = tmp_path / "twosets.yaml"
    config_path.write_text(TWOSETS_YAML, encoding="utf-8")
    project = load_config(config_path)
    excel_path = create_template(project, tmp_path / "twosets.xlsx", cases=2)

    workbook = load_workbook(excel_path)
    try:
        sheet = workbook["Local Parameter"]
        row_of = {sheet.cell(row=row, column=1).value: row for row in range(2, sheet.max_row + 1)}
        sheet.cell(row=1, column=5, value="case_ext")
        sheet.cell(row=1, column=6, value="case_int")
        sheet.cell(row=row_of["kind"], column=5, value="EXT")
        sheet.cell(row=row_of["kind"], column=6, value="INT")
        workbook.save(excel_path)
    finally:
        workbook.close()

    output = render_all(project, excel_path)
    assert [r.case_name for r in output.results["ext_code"]] == ["case_ext"]
    assert [r.case_name for r in output.results["int_code"]] == ["case_int"]
    assert output.results["ext_code"][0].lines == ["// EXT case_ext (L=340)"]
    assert output.skipped == {"ext_code": ["case_int"], "int_code": ["case_ext"]}
    assert output.skipped_total() == 2

    write_results(excel_path, project, output.results)
    workbook = load_workbook(excel_path)
    try:
        ext_sheet, int_sheet = workbook["Code EXT"], workbook["Code INT"]
        assert ext_sheet["B1"].value == "case_ext"
        assert ext_sheet["B2"].value == "// EXT case_ext (L=340)"
        assert ext_sheet["C1"].value is None  # 被过滤掉的 Case 不占列
        assert int_sheet["B1"].value == "case_int"
        assert int_sheet["B2"].value == "// INT case_int (L=340)"
    finally:
        workbook.close()


def test_case_filter_errors_are_reported() -> None:
    """case_filter 的语法错误 / 变量缺失都要给出明确的 RenderError。"""
    bad_syntax = TemplateDef.model_validate({"name": "tpl", "code": "x", "case_filter": "kind =="})
    with pytest.raises(RenderError, match="case_filter 语法错误"):
        compile_case_filter(bad_syntax)

    missing = TemplateDef.model_validate(
        {"name": "tpl", "code": "x", "case_filter": "nope == 'EXT'"}
    )
    expression = compile_case_filter(missing)
    with pytest.raises(RenderError, match="case_filter 求值失败"):
        case_matches(expression, {"case_name": "Case1"}, template_name="tpl")

    # case_filter 用到的变量也算"被引用"，validate 不应把它报成未使用
    assert "nope" in collect_variables(missing)
    # 空表达式一律通过
    assert case_matches(None, {"case_name": "Case1"}, template_name="tpl") is True
