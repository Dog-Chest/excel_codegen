"""Excel 批注：把"这一格填什么"写在变量名那一格上。

批注只加在 **A 列（变量名）**：取值格已经有数据有效性的输入提示，而 A 列是冻结的、
永远可见 —— 鼠标一放就知道这是什么、单位是什么、有没有约束、模板里怎么引用。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen import create_template, load_config

CONFIG = """\
version: 1

excel:
  output: "t.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: draft
      description: "吃水"
      unit: "m"
      type: float
      default: 20.559
      suffix: " m"
      min: 0
      max: 50
    - name: half
      description: "半吃水"
      type: float
      derived: "draft / 2"
  local:
    - name: kind
      description: "规则集归属"
      choices: ["EXT", "INT"]
      default: "EXT"

templates:
  - name: demo
    output_sheet: "Output"
    code: |
      // {{ kind }} {{ draft }}
"""


@pytest.fixture()
def project(tmp_path: Path):
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return load_config(path)


def comments_of(path: Path) -> dict[str, str]:
    """{工作表!单元格: 批注正文}。"""
    book = load_workbook(path)
    try:
        return {
            f"{sheet.title}!{cell.coordinate}": cell.comment.text
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.comment is not None
        }
    finally:
        book.close()


def test_comments_are_written_on_the_name_column(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=2, overwrite=True)
    found = comments_of(path)
    assert set(found) == {"Global Parameter!A2", "Global Parameter!A3", "Local Parameter!A2"}


def test_comment_carries_description_unit_and_constraints(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True)
    text = comments_of(path)["Global Parameter!A2"]
    assert "吃水" in text
    assert "单位：m" in text
    assert "类型：float" in text
    assert "约束：范围: 0 ~ 50" in text
    assert "前缀 / 后缀" in text and "' m'" in text
    assert "默认值：20.559" in text
    assert "{{ draft }}" in text  # 模板里怎么引用
    assert "Global 表 B 列" in text  # 填哪儿


def test_comment_for_derived_variable_says_auto(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True)
    text = comments_of(path)["Global Parameter!A3"]
    assert "自动计算：draft / 2" in text
    assert "不用手填" in text
    assert "填写位置" not in text  # 派生参数不用填


def test_comment_for_choices_lists_them(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True)
    text = comments_of(path)["Local Parameter!A2"]
    assert "可选: EXT / INT" in text
    assert "Local 表 E 列起" in text


def test_comments_can_be_switched_off(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True, include_comments=False)
    assert comments_of(path) == {}


def test_author_is_the_tool(project, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True)
    book = load_workbook(path)
    try:
        assert book["Global Parameter"]["A2"].comment.author == "excel_codegen"
    finally:
        book.close()


def test_unit_is_documentation_only(project) -> None:
    """unit 只写进批注与清单，**不会**出现在生成的文本里（要进文本请用 suffix）。"""
    assert project.global_variables[0].unit == "m"
    assert project.global_variables[0].suffix == " m"
