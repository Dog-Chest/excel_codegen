"""公式引擎测试：engine: excel 的编译、边界、写回与 check 语义。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen.excel_io import create_template, write_results
from excel_codegen.formula import FormulaError, compile_formulas, compile_line
from excel_codegen.models import ProjectConfig, RenderResult, load_config
from excel_codegen.renderer import render_all

FORMULA_YAML = """\
version: 1

excel:
  output: "formula.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output", "Output Vertical"]

variables:
  global:
    - name: baud
      description: "波特率"
      default: 115200
      type: int
    - name: mcu
      default: "STM32F103"
  local:
    - name: port
      default: "A"
      prefix: "GPIO"
      suffix: "_PORT"
    - name: mode
      default: "TX_RX"
      prefix: "MODE_"

templates:
  - name: uart_init
    output_sheet: "Output"
    start_cell: "B2"
    direction: "horizontal"
    engine: "excel"
    code: |
      // Case: {{ case_name }}
      UART_Init({{ baud }}, {{ port }}, {{ mode }});
      // 纯值: {{ port.value }} / {{ port.text }}, 前缀: {{ port.prefix }}, 后缀: {{ port.suffix }}
      // 常量行 with "quotes"
      // 模板: {{ template_name }}

  - name: pins
    output_sheet: "Output Vertical"
    start_cell: "B2"
    direction: "vertical"
    engine: "excel"
    code: |
      {{ case_name }},{{ port }}
"""


@pytest.fixture()
def formula_config(tmp_path: Path) -> ProjectConfig:
    path = tmp_path / "formula.yaml"
    path.write_text(FORMULA_YAML, encoding="utf-8")
    return load_config(path)


def _first_case_formulas(config: ProjectConfig) -> list[str]:
    template = config.templates[0]
    return compile_formulas(template, config, case_columns=[5])[0]


# --------------------------------------------------------------------------- #
# 编译
# --------------------------------------------------------------------------- #
def test_compile_line_uses_cell_references(formula_config: ProjectConfig) -> None:
    line = "UART_Init({{ baud }}, {{ port }}, {{ mode }});"
    formula = compile_line(line, config=formula_config, template_name="uart_init", case_column=5)

    assert formula.startswith('="UART_Init("')
    assert formula.endswith('&");"')
    # 全局 int：走 INDEX/MATCH；**不套 TEXT()** —— Excel 的 TEXT() 会把二进制尾巴
    # （20.559000000000001）打出来，而隐式转换走 General（最多 15 位有效数字），
    # 正好与 Python 侧 utils.to_text 的规则一致
    assert "INDEX('Global Parameter'!$B:$B,MATCH(\"baud\",'Global Parameter'!$A:$A,0))" in formula
    assert "TEXT(" not in formula
    # 局部：值列是相对列（横向布局右拉换 Case），前缀/后缀是绝对列
    assert "INDEX('Local Parameter'!E:E,MATCH(\"port\"" in formula
    assert "INDEX('Local Parameter'!$C:$C,MATCH(\"port\"" in formula
    assert "INDEX('Local Parameter'!$D:$D,MATCH(\"port\"" in formula
    # 空单元格保护：INDEX 对空单元格返回 0，必须用 ISBLANK
    assert formula.count("ISBLANK(") >= 6


def test_compile_defaults_are_inlined(formula_config: ProjectConfig) -> None:
    formula = compile_line(
        "{{ port }}", config=formula_config, template_name="t", case_column=5
    )
    assert '"GPIO"' in formula and '"_PORT"' in formula  # prefix / suffix 回落
    assert '"A"' in formula  # value 回落


def test_compile_case_name_and_template_name(formula_config: ProjectConfig) -> None:
    formula = compile_line(
        "{{ case_name }}/{{ template_name }}",
        config=formula_config,
        template_name="uart_init",
        case_column=6,
    )
    assert formula == "='Local Parameter'!F$1&\" /\"".replace(" ", "") or "\"uart_init\"" in formula
    assert "'Local Parameter'!F$1" in formula
    assert '"uart_init"' in formula


def test_compile_vertical_locks_case_column(formula_config: ProjectConfig) -> None:
    vertical = formula_config.templates[1]
    lines = compile_formulas(vertical, formula_config, case_columns=[5])[0]
    assert "'Local Parameter'!$E$1" in lines[0]  # 纵向：Case 表头锁死
    assert "INDEX('Local Parameter'!$E:$E," in lines[0]  # 值列也锁死

    horizontal = formula_config.templates[0]
    lines = compile_formulas(horizontal, formula_config, case_columns=[5])[0]
    assert "'Local Parameter'!E$1" in lines[0]  # 横向：相对列，右拉换 Case


def test_constant_line_and_quote_escaping(formula_config: ProjectConfig) -> None:
    formulas = _first_case_formulas(formula_config)
    # 常量行也写成公式（避免以 = 开头的文本被 Excel 当公式）
    assert formulas[3] == '="// 常量行 with ""quotes"""'
    assert "// 模板: " in formulas[4] and '"uart_init"' in formulas[4]


def test_dot_value_and_text_are_the_same(formula_config: ProjectConfig) -> None:
    def compile_one(code: str) -> str:
        template = formula_config.templates[0].model_copy(update={"code": code})
        return compile_formulas(template, formula_config, case_columns=[5])[0][0]

    assert compile_one("{{ port.value }}") == compile_one("{{ port.text }}")
    # 值列 + ISBLANK 保护：一次取空判断、两次取值
    assert compile_one("{{ port.value }}").count("INDEX('Local Parameter'!E:E,") == 3


@pytest.mark.parametrize(
    "code, keyword",
    [
        # 跨行的 {% if %}：一行模板 = 一个单元格，所以必须整段写在同一行内
        ("{% if baud %}\nX\n{% endif %}", "没有闭合"),
        ("{% for x in [1] %}\nX\n{% endfor %}", "不支持"),
        ("{{ port | upper }}", "过滤器"),
        ("{{ baud + 1 }}", "只支持"),
        ("{{ port.name }}", "不支持属性"),
        ("{{ nope }}", "未在 YAML 中定义"),
    ],
)
def test_unsupported_templates_report_line(
    formula_config: ProjectConfig, code: str, keyword: str
) -> None:
    """超出"纯替换"子集的写法必须报错，并带上出错的行内容。"""
    template = formula_config.templates[0]
    bad = template.model_copy(update={"code": code})
    with pytest.raises(FormulaError) as excinfo:
        compile_formulas(bad, formula_config, case_columns=[5])
    message = str(excinfo.value)
    assert keyword in message
    assert "行内容" in message
    assert code.splitlines()[0].strip()[:8] in message


def test_formula_length_guard(formula_config: ProjectConfig) -> None:
    template = formula_config.templates[0]
    bad = template.model_copy(update={"code": "{{ baud }}" * 900})
    with pytest.raises(FormulaError, match="公式过长"):
        compile_formulas(bad, formula_config, case_columns=[5])


# --------------------------------------------------------------------------- #
# 写回
# --------------------------------------------------------------------------- #
def test_write_results_writes_formulas(formula_config: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(formula_config, tmp_path / "formula.xlsx", cases=2)
    output = render_all(formula_config, path)
    write_results(path, formula_config, output.results)

    workbook = load_workbook(path)
    try:
        # 公式模式：Excel 打开即重算
        assert workbook.calculation.fullCalcOnLoad is True
        sheet = workbook["Output"]
        assert sheet["B1"].value == "Case1"  # Case 表头仍是文本
        assert sheet["B2"].data_type == "f"
        assert str(sheet["B2"].value).startswith('="// Case: "')
        assert sheet["C2"].data_type == "f"
        assert "F$1" in str(sheet["C2"].value)  # 第二个 Case 指向 F 列

        vertical = workbook["Output Vertical"]
        assert vertical["A2"].value == "Case1"
        assert vertical["B2"].data_type == "f"
    finally:
        workbook.close()

    # 快照模式（同一个模板改回 snapshot）写的是文本
    snapshot = formula_config.templates[0].model_copy(update={"engine": "snapshot"})
    formulas_config = formula_config.model_copy(
        update={"templates": [snapshot, formula_config.templates[1]]}
    )
    text_output = render_all(formulas_config, path)
    write_results(path, formulas_config, text_output.results)
    workbook = load_workbook(path)
    try:
        assert workbook["Output"]["B2"].data_type == "s"
        assert workbook["Output"]["B2"].value == "// Case: Case1"
    finally:
        workbook.close()


def test_python_render_matches_formula_inputs(formula_config: ProjectConfig, tmp_path: Path) -> None:
    """公式模式不影响 Python 侧渲染：导出文件/指纹仍然来自 Jinja 渲染。"""
    path = create_template(formula_config, tmp_path / "formula.xlsx", cases=1)
    output = render_all(formula_config, path)
    assert output.results["uart_init"][0].lines[1] == "UART_Init(115200, GPIOA_PORT, MODE_TX_RX);"
    # 每条渲染行都对应一条公式
    formulas = compile_formulas(formula_config.templates[0], formula_config, case_columns=[5])[0]
    assert len(formulas) == len(output.results["uart_init"][0].lines)
