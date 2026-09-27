"""excel_io 模块测试：模板生成、读取填写内容、写回渲染结果。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen.excel_io import (
    GLOBAL_HEADERS,
    LOCAL_HEADERS,
    check_required_sheets,
    create_template,
    input_fingerprint,
    load_workbook_file,
    output_fingerprint,
    read_cases,
    read_global_values,
    write_results,
)
from excel_codegen.models import FIRST_CASE_COLUMN, ProjectConfig, RenderResult, load_config
from excel_codegen.utils import ConfigError, ExcelError, VarValue


# --------------------------------------------------------------------------- #
# 生成模板
# --------------------------------------------------------------------------- #
def test_create_template_layout(tmp_path: Path, project: ProjectConfig) -> None:
    path = create_template(project, tmp_path / "template.xlsx", cases=3, overwrite=True)

    assert path.exists()
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == [
            "HOWTO",
            "Global Parameter",
            "Local Parameter",
            "Output",
            "Output Vertical",
            "Template",
        ]
        assert workbook["Template"].sheet_state == "hidden"
        # 隐藏的 Template 表第一行是"只读参考"提示（FINDINGS #8.4）
        assert "只读参考" in str(workbook["Template"]["A1"].value)
        # HOWTO 表放在第一张，写清"下一步跑什么"（FINDINGS #8.2）
        howto = workbook["HOWTO"]
        howto_text = "\n".join(
            str(cell.value) for (cell,) in howto.iter_rows(min_col=1, max_col=1) if cell.value
        )
        assert "怎么用" in howto_text
        assert "render" in howto_text and "--write-excel" in howto_text
        assert "check" in howto_text

        global_sheet = workbook["Global Parameter"]
        assert [global_sheet.cell(row=1, column=i).value for i in range(1, 6)] == list(GLOBAL_HEADERS)
        assert global_sheet["A2"].value == "baud"
        assert global_sheet["B2"].value == 115200
        assert global_sheet["C2"].value == "波特率"
        assert global_sheet["A3"].value == "mcu"
        assert global_sheet["B3"].value == "STM32F103"

        local_sheet = workbook["Local Parameter"]
        assert [local_sheet.cell(row=1, column=i).value for i in range(1, 5)] == list(LOCAL_HEADERS)
        assert local_sheet.cell(row=1, column=FIRST_CASE_COLUMN).value == "Case1"
        assert local_sheet.cell(row=1, column=FIRST_CASE_COLUMN + 1).value == "Case2"
        assert local_sheet.cell(row=1, column=FIRST_CASE_COLUMN + 2).value == "Case3"
        assert local_sheet["A2"].value == "port"
        assert local_sheet["C2"].value == "GPIO"
        assert local_sheet["D2"].value == "_PORT"
        # 每个 Case 列都预填默认值，用户可直接覆盖
        assert local_sheet.cell(row=2, column=FIRST_CASE_COLUMN).value == "A"
        assert local_sheet.cell(row=2, column=FIRST_CASE_COLUMN + 1).value == "A"

        # Template 隐藏表保存了模板原文
        template_text = "\n".join(
            str(row[0].value)
            for row in workbook["Template"].iter_rows(min_col=2, max_col=2)
            if row[0].value
        )
        assert "UART_Init" in template_text
        assert "uart_summary" in template_text
    finally:
        workbook.close()


def test_create_template_custom_case_names(tmp_path: Path, project: ProjectConfig) -> None:
    path = create_template(
        project, tmp_path / "t.xlsx", cases=["UART1", "UART2"], include_template_sheet=False
    )
    workbook = load_workbook(path)
    try:
        assert "Template" not in workbook.sheetnames
        local_sheet = workbook["Local Parameter"]
        assert local_sheet.cell(row=1, column=FIRST_CASE_COLUMN).value == "UART1"
        assert local_sheet.cell(row=1, column=FIRST_CASE_COLUMN + 1).value == "UART2"
    finally:
        workbook.close()


def test_create_template_refuses_overwrite(tmp_path: Path, project: ProjectConfig) -> None:
    create_template(project, tmp_path / "t.xlsx", cases=1)
    with pytest.raises(ExcelError, match="已存在"):
        create_template(project, tmp_path / "t.xlsx", cases=1)
    # --force 语义
    assert create_template(project, tmp_path / "t.xlsx", cases=1, overwrite=True).exists()


def test_create_template_rejects_invalid_case_count(project: ProjectConfig, tmp_path: Path) -> None:
    with pytest.raises(ExcelError, match="至少为 1"):
        create_template(project, tmp_path / "t.xlsx", cases=0)


# --------------------------------------------------------------------------- #
# 读取
# --------------------------------------------------------------------------- #
def test_read_global_values_blank_cell_falls_back_to_default(
    workbook_path: Path, project: ProjectConfig
) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Global Parameter"]
        sheet["B2"] = None  # 清空 baud 的取值
        sheet["B3"] = "STM32H743"
        values = read_global_values(workbook, project)

        assert values["baud"].value == 115200  # 回落 YAML 默认值
        assert values["mcu"].value == "STM32H743"
        assert str(values["baud"]) == "115200"
    finally:
        workbook.close()


def test_read_global_values_applies_prefix_suffix(project: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Global Parameter"]
        sheet["D2"] = "BAUD_"
        sheet["E2"] = "U"
        values = read_global_values(workbook, project)
        assert values["baud"].prefix == "BAUD_"
        assert values["baud"].suffix == "U"
        assert str(values["baud"]) == "BAUD_115200U"
    finally:
        workbook.close()


def test_read_cases_per_column(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        sheet.cell(row=2, column=FIRST_CASE_COLUMN, value="A")
        sheet.cell(row=2, column=FIRST_CASE_COLUMN + 1, value="C")
        sheet.cell(row=3, column=FIRST_CASE_COLUMN, value="MODE_TX")
        sheet.cell(row=3, column=FIRST_CASE_COLUMN + 1).value = None  # 清空 -> 回落默认值

        cases = read_cases(workbook, project)

        assert [case.name for case in cases] == ["Case1", "Case2"]
        assert [case.index for case in cases] == [1, 2]
        assert cases[0].values["port"].value == "A"
        assert str(cases[0].values["port"]) == "GPIOA_PORT"
        assert cases[1].values["port"].value == "C"
        assert str(cases[1].values["port"]) == "GPIOC_PORT"
        assert cases[0].values["mode"].value == "MODE_TX"
        assert cases[1].values["mode"].value == "MODE_TX_RX"  # 空单元格 -> YAML 默认值
    finally:
        workbook.close()


def test_read_cases_without_case_column_raises(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook["Local Parameter"]
        for column in range(FIRST_CASE_COLUMN, sheet.max_column + 1):
            sheet.cell(row=1, column=column).value = None
        with pytest.raises(ExcelError, match="没有 Case 列"):
            read_cases(workbook, project)
    finally:
        workbook.close()


def test_missing_sheet_reports_all_missing(workbook_path: Path, project: ProjectConfig) -> None:
    workbook = load_workbook(workbook_path)
    try:
        del workbook["Local Parameter"]
        del workbook["Output"]
        with pytest.raises(ExcelError) as excinfo:
            check_required_sheets(workbook, project)
        message = str(excinfo.value)
        assert "Local Parameter" in message
        assert "Output" in message
    finally:
        workbook.close()


def test_load_workbook_file_missing(tmp_path: Path) -> None:
    with pytest.raises(ExcelError, match="不存在"):
        load_workbook_file(tmp_path / "nope.xlsx")


def test_duplicate_variable_name_reported(project: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Global Parameter"]
        sheet["A4"] = "baud"
        sheet["B4"] = 9600
        with pytest.raises(ExcelError, match="重复"):
            read_global_values(workbook, project)
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 写入结果
# --------------------------------------------------------------------------- #
def test_write_results_horizontal(project: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    rendered = {
        "uart_init": [
            RenderResult("uart_init", "Case1", "line1\nline2"),
            RenderResult("uart_init", "Case2", "only"),
        ]
    }
    write_results(path, project, rendered)

    workbook = load_workbook(path)
    try:
        sheet = workbook["Output"]
        assert sheet["B1"].value == "Case1"  # Case 表头写在 start_cell 上一行
        assert sheet["C1"].value == "Case2"
        assert sheet["B2"].value == "line1"
        assert sheet["B3"].value == "line2"
        assert sheet["C2"].value == "only"
        assert sheet["C3"].value is None

        # 重新渲染成 1 个 Case，旧的第二列必须被清理
        write_results(path, project, {"uart_init": rendered["uart_init"][:1]})
        workbook2 = load_workbook(path)
        try:
            sheet2 = workbook2["Output"]
            assert sheet2["B2"].value == "line1"
            assert sheet2["C1"].value is None
            assert sheet2["C2"].value is None
        finally:
            workbook2.close()
    finally:
        workbook.close()


def test_write_results_vertical(project: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    rendered = {
        "uart_summary": [
            RenderResult("uart_summary", "Case1", "a|b|c"),
            RenderResult("uart_summary", "Case2", "d|e|f"),
        ]
    }
    write_results(path, project, rendered)

    workbook = load_workbook(path)
    try:
        sheet = workbook["Output Vertical"]
        assert sheet["A2"].value == "Case1"  # Case 名写在 start_cell 左侧一列
        assert sheet["B2"].value == "a|b|c"
        assert sheet["A3"].value == "Case2"
        assert sheet["B3"].value == "d|e|f"
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# 回归：FINDINGS.md 里逐条修复的问题
# --------------------------------------------------------------------------- #
def test_prefix_suffix_keep_leading_space(tmp_path: Path, config_text: str) -> None:
    """FINDINGS #2：suffix: " m" 不能被 strip 成 "m"（否则 `340 m` 变 `340m`）。"""
    config_path = tmp_path / "spaces.yaml"
    config_path.write_text(
        config_text.replace('suffix: "_PORT"', 'suffix: " m"'), encoding="utf-8"
    )
    project = load_config(config_path)
    path = create_template(project, tmp_path / "t.xlsx", cases=1)

    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        assert sheet["D2"].value == " m"  # YAML 里写了什么，单元格里就是什么
        values = read_global_values(workbook, project)
        cases = read_cases(workbook, project)
    finally:
        workbook.close()

    port = cases[0].values["port"]
    assert port.suffix == " m"
    assert str(port) == "GPIOA m"
    assert values["baud"].suffix == ""

    # 空格子仍然回落 YAML 默认值
    workbook = load_workbook(path)
    try:
        workbook["Local Parameter"]["D2"] = None
        cases = read_cases(workbook, project)
    finally:
        workbook.close()
    assert cases[0].values["port"].suffix == " m"  # 回落 YAML 的 " m"，不是 "m"


def test_rerender_clears_older_longer_output(project: ProjectConfig, tmp_path: Path) -> None:
    """FINDINGS #5 症状 A：8 行渲染之后用 3 行渲染覆盖，旧行不能残留。"""
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    long_text = "\n".join(f"LONG line {index}" for index in range(1, 9))
    write_results(path, project, {"uart_init": [RenderResult("uart_init", "Case1", long_text)]})

    short_text = "\n".join(f"short {index}" for index in range(1, 4))
    write_results(path, project, {"uart_init": [RenderResult("uart_init", "Case1", short_text)]})

    workbook = load_workbook(path)
    try:
        sheet = workbook["Output"]
        column = [sheet.cell(row=row, column=2).value for row in range(1, 12)]
        assert [value for value in column if value and "LONG" in str(value)] == []
        assert column[1:4] == ["short 1", "short 2", "short 3"]
        assert column[4] is None
    finally:
        workbook.close()


def test_vertical_rerender_clears_extra_cases(project: ProjectConfig, tmp_path: Path) -> None:
    """FINDINGS #5 症状 B：vertical 布局 8 Case 减到 1 个，旧 Case 行不能残留。"""
    path = create_template(project, tmp_path / "t.xlsx", cases=8)
    many = [
        RenderResult("uart_summary", f"Case{index}", f"line {index}") for index in range(1, 9)
    ]
    write_results(path, project, {"uart_summary": many})
    write_results(path, project, {"uart_summary": many[:1]})

    workbook = load_workbook(path)
    try:
        sheet = workbook["Output Vertical"]
        column = [sheet.cell(row=row, column=1).value for row in range(1, 12)]
        assert [value for value in column if value and value != "Case1"] == []
    finally:
        workbook.close()


def test_write_results_records_fingerprints(project: ProjectConfig, tmp_path: Path) -> None:
    """FINDINGS #8.3：写回时留下时间/参数指纹/输出指纹，且能被 read_metadata 读回。"""
    from excel_codegen.excel_io import read_metadata

    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    results = {"uart_init": [RenderResult("uart_init", "Case1", "a\nb")]}
    write_results(path, project, results, command="excel-codegen render --write-excel")

    workbook = load_workbook(path)
    try:
        metadata = read_metadata(workbook, project)
        expected_input = input_fingerprint(
            read_global_values(workbook, project), read_cases(workbook, project)
        )
        assert metadata["参数指纹"] == expected_input
        assert metadata["输出指纹"] == output_fingerprint(results)
        assert metadata["case 数"] == "1"
        assert metadata["cases"] == "Case1"
        assert metadata["时间"]
        # HOWTO 表也要有人读的那一份
        howto = "\n".join(
            str(cell.value) for (cell,) in workbook["HOWTO"].iter_rows(min_col=1, max_col=1)
            if cell.value
        )
        assert metadata["参数指纹"] in howto
        assert "尚未渲染" not in howto
    finally:
        workbook.close()

    # 参数改了 -> 指纹必须变（否则 check 检测不到过期）
    workbook = load_workbook(path)
    try:
        workbook["Global Parameter"]["B2"] = 9600
        workbook.save(path)
    finally:
        workbook.close()
    workbook = load_workbook(path)
    try:
        after = input_fingerprint(read_global_values(workbook, project), read_cases(workbook, project))
    finally:
        workbook.close()
    assert after != expected_input


def test_input_fingerprint_reacts_to_prefix_suffix(project: ProjectConfig, tmp_path: Path) -> None:
    """前缀/后缀也是参数：改了它，指纹必须变。"""
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        before = input_fingerprint(read_global_values(workbook, project), read_cases(workbook, project))
        workbook["Local Parameter"]["C2"] = "PORT_"
        workbook.save(path)
    finally:
        workbook.close()
    workbook = load_workbook(path)
    try:
        after = input_fingerprint(read_global_values(workbook, project), read_cases(workbook, project))
    finally:
        workbook.close()
    assert before != after


def test_howto_sheet_can_be_disabled(project: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(project, tmp_path / "t.xlsx", cases=1, include_howto_sheet=False)
    workbook = load_workbook(path)
    try:
        assert "HOWTO" not in workbook.sheetnames
    finally:
        workbook.close()


def test_duplicate_yaml_key_is_rejected(tmp_path: Path, config_text: str) -> None:
    """FINDINGS #6：YAML 同一个键写两遍必须立刻报错，而不是静默丢一份。"""
    broken = config_text + "\nvariables:\n  local: []\n"
    config_path = tmp_path / "dup.yaml"
    config_path.write_text(broken, encoding="utf-8")

    with pytest.raises(ConfigError) as excinfo:
        load_config(config_path)
    message = str(excinfo.value)
    assert "YAML 键重复" in message
    assert "variables" in message


def test_read_metadata_howto_fallback_uses_same_keys(tmp_path: Path, config_text: str) -> None:
    """关掉 Template 表时，元信息只存在于 HOWTO 表 —— 键名必须和 Template 表路径一致。"""
    from excel_codegen.excel_io import read_metadata

    config_path = tmp_path / "howto_only.yaml"
    config_path.write_text(
        config_text.replace('template_sheet: "Template"', "template_sheet: null"),
        encoding="utf-8",
    )
    project = load_config(config_path)
    path = create_template(project, tmp_path / "t.xlsx", cases=1)
    write_results(path, project, {"uart_init": [RenderResult("uart_init", "Case1", "a")]})

    workbook = load_workbook(path)
    try:
        assert "Template" not in workbook.sheetnames
        metadata = read_metadata(workbook, project)
        assert set(metadata) >= {"时间", "参数指纹", "输出指纹"}
        assert metadata["参数指纹"] == input_fingerprint(
            read_global_values(workbook, project), read_cases(workbook, project)
        )
        assert metadata["输出指纹"] == output_fingerprint(
            {"uart_init": [RenderResult("uart_init", "Case1", "a")]}
        )
    finally:
        workbook.close()


def test_blank_global_cell_still_falls_back(workbook_path: Path, project: ProjectConfig) -> None:
    """回落语义没变：只有真空白才回落 YAML 默认值。"""
    workbook = load_workbook(workbook_path)
    try:
        workbook["Global Parameter"]["B2"] = None
        values = read_global_values(workbook, project)
        assert values["baud"].value == 115200
    finally:
        workbook.close()
