"""``allow_blank``：让"某类型才有的字段"能正当地留空。

要解决的问题（改进建议 §2.2）
----------------------------
`card_type` 只有银行卡那几行才有，但"声明了约束就不许为空"，于是只能往 ``choices``
里塞一个空串：``choices: ["", "储蓄卡", "信用卡"]``。副作用是 Excel 下拉列表第一项是空的
（读起来像"允许空"），而文档又说空值不合格 —— 两边说法冲突；而且这个绕法**没有任何文档提示**。

``allow_blank: true`` 直接表达"这个变量可以留空"，下拉列表里也不再出现那个空选项。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from spreadsheet_codegen.excel_io import check_value_constraints, create_template
from spreadsheet_codegen.models import CaseData, ProjectConfig, VariableDef, load_config
from spreadsheet_codegen.utils import VarValue

CARD_YAML = """\
version: 1

excel:
  output: "vault.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  local:
    - name: kind
      default: "银行卡"
      type: string
    - name: card_type
      description: "只有银行卡才有"
__EXTRA__

templates:
  - name: row
    output_sheet: "Output"
    start_cell: "B2"
    code: |
      {{ kind }}/{{ card_type }}
"""


def _config(tmp_path: Path, extra: str) -> ProjectConfig:
    path = tmp_path / "vault.yaml"
    path.write_text(CARD_YAML.replace("__EXTRA__", extra), encoding="utf-8")
    return load_config(path)


# --------------------------------------------------------------------------- #
# 判定层
# --------------------------------------------------------------------------- #
def test_default_still_rejects_blank() -> None:
    """默认行为不变：声明了约束，空值就不合格（这治的是"新插空列悄悄落进规则集"）。"""
    variable = VariableDef.model_validate({"name": "card_type", "choices": ["储蓄卡"]})
    assert variable.allow_blank is False
    assert variable.value_problem("") is not None


def test_allow_blank_accepts_blank_but_not_garbage() -> None:
    """``allow_blank`` 只放行空值，别的取值照旧要合规。"""
    variable = VariableDef.model_validate({"name": "card_type", "choices": ["储蓄卡"], "allow_blank": True})
    assert variable.value_problem("") is None
    assert variable.value_problem("   ") is None  # 只敲空格也算没填
    assert variable.value_problem("储蓄卡") is None
    assert variable.value_problem("白金卡") is not None


def test_allow_blank_applies_to_all_constraint_kinds() -> None:
    """``min`` / ``max`` / ``pattern`` 也认这个开关。"""
    ranged = VariableDef.model_validate({"name": "n", "type": "int", "min": 1, "max": 8, "allow_blank": True})
    assert ranged.value_problem("") is None
    assert ranged.value_problem("9") is not None

    patterned = VariableDef.model_validate({"name": "mcu", "pattern": "STM32.*", "allow_blank": True})
    assert patterned.value_problem("") is None
    assert patterned.value_problem("ATMEGA") is not None


def test_allow_blank_without_any_constraint_is_rejected() -> None:
    """没有约束时 ``allow_blank`` 没有意义（本来就没人拦空值）→ 配置期报错。"""
    with pytest.raises(ValueError, match="没有任何取值约束"):
        VariableDef.model_validate({"name": "x", "allow_blank": True})


def test_constraint_text_mentions_blank_and_hides_the_empty_option() -> None:
    """提示文字写清"可以留空"，且选项里不再有那个空串。"""
    variable = VariableDef.model_validate(
        {"name": "card_type", "choices": ["", "储蓄卡", "信用卡"], "allow_blank": True}
    )
    assert "可以留空" in variable.constraint_text
    assert not variable.constraint_text.startswith("可选:  /")


# --------------------------------------------------------------------------- #
# Excel 侧：下拉列表
# --------------------------------------------------------------------------- #
def test_dropdown_does_not_offer_a_blank_option(tmp_path: Path) -> None:
    """``choices`` 里的空串只是校验层的概念 —— 下拉列表里不该出现空选项。"""
    config = _config(tmp_path, '      choices: ["", "储蓄卡", "信用卡"]\n      allow_blank: true')
    path = create_template(config, tmp_path / "vault.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        validations = sheet.data_validations.dataValidation
        assert len(validations) == 1
        formula = validations[0].formula1
        assert formula == '"储蓄卡,信用卡"', f"下拉列表里不该有空的第一个选项：{formula}"
        # Excel 的下拉本身总允许留空；工具侧的判据才决定空值合不合法
        assert validations[0].allow_blank is True
    finally:
        workbook.close()


def test_choices_with_only_empty_strings_is_an_error(tmp_path: Path) -> None:
    """``choices: [""]`` 不是约束 —— 直接报错并指路 ``allow_blank``。"""
    config = _config(tmp_path, '      choices: [""]')
    with pytest.raises(Exception, match="allow_blank"):
        create_template(config, tmp_path / "vault.xlsx", cases=1)


# --------------------------------------------------------------------------- #
# 运行期：校验
# --------------------------------------------------------------------------- #
def test_check_value_constraints_passes_for_blank_when_allowed(tmp_path: Path) -> None:
    config = _config(tmp_path, '      choices: ["储蓄卡", "信用卡"]\n      allow_blank: true')
    cases = [
        CaseData(name="bank", column=5, values={"kind": VarValue("银行卡"), "card_type": VarValue("储蓄卡")}),
        CaseData(name="site", column=6, values={"kind": VarValue("网站登录"), "card_type": VarValue("")}),
    ]
    check_value_constraints(config, {}, cases)  # 不抛异常 = 通过


def test_check_value_constraints_still_rejects_blank_by_default(tmp_path: Path) -> None:
    config = _config(tmp_path, '      choices: ["储蓄卡", "信用卡"]')
    cases = [CaseData(name="site", column=6, values={"kind": VarValue("网站登录"), "card_type": VarValue("")})]
    with pytest.raises(Exception, match="不在允许列表"):
        check_value_constraints(config, {}, cases)


# --------------------------------------------------------------------------- #
# 与 default / fallback 的搭配（"可选字段"的完整写法）
# --------------------------------------------------------------------------- #
def test_optional_field_recipe(tmp_path: Path) -> None:
    """ "某类型才有的字段"的推荐写法：``allow_blank`` + ``fallback: false``。

    ``default`` 仍然预填进新表当提示；清空即表示"这条记录没有这个字段"；
    而且空值别被取值约束拦下。
    """
    config = _config(
        tmp_path,
        '      choices: ["储蓄卡", "信用卡"]\n      allow_blank: true\n      fallback: false\n      default: "储蓄卡"',
    )
    path = create_template(config, tmp_path / "vault.xlsx", cases=["bank", "site"])

    from openpyxl import load_workbook as _load

    from spreadsheet_codegen.renderer import render_all

    workbook = _load(path)
    try:
        sheet = workbook["Local Parameter"]
        row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(row=r, column=1).value == "card_type")
        assert sheet.cell(row=row, column=5).value == "储蓄卡", "default 仍然预填（当提示）"
        sheet.cell(row=row, column=6).value = None  # site 没有这个字段
        workbook.save(path)
    finally:
        workbook.close()

    output = render_all(config, path)
    texts = {result.case_name: result.text.strip() for result in output.results["row"]}
    assert texts["bank"] == "银行卡/储蓄卡"
    assert texts["site"] == "银行卡/", "清空之后就是空的，不回落默认值"


def test_legacy_empty_choice_still_works_and_is_flagged(tmp_path: Path) -> None:
    """老绕法（``choices`` 含空串、不写 ``allow_blank``）照常能跑，但会给一条改写提示。"""
    config = _config(tmp_path, '      choices: ["", "储蓄卡", "信用卡"]')
    from spreadsheet_codegen.renderer import render_all

    path = create_template(config, tmp_path / "vault.xlsx", cases=1)
    output = render_all(config, path)
    blank_ok = any("choices 里有空串" in warning and "allow_blank" in warning for warning in output.warnings)
    assert blank_ok, f"应当提示改用 allow_blank：{output.warnings}"
