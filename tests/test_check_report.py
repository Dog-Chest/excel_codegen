"""``check`` 的差异报告：列出**全部**不同行，以及给 CI 消费的 ``--json``。

只报第一处差异在 CI 里很难定位 —— 想知道"一共差多少、都差在哪"，
所以差异按行汇总（每处最多列 5 行，其余折成一句）。``--json`` 让下游脚本 / 看板能直接消费，
退出码与表格模式一致（过期 = 1），两种模式可以互换。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from spreadsheet_codegen import create_template, load_config, render_all
from spreadsheet_codegen.cli import _differences, app
from spreadsheet_codegen.excel_io import write_results

runner = CliRunner()

CONFIG = """\
version: 1

excel:
  output: "t.xlsx"
  template_sheet: "Template"
  howto_sheet: "HOWTO"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: n
      default: 1
  local:
    - name: tag
      default: "A"

templates:
  - name: demo
    output_sheet: "Output"
    start_cell: "B2"
    code: |
      line1 {{ n }}
      line2 {{ tag }}
      line3
      line4
      line5
      line6
      line7
"""


@pytest.fixture()
def project_path(tmp_path: Path) -> Path:
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return path


def make_fresh(project_path: Path) -> Path:
    project = load_config(project_path)
    excel = project_path.parent / "t.xlsx"
    create_template(project, excel, cases=1, overwrite=True, include_scripts=False)
    write_results(excel, project, render_all(project, excel).results)
    return excel


def damage(project_path: Path, *, rows: list[int]) -> Path:
    """把 Output 表 B 列的指定几行改成垃圾。"""
    excel = make_fresh(project_path)
    book = load_workbook(excel)
    try:
        for row in rows:
            book["Output"].cell(row=row, column=2, value=f"// broken {row}")
        book.save(excel)
    finally:
        book.close()
    return excel


# --------------------------------------------------------------------------- #
# 差异汇总
# --------------------------------------------------------------------------- #
def test_differences_lists_every_changed_line() -> None:
    got = _differences(["a", "X", "c", "Y"], ["a", "b", "c", "d"], limit=5)
    assert len(got) == 2
    assert "第 2 行不同" in got[0] and "第 4 行不同" in got[1]


def test_differences_caps_and_summarises() -> None:
    got = _differences(list("XXXXXXXXXX"), list("YYYYYYYYYY"), limit=3)
    assert len(got) == 4
    assert "还有 7 行不同" in got[-1]


def test_differences_reports_length_mismatch() -> None:
    assert "行数不同" in _differences(["a"], ["a", "b"])[0]
    assert "行数不同" in _differences(["a", "b"], ["a"])[0]


def test_differences_on_identical_input() -> None:
    assert _differences(["a"], ["a"], what="公式") == ["公式不同"]


# --------------------------------------------------------------------------- #
# CLI：表格模式报全部差异
# --------------------------------------------------------------------------- #
def test_check_lists_all_changed_lines(project_path: Path) -> None:
    excel = damage(project_path, rows=[4, 6, 8])
    result = runner.invoke(app, ["check", "-c", str(project_path), "-x", str(excel)])
    assert result.exit_code == 1
    # rich 会在中文之间折行，先把空白去掉再找
    text = "".join(((result.output or "") + (getattr(result, "stderr", "") or "")).split())
    # start_cell=B2 且写 Case 表头：Excel 第 3 行 = 模板第 2 行 → 破坏 4/6/8 行就是模板第 3/5/7 行
    for index in (3, 5, 7):
        assert f"第{index}行不同" in text


# --------------------------------------------------------------------------- #
# CLI：--json
# --------------------------------------------------------------------------- #
def test_check_json_is_pure_json_when_fresh(project_path: Path) -> None:
    excel = make_fresh(project_path)
    result = runner.invoke(app, ["check", "-c", str(project_path), "-x", str(excel), "--json"])
    assert result.exit_code == 0
    payload = json.loads(result.output)  # 能解析 = stdout 上没有混进表格
    assert payload["ok"] is True
    assert payload["problems"] == []
    assert payload["current"]["output_fingerprint"]


def test_check_json_reports_problems_and_exit_code(project_path: Path) -> None:
    excel = damage(project_path, rows=[5, 7])
    result = runner.invoke(app, ["check", "-c", str(project_path), "-x", str(excel), "--json"])
    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert len(payload["problems"]) == 1
    assert "第 4 行不同" in payload["problems"][0]
    assert "第 6 行不同" in payload["problems"][0]


def test_check_json_marks_parameter_drift(project_path: Path) -> None:
    excel = make_fresh(project_path)
    book = load_workbook(excel)
    try:
        book["Global Parameter"]["B2"] = 42  # 改参数但不重渲染
        book.save(excel)
    finally:
        book.close()
    result = runner.invoke(app, ["check", "-c", str(project_path), "-x", str(excel), "--json"])
    payload = json.loads(result.output)
    assert payload["drift"] is True
    assert payload["recorded"]["input_fingerprint"] != payload["current"]["input_fingerprint"]


def test_check_json_has_no_recorded_block_when_never_rendered(project_path: Path) -> None:
    project = load_config(project_path)
    excel = project_path.parent / "t.xlsx"
    create_template(project, excel, cases=1, overwrite=True, include_scripts=False)
    result = runner.invoke(app, ["check", "-c", str(project_path), "-x", str(excel), "--json"])
    payload = json.loads(result.output)
    assert payload["recorded"]["time"] == ""
    assert payload["drift"] is False
