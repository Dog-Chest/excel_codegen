"""内置示例（`spreadsheet_codegen/examples/`）与 `spreadsheet-codegen examples` 命令。

0.9.0 起示例是**包数据**：wheel 与 sdist 里都有一份，所以"装了 pip 包、没克隆仓库"的人
也能 `spreadsheet-codegen examples --copy ./demo` 拿到完整示例（含已填好样例参数的工作簿）。
这里守住三件事：

1. **清单不漂移**：`examples/` 下的目录与 `example_pack.EXAMPLES` 一一对应（新增目录忘了
   登记、或登记了却删了目录，都会红）；
2. **示例本身是好的**：每个示例的入口 YAML 能被 `load_config` 读、工作簿文件在；
3. **拷出来仍然能用**：`--copy` 出来的目录里 `template_file` 这类相对路径照样解析得开
   （这正是"自带示例"最容易坏的地方）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from spreadsheet_codegen.cli import app
from spreadsheet_codegen.example_pack import EXAMPLES, copy_examples, examples_root
from spreadsheet_codegen.models import load_config
from spreadsheet_codegen.utils import CodeGenError

runner = CliRunner()


def _output(result) -> str:
    return (result.output or "") + (getattr(result, "stderr", "") or "")


# --------------------------------------------------------------------------- #
# 1. 清单与实际目录一一对应
# --------------------------------------------------------------------------- #
def test_every_example_directory_is_registered() -> None:
    on_disk = sorted(item.name for item in examples_root().iterdir() if item.is_dir())
    assert sorted(item.name for item in EXAMPLES) == on_disk, (
        "examples/ 下的目录与 example_pack.EXAMPLES 不一致：两边要同步改"
    )


@pytest.mark.parametrize("example", EXAMPLES, ids=lambda item: item.name)
def test_example_files_exist_and_load(example) -> None:
    """入口 YAML 可解析、工作簿在 —— 示例不能是"半份"的。"""
    root = examples_root()
    yaml_path = root / example.entry
    book = root / example.workbook
    assert yaml_path.is_file(), f"缺入口 YAML：{example.entry}"
    assert book.is_file(), f"缺工作簿：{example.workbook}"

    config = load_config(yaml_path)
    assert config.templates, f"{example.name} 一个模板都没有"


def test_shipped_workbooks_carry_no_old_name() -> None:
    """工作簿里烘焙进去的文字也不能留着旧名字。

    HOWTO 表会告诉用户"跑 excel-codegen render …"，Template 表里还藏着机器读的元信息标记。
    改名时如果只改代码不改这些二进制，用户打开示例就会被指去一个**不存在**的命令。
    """
    stale = ("excel-codegen", "excel_codegen")
    offenders: list[str] = []
    for path in sorted(examples_root().rglob("*.xlsx")):
        workbook = load_workbook(path)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and any(token in cell.value for token in stale):
                            offenders.append(f"{path.name}:{worksheet.title}!{cell.coordinate}")
        finally:
            workbook.close()
    assert not offenders, f"随包工作簿里还残留旧名字：{offenders}"


# --------------------------------------------------------------------------- #
# 2. 列出
# --------------------------------------------------------------------------- #
def test_examples_lists_all_examples() -> None:
    result = runner.invoke(app, ["examples"])
    assert result.exit_code == 0, _output(result)
    text = _output(result)
    for example in EXAMPLES:
        assert example.name in text
        assert example.entry in text


# --------------------------------------------------------------------------- #
# 3. 复制
# --------------------------------------------------------------------------- #
def test_copy_all_examples_into_target(tmp_path: Path) -> None:
    dest = tmp_path / "demo"
    result = runner.invoke(app, ["examples", "--copy", str(dest)])
    assert result.exit_code == 0, _output(result)

    for example in EXAMPLES:
        assert (dest / example.name).is_dir(), f"没拷出 {example.name}"
        assert (dest / example.entry).is_file()
        assert (dest / example.workbook).is_file()


def test_copied_example_is_self_contained(tmp_path: Path) -> None:
    """拷出来的示例必须**自己就能用**：带 template_file 的相对路径要解析得开。"""
    dest = tmp_path / "demo"
    copy_examples(dest, only="abs_fpi")
    copied = dest / "abs_fpi" / "abs_fpi_internal.yaml"

    config = load_config(copied)
    using_files = [t for t in config.templates if t.template_file]
    assert using_files, "这个示例本来就该用 template_file（否则测不到相对路径）"
    for template in using_files:
        resolved = (copied.parent / template.template_file).resolve()
        assert resolved.is_file(), f"template_file 解析不到：{template.template_file}"


def test_copy_only_one_example(tmp_path: Path) -> None:
    dest = tmp_path / "demo"
    result = runner.invoke(app, ["examples", "--copy", str(dest), "--only", "nastran"])
    assert result.exit_code == 0, _output(result)
    assert (dest / "nastran").is_dir()
    assert not (dest / "basic").exists()


def test_copy_refuses_to_clobber_without_force(tmp_path: Path) -> None:
    dest = tmp_path / "demo"
    assert runner.invoke(app, ["examples", "--copy", str(dest)]).exit_code == 0

    again = runner.invoke(app, ["examples", "--copy", str(dest)])
    assert again.exit_code != 0
    assert "--force" in _output(again)

    forced = runner.invoke(app, ["examples", "--copy", str(dest), "--force"])
    assert forced.exit_code == 0, _output(forced)


def test_unknown_example_name_is_rejected(tmp_path: Path) -> None:
    result = runner.invoke(app, ["examples", "--copy", str(tmp_path / "demo"), "--only", "nope"])
    assert result.exit_code != 0
    assert "nope" in _output(result)
    # 报错要把可选项列出来，不然用户不知道写什么
    assert "basic" in _output(result)


def test_copy_examples_rejects_unknown_name(tmp_path: Path) -> None:
    with pytest.raises(CodeGenError, match="没有这个示例"):
        copy_examples(tmp_path / "never-written", only="nope")
