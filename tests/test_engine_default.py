"""默认引擎：**公式模式（`engine: excel`）**。

本工具的用法是"写一个 YAML → 生成一本 Excel → 之后就在 Excel 里干活"：
工作簿必须**独立可用** —— 改一格参数，输出自己就变了，不需要再跑命令。
所以模板不写 `engine` 时默认就是公式模式；快照模式要**显式**声明。

顺带钉住配套的两件事：

* 公式模式表达不了的模板（过滤器 / `{% for %}` / 多行 `{% if %}`）会在
  **建表时就报错**（`init` / `create_template`），而不是等你把参数填完再炸 ——
  报错里要明确写出"加 engine: snapshot"；
* 显式 `engine: snapshot` 的老行为不变（过滤器、循环、`{% include %}` 照常可用）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen import create_template, render_all, write_results
from excel_codegen.formula import FormulaError
from excel_codegen.models import ProjectConfig, load_config
from excel_codegen.utils import ConfigError

HEAD = """\
version: 1

excel:
  output: "t.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: L
      type: float
      default: 340
      suffix: " m"
  local:
    - name: draft
      type: float
      default: 20.5
      suffix: " m"
"""


def _write(tmp_path: Path, templates: str, name: str = "p.yaml") -> Path:
    path = tmp_path / name
    path.write_text(HEAD + "templates:\n" + templates, encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# 默认值
# --------------------------------------------------------------------------- #
def test_default_engine_is_excel(tmp_path: Path) -> None:
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    code: |\n      L = {{ L }}\n')
    config = load_config(path)
    assert config.templates[0].engine == "excel"


def test_default_engine_writes_formulas(tmp_path: Path) -> None:
    """默认生成的 Output 表里是**公式**，不是文本快照 —— 工作簿因此独立可用。"""
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    code: |\n      L = {{ L }}\n')
    config = load_config(path)
    book = create_template(config, tmp_path / "t.xlsx", cases=2, overwrite=True, include_scripts=False)
    write_results(book, config, render_all(config, book).results)

    workbook = load_workbook(book)
    try:
        value = workbook["Output"]["B2"].value
        assert isinstance(value, str) and value.startswith("="), value
        assert "INDEX" in value and "MATCH" in value  # 按变量名定位，不是写死的快照
    finally:
        workbook.close()


def test_snapshot_still_available_when_asked_for(tmp_path: Path) -> None:
    path = _write(
        tmp_path,
        '  - name: t\n    output_sheet: "Output"\n    engine: "snapshot"\n    code: |\n      L = {{ L }}\n',
    )
    config = load_config(path)
    assert config.templates[0].engine == "snapshot"
    book = create_template(config, tmp_path / "t.xlsx", cases=2, overwrite=True, include_scripts=False)
    write_results(book, config, render_all(config, book).results)

    workbook = load_workbook(book)
    try:
        assert workbook["Output"]["B2"].value == "L = 340 m"
    finally:
        workbook.close()


def test_unknown_engine_is_rejected(tmp_path: Path) -> None:
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    engine: "nope"\n    code: |\n      x\n')
    with pytest.raises(ConfigError):
        load_config(path)


# --------------------------------------------------------------------------- #
# 快照模式表达得了、公式模式表达不了的写法
# --------------------------------------------------------------------------- #
def test_snapshot_can_use_filters_and_loops(tmp_path: Path) -> None:
    """过滤器 / 循环只有快照模式能做 —— 显式声明之后照常可用。"""
    path = _write(
        tmp_path,
        '  - name: t\n    output_sheet: "Output"\n    engine: "snapshot"\n'
        "    code: |\n"
        "      {% for i in [1, 2] %}row {{ i }}: {{ L.value | pvs('', ' m') }}\n"
        "      {% endfor %}",
    )
    config = load_config(path)
    book = create_template(config, tmp_path / "t.xlsx", cases=1, overwrite=True, include_scripts=False)
    lines = [line for line in render_all(config, book).results["t"][0].lines if line]
    assert lines == ["row 1: 340 m", "row 2: 340 m"]


# --------------------------------------------------------------------------- #
# 建表时就报错（别等填完参数才发现模板用不了）
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("code", "hint"),
    [
        ("{{ L | pvs('', ' m') }}", "过滤器"),
        ("{% for i in [1] %}x{% endfor %}", "for"),
        ("{% if L %}\n多行\n{% endif %}", "if"),
    ],
)
def test_formula_default_rejects_unsupported_template_at_init(tmp_path: Path, code: str, hint: str) -> None:
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    code: |\n' + _indent(code) + "\n")
    config = load_config(path)

    with pytest.raises(FormulaError) as excinfo:
        create_template(config, tmp_path / "t.xlsx", cases=1, overwrite=True, include_scripts=False)
    message = str(excinfo.value)
    assert hint in message, message
    assert "engine: snapshot" in message, message
    # 报错发生在建表阶段：文件不该被生出来
    assert not (tmp_path / "t.xlsx").exists()


def test_unsupported_template_only_fails_for_formula_mode(tmp_path: Path) -> None:
    """同一段模板显式写成快照模式就没问题 —— 报错是"引擎选错了"，不是"模板写错了"。"""
    path = _write(
        tmp_path,
        '  - name: t\n    output_sheet: "Output"\n    engine: "snapshot"\n    code: |\n'
        "      {{ L | pvs('', ' m') }}\n",
    )
    config = load_config(path)
    create_template(config, tmp_path / "t.xlsx", cases=1, overwrite=True, include_scripts=False)
    assert (tmp_path / "t.xlsx").exists()


def _indent(code: str) -> str:
    return "".join(f"      {line}\n" for line in code.splitlines()).rstrip("\n")


# --------------------------------------------------------------------------- #
# 仓库自带的示例也遵守这条默认
# --------------------------------------------------------------------------- #
def test_shipped_formula_example_needs_no_engine_line(tmp_path: Path) -> None:
    """`examples/example_formula.yaml` 里那两处 engine: excel 现在只是"写出来更清楚"。"""
    here = Path(__file__).resolve().parents[1]
    example = here / "examples" / "example_formula.yaml"
    if not example.exists():  # pragma: no cover - sdist 里一定带着
        pytest.skip("示例不在")
    config: ProjectConfig = load_config(example)
    assert all(template.engine == "excel" for template in config.templates)

    # 去掉 engine 行，语义应当完全不变（因为默认就是 excel）
    stripped = tmp_path / "no_engine.yaml"
    stripped.write_text(
        "\n".join(line for line in example.read_text(encoding="utf-8").splitlines() if "engine:" not in line) + "\n",
        encoding="utf-8",
    )
    without = load_config(stripped)
    assert [t.engine for t in without.templates] == [t.engine for t in config.templates]


# --------------------------------------------------------------------------- #
# init 生成完就能用（不用再跑第二条命令）
# --------------------------------------------------------------------------- #
def _cli():
    from typer.testing import CliRunner

    from excel_codegen.cli import app

    return CliRunner(), app


def test_init_prefills_the_output_so_the_workbook_is_ready(tmp_path: Path) -> None:
    """`init` 之后直接打开 Excel 就能用：输出表里已经是活公式，check 也立刻一致。"""
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    code: |\n      L = {{ L }}\n')
    book = tmp_path / "t.xlsx"
    runner, app = _cli()

    result = runner.invoke(app, ["init", "-c", str(path), "-o", str(book), "--cases", "2"])
    assert result.exit_code == 0, result.output
    assert "已写入公式" in result.output
    assert "不用再跑命令" in result.output

    workbook = load_workbook(book)
    try:
        value = workbook["Output"]["B2"].value
        assert isinstance(value, str) and value.startswith("=") and "INDEX" in value
    finally:
        workbook.close()

    # 生成完就一致 —— 不需要先手跑一次 render
    result = runner.invoke(app, ["check", "-c", str(path), "-x", str(book)])
    assert result.exit_code == 0, result.output


def test_init_can_skip_prefill(tmp_path: Path) -> None:
    """--no-prerender：只要骨架（例如想先自己填一遍参数）。"""
    path = _write(tmp_path, '  - name: t\n    output_sheet: "Output"\n    code: |\n      L = {{ L }}\n')
    book = tmp_path / "t.xlsx"
    runner, app = _cli()

    result = runner.invoke(app, ["init", "-c", str(path), "-o", str(book), "--cases", "2", "--no-prerender"])
    assert result.exit_code == 0, result.output
    assert "还没有写输出" in result.output

    workbook = load_workbook(book)
    try:
        assert workbook["Output"]["B2"].value is None
    finally:
        workbook.close()


def test_init_prefill_is_best_effort(tmp_path: Path) -> None:
    """默认值过不了跨变量校验时**不能**让 init 失败：骨架照给，只跳过预填并说明原因。"""
    body = (
        "variables:\n"
        "  global: []\n"
        "  local:\n"
        "    - name: inner\n"
        "      type: float\n"
        "      default: 30\n"
        "    - name: outer\n"
        "      type: float\n"
        "      default: 10\n"
        "asserts:\n"
        '  - "inner.value <= outer.value"\n'  # 默认值 30 > 10，填表时才会被发现
        'templates:\n  - name: t\n    output_sheet: "Output"\n    code: |\n'
        "      inner={{ inner }} outer={{ outer }}\n"
    )
    path = tmp_path / "p.yaml"
    path.write_text(
        'version: 1\nexcel:\n  output: "t.xlsx"\n  template_sheet: null\n  howto_sheet: null\n'
        '  sheets:\n    global: "Global Parameter"\n    local: "Local Parameter"\n    outputs: ["Output"]\n' + body,
        encoding="utf-8",
    )
    book = tmp_path / "t.xlsx"
    runner, app = _cli()

    result = runner.invoke(app, ["init", "-c", str(path), "-o", str(book), "--cases", "1"])
    assert result.exit_code == 0, result.output
    assert "跳过预填输出" in result.output
    assert book.exists()
    # 骨架照样能用：填好参数之后 render 会过
    workbook = load_workbook(book)
    try:
        sheet = workbook["Local Parameter"]
        rows = {sheet.cell(row=r, column=1).value: r for r in range(2, sheet.max_row + 1)}
        sheet.cell(row=rows["outer"], column=5, value=40)
        workbook.save(book)
    finally:
        workbook.close()
    result = runner.invoke(app, ["render", "-c", str(path), "-x", str(book), "--write-excel"])
    assert result.exit_code == 0, result.output
