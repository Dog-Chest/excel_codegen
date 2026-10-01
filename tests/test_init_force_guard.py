"""``init --force`` 的防呆：重建会盖掉已填的参数，所以"人填过东西"时要显式 ``--yes``。

背景（改进建议 §3.6）
--------------------
``--force`` 是**数据丢失点**：它会重建工作簿，已填的参数全丢（文档 Q3 也承认了），
但没有任何防呆 —— 手滑一次就得重填。判据是**参数指纹**：等于"全部取 YAML 默认值"的指纹时，
这本表就是刚生成的、没人动过，覆盖它没有代价（所以老用法一条命令照常跑完）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from spreadsheet_codegen.cli import app
from spreadsheet_codegen.excel_io import default_input_fingerprint, workbook_deviates_from_defaults
from spreadsheet_codegen.models import load_config

runner = CliRunner()

CONFIG = """\
version: 1

excel:
  output: "book.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: baud
      default: 115200
  local:
    - name: port
      default: "A"

templates:
  - name: demo
    output_sheet: "Output"
    start_cell: "B2"
    code: |
      {{ port }} @ {{ baud }}
"""


@pytest.fixture()
def config_path(tmp_path: Path) -> Path:
    path = tmp_path / "demo.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return path


def _set_global(path: Path, value: object) -> None:
    workbook = load_workbook(path)
    try:
        workbook["Global Parameter"]["B2"] = value
        workbook.save(path)
    finally:
        workbook.close()


def _global_value(path: Path) -> object:
    workbook = load_workbook(path)
    try:
        return workbook["Global Parameter"]["B2"].value
    finally:
        workbook.close()


def _init(config_path: Path, excel: Path, *extra: str):
    return runner.invoke(
        app,
        ["init", "-c", str(config_path), "-o", str(excel), "--cases", "2", *extra],
    )


# --------------------------------------------------------------------------- #
# 库接口
# --------------------------------------------------------------------------- #
def test_fresh_workbook_looks_like_defaults(config_path: Path, tmp_path: Path) -> None:
    """刚生成的工作簿 = 全默认值 —— 不该被当成"人填过东西"。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    config = load_config(config_path)
    assert default_input_fingerprint(config, cases=["Case1", "Case2"])
    assert workbook_deviates_from_defaults(excel, config) is None


def test_edited_workbook_is_reported(config_path: Path, tmp_path: Path) -> None:
    """改过参数的工作簿要能报出"人填过东西"，并带上两个指纹便于核对。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    _set_global(excel, 9600)

    note = workbook_deviates_from_defaults(excel, load_config(config_path))
    assert note is not None
    assert "2 个 Case" in note
    assert "指纹" in note


def test_missing_workbook_is_not_a_problem(config_path: Path, tmp_path: Path) -> None:
    """文件不存在（或不是工作簿）时不该在这里报错 —— 交给 init 自己的逻辑。"""
    assert workbook_deviates_from_defaults(tmp_path / "nope.xlsx", load_config(config_path)) is None


# --------------------------------------------------------------------------- #
# CLI 行为
# --------------------------------------------------------------------------- #
def test_force_on_untouched_workbook_still_works(config_path: Path, tmp_path: Path) -> None:
    """没人动过的表：`--force` 照旧一条命令跑完（不要求 --yes）。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    result = _init(config_path, excel, "--force")
    assert result.exit_code == 0, result.output


def test_force_on_edited_workbook_is_blocked(config_path: Path, tmp_path: Path) -> None:
    """人填过参数：拦住，**并且一个字节都不改**（原值还在）。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    _set_global(excel, 9600)

    result = _init(config_path, excel, "--force")
    assert result.exit_code == 2  # 配置/取值类错误
    output = (result.output or "") + (getattr(result, "stderr", "") or "")
    assert "人填过东西" in output
    assert "--yes" in output
    assert _global_value(excel) == 9600, "被拦住时不该覆盖"


def test_yes_lets_the_rebuild_through(config_path: Path, tmp_path: Path) -> None:
    """显式 --yes 才真的重建（参数回到 YAML 默认值）。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    _set_global(excel, 9600)

    result = _init(config_path, excel, "--force", "--yes")
    assert result.exit_code == 0, result.output
    assert _global_value(excel) == 115200


def test_yes_alone_is_harmless(config_path: Path, tmp_path: Path) -> None:
    """只加 --yes（没有 --force）不影响任何事 —— 目标存在时照样按"已存在"报错。"""
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    result = _init(config_path, excel, "--yes")
    assert result.exit_code == 1
    assert "已存在" in ((result.output or "") + (getattr(result, "stderr", "") or ""))


def test_render_record_does_not_make_it_trustworthy(config_path: Path, tmp_path: Path) -> None:
    """人改过参数**并且重渲染过**时，照样要拦。

    渲染记录里存的是"上次渲染时的参数"——它等于当前参数只能说明"渲染过"，
    **完全不能说明参数还是 YAML 默认值**。早先拿它当短路条件，于是
    "改了参数 → 渲染 → `init --force`"这条路会绕过防呆，把用户填的值静默换回默认值
    （`--force` 是数据丢失点，这正是防呆要挡的那一下）。
    """
    excel = tmp_path / "book.xlsx"
    assert _init(config_path, excel).exit_code == 0
    _set_global(excel, 9600)
    assert runner.invoke(app, ["render", "-c", str(config_path), "-x", str(excel), "-w", "--no-show"]).exit_code == 0

    config = load_config(config_path)
    assert workbook_deviates_from_defaults(excel, config) is not None  # 仍然偏离默认值
    result = _init(config_path, excel, "--force")
    assert result.exit_code == 2
    assert _global_value(excel) == 9600, "被拦住时不该覆盖"


# --------------------------------------------------------------------------- #
# 有「渲染记录」的工作簿（Template / HOWTO 表打开）—— 上面那条路径的真正来源
# --------------------------------------------------------------------------- #
#: 与 CONFIG 的差别只在**开着渲染记录**（Template / HOWTO 表）。默认配置就是这样，
#: 所以这条路径才是用户真正会走的；上面那个 fixture 把两张表关了，短路分支从没被走到。
RECORDED_CONFIG = """\
version: 1

excel:
  output: "book.xlsx"
  template_sheet: "Template"
  howto_sheet: "HOWTO"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: baud
      default: 115200
  local:
    - name: port
      default: "A"

templates:
  - name: demo
    output_sheet: "Output"
    start_cell: "B2"
    code: |
      {{ port }} @ {{ baud }}
"""


@pytest.fixture()
def recorded_config_path(tmp_path: Path) -> Path:
    path = tmp_path / "recorded.yaml"
    path.write_text(RECORDED_CONFIG, encoding="utf-8")
    return path


def test_recorded_fresh_workbook_is_not_blocked(recorded_config_path: Path, tmp_path: Path) -> None:
    """刚生成（没人动过）→ `--force` 照旧一条命令跑完，哪怕工作簿里有渲染记录。"""
    excel = tmp_path / "book.xlsx"
    assert _init(recorded_config_path, excel).exit_code == 0
    assert workbook_deviates_from_defaults(excel, load_config(recorded_config_path)) is None
    result = _init(recorded_config_path, excel, "--force")
    assert result.exit_code == 0, result.output


def test_recorded_edited_workbook_is_blocked_after_render(recorded_config_path: Path, tmp_path: Path) -> None:
    """**本文件最要紧的一条**：有渲染记录时，"改了参数 → 渲染 → `init --force`"必须仍然被拦。

    回归的是一个静默数据丢失：改 Global B2=9600 → render → `init --force`（不带 `--yes`）
    放行并把 B2 悄悄变回 115200。防呆失效的原因见上面那条测试的 docstring。
    """
    excel = tmp_path / "book.xlsx"
    assert _init(recorded_config_path, excel).exit_code == 0
    _set_global(excel, 9600)
    assert (
        runner.invoke(app, ["render", "-c", str(recorded_config_path), "-x", str(excel), "-w", "--no-show"]).exit_code
        == 0
    )
    assert _global_value(excel) == 9600  # 渲染不会改参数，只是把指纹记下来

    result = _init(recorded_config_path, excel, "--force")
    assert result.exit_code == 2, "渲染记录不能当成'参数还是默认值'的证据"
    output = (result.output or "") + (getattr(result, "stderr", "") or "")
    assert "人填过东西" in output
    assert _global_value(excel) == 9600, "被拦住时不该覆盖"


# --------------------------------------------------------------------------- #
# 基线要按"init 真的会写成什么样"算，否则刚生成的工作簿会被误判成"人填过东西"
# --------------------------------------------------------------------------- #
FANCY_CONFIG = """\
version: 1

excel:
  output: "book.xlsx"
  template_sheet: "Template"
  howto_sheet: "HOWTO"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: baud
      default: 115200
    - name: note
      default: "预填进去的提示值"
    - name: optional
      default: "不该被写进新表"
      prefill: false
    - name: doubled
      derived: "baud * 2"
  local:
    - name: port
      default: "A"

templates:
  - name: demo
    output_sheet: "Output"
    start_cell: "B2"
    code: |
      {{ port }} @ {{ baud }}
"""


@pytest.fixture()
def fancy_config_path(tmp_path: Path) -> Path:
    path = tmp_path / "fancy.yaml"
    path.write_text(FANCY_CONFIG, encoding="utf-8")
    return path


def test_fresh_workbook_with_prefill_false_is_not_blocked(fancy_config_path: Path, tmp_path: Path) -> None:
    """``prefill: false`` 的格子 init 根本不写（读回来是空串）—— 基线必须照这个口径，
    否则"刚生成的工作簿"会被误判成"人填过东西"，`--force` 平白要求 `--yes`。"""
    excel = tmp_path / "book.xlsx"
    assert _init(fancy_config_path, excel).exit_code == 0
    assert workbook_deviates_from_defaults(excel, load_config(fancy_config_path)) is None


def test_derived_params_do_not_look_like_edits(fancy_config_path: Path, tmp_path: Path) -> None:
    """派生参数是**算出来的**，不是"人填的东西" —— 不能因为它让基线对不上。

    基线是"直接取 default + prefix/suffix"，而读工作簿时派生值是现算的；两套口径只要
    差一点，带派生参数的项目就永远被判成"人填过东西"。
    """
    excel = tmp_path / "book.xlsx"
    assert _init(fancy_config_path, excel).exit_code == 0
    assert workbook_deviates_from_defaults(excel, load_config(fancy_config_path)) is None
    result = _init(fancy_config_path, excel, "--force")
    assert result.exit_code == 0, result.output


def test_editing_a_prefill_false_cell_counts_as_an_edit(fancy_config_path: Path, tmp_path: Path) -> None:
    """反过来的那一半：``prefill: false`` 的空格子被填上东西时，照样要算"人填过东西"。"""
    excel = tmp_path / "book.xlsx"
    assert _init(fancy_config_path, excel).exit_code == 0
    # Global 表：B2=baud、B3=note、B4=optional、B5=doubled
    workbook = load_workbook(excel)
    try:
        workbook["Global Parameter"]["B4"] = "手工填的"
        workbook.save(excel)
    finally:
        workbook.close()
    assert workbook_deviates_from_defaults(excel, load_config(fancy_config_path)) is not None
