"""``spreadsheet-codegen doctor``：一条命令体检环境 / 配置 / 工作簿。

它的价值在"把散在各处的常见坑一次说清"，所以测试关注三件事：
① 正常项目退出码 0 且该报的"值得留意"都报了；② 真有问题时给出**分类退出码**（配置错 2 /
工作簿打不开 4 / 其他 1）且**不甩 traceback**；
③ 只该是警告的事（没 uv、没渲染记录、没约束）不能升级成错误。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from spreadsheet_codegen import create_template, load_config, render_all
from spreadsheet_codegen.cli import app
from spreadsheet_codegen.excel_io import write_results

runner = CliRunner()

CONFIG = """\
version: 1

excel:
  output: "t.xlsx"
  template_sheet: "Template"
  howto_sheet: "HOWTO"
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: d_tank
      type: float
      default: 24
      min: 1
  local:
    - name: draft
      type: float
      default: 20
      min: 0
    - name: unused_one
      default: "x"

asserts:
  - "draft.value <= d_tank.value"

templates:
  - name: demo
    output_sheet: "Output"
    code: |
      // {{ draft }}
"""


def text_of(result) -> str:
    return ((result.output or "") + (getattr(result, "stderr", "") or "")).replace("\n", "")


@pytest.fixture()
def config_path(tmp_path: Path) -> Path:
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    return path


def run_doctor(path: Path, excel: Path | None = None):
    args = ["doctor", "-c", str(path)]
    if excel is not None:
        args += ["-x", str(excel)]
    return runner.invoke(app, args)


def make_book(config_path: Path, *, render: bool = False) -> Path:
    config = load_config(config_path)
    excel = config_path.parent / "t.xlsx"
    create_template(config, excel, cases=["A", "B"], overwrite=True, include_scripts=False)
    if render:
        write_results(excel, config, render_all(config, excel).results)
    return excel


# --------------------------------------------------------------------------- #
# 正常路径
# --------------------------------------------------------------------------- #
def test_doctor_on_healthy_project(config_path: Path) -> None:
    excel = make_book(config_path, render=True)
    result = run_doctor(config_path, excel)
    assert result.exit_code == 0, text_of(result)
    text = text_of(result)
    assert "Python" in text and "依赖" in text
    assert "加载成功" in text
    assert "与当前一致" in text
    assert "体检完成" in text


def test_doctor_excel_defaults_to_config_output(config_path: Path, monkeypatch) -> None:
    """不传 -x 时用配置里的 excel.output；它**相对配置文件所在目录**解析（与 init / render 一致）。"""
    make_book(config_path, render=True)
    monkeypatch.chdir(config_path.parent)
    result = run_doctor(config_path)  # 不传 -x
    assert result.exit_code == 0
    assert "与当前一致" in text_of(result)


def test_doctor_finds_config_output_from_another_cwd(config_path: Path, monkeypatch) -> None:
    """从别的目录跑 doctor：``excel.output`` 相对**配置文件**解析，照样找得到工作簿。

    这条是 1.4 的回归：以前它相对 CWD 解析，换个目录跑就报"工作簿不存在"。
    """
    make_book(config_path, render=True)
    monkeypatch.chdir(config_path.parent.parent)
    result = run_doctor(config_path)  # 不传 -x，CWD 也不是配置目录
    assert result.exit_code == 0, text_of(result)
    assert "与当前一致" in text_of(result)
    assert "不存在" not in text_of(result)


def test_missing_workbook_is_only_a_warning(config_path: Path) -> None:
    result = run_doctor(config_path)  # 还没 init
    assert result.exit_code == 0
    assert "不存在" in text_of(result)


def test_missing_render_record_is_only_a_warning(config_path: Path) -> None:
    excel = make_book(config_path, render=False)
    result = run_doctor(config_path, excel)
    assert result.exit_code == 0
    assert "还没跑过" in text_of(result)


def test_missing_constraints_is_reported_as_a_hint(tmp_path: Path) -> None:
    path = tmp_path / "bare.yaml"
    path.write_text(
        CONFIG.replace("      min: 1\n", "")
        .replace("      min: 0\n", "")
        .replace('asserts:\n  - "draft.value <= d_tank.value"\n', "asserts: []\n"),
        encoding="utf-8",
    )
    excel = make_book(path, render=True)
    result = run_doctor(path, excel)
    assert result.exit_code == 0
    assert "一条约束都没有" in text_of(result)


def test_unused_variable_is_reported(tmp_path: Path) -> None:
    excel = make_book(config_path=tmp_path / "t.yaml") if False else None  # 占位，见下
    path = tmp_path / "u.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    excel = make_book(path, render=True)
    result = run_doctor(path, excel)
    assert result.exit_code == 0
    text = text_of(result)
    assert "unused_one" in text
    # 保留名（case_name / template_name）不该被当成"定义了没人用"
    assert "template_name" not in text


# --------------------------------------------------------------------------- #
# 真问题：给出分类退出码，但必须是可读结论而不是 traceback
# --------------------------------------------------------------------------- #
def test_value_violation_is_an_error_not_a_crash(tmp_path: Path) -> None:
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    excel = make_book(path, render=True)

    book = load_workbook(excel)
    try:
        book["Local Parameter"]["E2"] = 99  # draft 超过 d_tank
        book.save(excel)
    finally:
        book.close()

    result = run_doctor(path, excel)
    assert result.exit_code == 2, "取值不满足 asserts 属于『你给的东西不对』（2），不是通用失败"
    text = text_of(result)
    assert "asserts" in text
    assert "Traceback" not in text
    assert "不满足" in text


def test_broken_yaml_is_reported_cleanly(tmp_path: Path) -> None:
    path = tmp_path / "bad.yaml"
    path.write_text("version: 1\nvariables: {global: [], local: []}\n", encoding="utf-8")
    result = run_doctor(path)
    assert result.exit_code == 2, "YAML 加载失败 = 配置错（2）"
    text = text_of(result)
    assert "YAML" in text and "ERROR" in text
    assert "Traceback" not in text


def test_unreadable_workbook_is_reported_not_a_traceback(tmp_path: Path) -> None:
    """工作簿读不了（被占用 / 写坏了）时必须是报告里的一行结论 + 退出码 4。

    这条曾经真的坏过：``load_workbook_file`` 在 ``try`` 之外，openpyxl 的异常
    一路冒到 Typer，用户看到的是整页 traceback —— 而"工作簿坏了"恰恰是用户来跑
    ``doctor`` 最常见的原因，等于在最需要它的时候失效。
    """
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    (tmp_path / "t.xlsx").write_bytes(b"this is not an xlsx")

    result = run_doctor(path)
    assert result.exit_code == 4, "打不开工作簿属于环境问题（4）"
    text = text_of(result)
    assert "Traceback" not in text
    assert "无法读取" in text


def test_workbook_missing_sheets_is_reported_not_a_traceback(tmp_path: Path) -> None:
    """结构不对（少了参数表）同样只是报告里的一行 ERROR，退出码 1。"""
    path = tmp_path / "t.yaml"
    path.write_text(CONFIG, encoding="utf-8")
    make_book(path)  # 先按配置生成一本，再换成"表名不对"的那本

    wrong = Workbook()
    wrong.active.title = "Nope"
    wrong.save(tmp_path / "t.xlsx")

    result = run_doctor(path)
    assert result.exit_code == 1, "工作簿结构与配置对不上 = 通用失败（ExcelError -> 1）"
    text = text_of(result)
    assert "Traceback" not in text
    assert "缺少工作表" in text


def test_missing_fragment_is_reported_as_template_error(tmp_path: Path) -> None:
    path = tmp_path / "f.yaml"
    path.write_text(CONFIG.replace("      // {{ draft }}\n", '      {% include "nope.j2" %}\n'), encoding="utf-8")
    result = run_doctor(path)
    assert result.exit_code == 1
    text = text_of(result)
    assert "找不到片段" in text
    assert "Traceback" not in text
