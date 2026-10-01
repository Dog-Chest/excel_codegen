"""第三层作用域：成员表（船 → 工况 → 舱/设备）。

没有它的时候，一个被多个工况引用的舱只能把参数**按工况摊平**（同名舱在每个 Case 列里
各写一遍，改一个舱的尺寸要改 N 列）。有了它，舱的参数只写一遍，Case 用一个"指针变量"
（``group.key``）指向自己用哪个成员。

布局与 Global / Local 都不同：**一行一个成员，B 列起一个变量一列**（表头是变量名）——
这是工程师写舱容表的习惯，也让公式模式能用同一形态的 ``INDEX/MATCH`` 定位。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from spreadsheet_codegen import create_template, load_config, read_group_members, render_all
from spreadsheet_codegen.cli import app
from spreadsheet_codegen.excel_io import load_workbook_file, write_results
from spreadsheet_codegen.utils import ConfigError, ExcelError, InputError

runner = CliRunner()

BASE = """\
version: 1

excel:
  output: "g.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: rho
      type: float
      default: 1025
  local:
    - name: tank_ref
      choices: ["WBT6", "WBT7"]
      default: "WBT6"
  group:
    sheet: "Tank Data"
    key: tank_ref
    members: ["WBT6", "WBT7"]
    variables:
      - name: l_tank
        type: float
        default: 42
        min: 0
        unit: "m"
      - name: h_tank
        type: float
        default: 32

templates:
  - name: demo
    output_sheet: "Output"
    __ENGINE__code: |
      // {{ case_name }} tank={{ tank_ref }} L={{ l_tank }} H={{ h_tank }} rho={{ rho }}
"""


def full_config(engine: str = "") -> str:
    """把 engine 占位替换掉（不传就是快照模式）。"""
    return BASE.replace("__ENGINE__", f"engine: {engine}\n    " if engine else "")


def write_config(tmp_path: Path, text: str, name: str = "g.yaml"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return load_config(path)


def make_config(tmp_path: Path, *, engine: str = ""):
    return write_config(tmp_path, full_config(engine))


def make_book(tmp_path: Path, *, engine: str = "") -> Path:
    config = make_config(tmp_path, engine=engine)
    return create_template(config, tmp_path / "g.xlsx", cases=["A", "B"], overwrite=True, include_scripts=False)


def set_member(book: Path, member: str, **values) -> None:
    """按成员名改成员表里的值（列名取自表头）。"""
    wb = load_workbook(book)
    try:
        ws = wb["Tank Data"]
        headers = {ws.cell(row=1, column=c).value: c for c in range(2, ws.max_column + 1)}
        target = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == member)
        for name, value in values.items():
            ws.cell(row=target, column=headers[name], value=value)
        wb.save(book)
    finally:
        wb.close()


def set_case(book: Path, column: str, row: int, value) -> None:
    wb = load_workbook(book)
    try:
        wb["Local Parameter"][f"{column}{row}"] = value
        wb.save(book)
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# 配置期
# --------------------------------------------------------------------------- #
def test_group_is_optional(tmp_path: Path) -> None:
    head, rest = full_config().split("  group:", 1)
    path = tmp_path / "plain.yaml"
    path.write_text(head + "templates:" + rest.split("templates:", 1)[1], encoding="utf-8")
    config = load_config(path)
    assert config.group is None
    assert config.group_variables == []


def test_key_must_be_a_local_variable(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as excinfo:
        write_config(tmp_path, full_config().replace("key: tank_ref", "key: rho"))
    assert "不是 local 变量" in str(excinfo.value)


def test_group_variable_cannot_be_derived(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as excinfo:
        write_config(
            tmp_path,
            full_config().replace(
                "      - name: h_tank\n        type: float\n        default: 32\n",
                "      - name: h_tank\n        derived: 'l_tank * 2'\n",
            ),
        )
    assert "derived" in str(excinfo.value)


def test_group_name_cannot_clash_with_global_or_local(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as excinfo:
        write_config(tmp_path, full_config().replace("      - name: h_tank\n", "      - name: rho\n"))
    assert "重名" in str(excinfo.value)


def test_group_sheet_cannot_clash(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as excinfo:
        write_config(tmp_path, full_config().replace('sheet: "Tank Data"', 'sheet: "Local Parameter"'))
    assert "与 Global / Local / Output 表名冲突" in str(excinfo.value)


def test_empty_group_variables_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as excinfo:
        write_config(
            tmp_path,
            full_config().split("    variables:\n")[0] + "    variables: []\n\ntemplates:\n  - name: demo\n"
            '    output_sheet: "Output"\n    code: |\n      x\n',
        )
    assert "不能为空" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# Excel 结构
# --------------------------------------------------------------------------- #
def test_member_sheet_layout(tmp_path: Path) -> None:
    """一行一个成员，B 列起一个变量一列（表头是变量名）。"""
    book = make_book(tmp_path)
    wb = load_workbook(book)
    try:
        ws = wb["Tank Data"]
        assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["成员", "l_tank", "h_tank"]
        assert [ws.cell(row=r, column=1).value for r in (2, 3)] == ["WBT6", "WBT7"]
        assert ws.cell(row=2, column=2).value == 42
        # 表头带批注（描述 / 单位 / 约束 / 填哪儿）
        comment = ws.cell(row=1, column=2).comment
        assert comment is not None
        assert "单位：m" in comment.text
        assert "范围: >= 0" in comment.text
        # 取值约束写成了数据有效性
        assert ws.data_validations.dataValidation
    finally:
        wb.close()


def test_members_can_be_added_by_inserting_rows(tmp_path: Path) -> None:
    """成员表是"从表里发现"的 —— 插一行就多一个成员，不用改 YAML。"""
    book = make_book(tmp_path)
    wb = load_workbook(book)
    try:
        ws = wb["Tank Data"]
        ws.cell(row=4, column=1, value="COT1")
        ws.cell(row=4, column=2, value=99)
        ws.cell(row=4, column=3, value=11)
        wb.save(book)
    finally:
        wb.close()

    config = make_config(tmp_path)
    wb = load_workbook_file(book)
    try:
        members = read_group_members(wb, config)
    finally:
        wb.close()
    assert set(members) == {"WBT6", "WBT7", "COT1"}
    assert members["COT1"]["l_tank"].text == "99"


# --------------------------------------------------------------------------- #
# 渲染
# --------------------------------------------------------------------------- #
def test_each_case_resolves_its_own_member(tmp_path: Path) -> None:
    book = make_book(tmp_path)
    set_member(book, "WBT7", l_tank=14.5, h_tank=10.15)
    set_case(book, "F", 2, "WBT7")  # Case B 用 WBT7

    config = make_config(tmp_path)
    output = render_all(config, book)
    assert output.results["demo"][0].lines == ["// A tank=WBT6 L=42 H=32 rho=1025"]
    assert output.results["demo"][1].lines == ["// B tank=WBT7 L=14.5 H=10.15 rho=1025"]


def test_empty_key_is_reported(tmp_path: Path) -> None:
    """key 没有任何回落时（没有 choices、也没有 default）要明确报"是空的"。

    注意：如果 key 声明了 ``choices``，空值会先被取值约束拦住 —— 那条报错更具体，
    所以这里特意把 choices 也去掉，专门测 member_of 的那条分支。
    """
    book = make_book(tmp_path)
    set_case(book, "E", 2, "")
    path = tmp_path / "no_default.yaml"
    path.write_text(
        full_config().replace('      choices: ["WBT6", "WBT7"]\n      default: "WBT6"\n', ""),
        encoding="utf-8",
    )
    config2 = load_config(path)
    with pytest.raises(ExcelError) as excinfo:
        render_all(config2, book)
    assert "是空的" in str(excinfo.value)


def test_unknown_member_is_reported(tmp_path: Path) -> None:
    book = make_book(tmp_path)
    set_case(book, "E", 2, "NOPE")
    config = make_config(tmp_path)
    with pytest.raises(InputError) as excinfo:
        render_all(config, book)
    message = str(excinfo.value)
    assert "NOPE" in message
    assert "WBT6" in message  # 可选成员列表


def test_group_constraints_are_checked(tmp_path: Path) -> None:
    book = make_book(tmp_path)
    set_member(book, "WBT7", l_tank=-5)  # min: 0
    config = make_config(tmp_path)
    with pytest.raises(InputError) as excinfo:
        render_all(config, book)
    assert "Tank Data" in str(excinfo.value)
    assert "WBT7" in str(excinfo.value)


def test_global_and_local_can_use_group_values_in_expressions(tmp_path: Path) -> None:
    """成员值进了上下文，模板里照常引用（含 case_filter / asserts）。"""
    path = tmp_path / "x.yaml"
    path.write_text(
        full_config().replace("templates:", 'asserts:\n  - "l_tank.value > 0"\ntemplates:'),
        encoding="utf-8",
    )
    config = load_config(path)
    book = create_template(config, tmp_path / "g.xlsx", cases=["A"], overwrite=True, include_scripts=False)
    set_member(book, "WBT6", l_tank=0)
    with pytest.raises(InputError) as excinfo:
        render_all(config, book)
    assert "asserts" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# 公式模式
# --------------------------------------------------------------------------- #
def test_formula_mode_uses_nested_lookup(tmp_path: Path) -> None:
    book = make_book(tmp_path, engine="excel")
    set_member(book, "WBT7", l_tank=14.5)
    set_case(book, "F", 2, "WBT7")

    config = make_config(tmp_path, engine="excel")
    write_results(book, config, render_all(config, book).results)

    wb = load_workbook(book)
    try:
        formula = wb["Output"].cell(row=2, column=2).value
    finally:
        wb.close()
    assert "Tank Data" in formula
    assert "MATCH(" in formula  # 成员 → 行 的嵌套查找

    result = runner.invoke(app, ["check", "-c", str(tmp_path / "g.yaml"), "-x", str(book)])
    assert result.exit_code == 0, (result.output or "") + str(getattr(result, "stderr", ""))


def test_formula_check_notices_member_change(tmp_path: Path) -> None:
    """改成员表里的值 → 公式模式改参数免重跑，所以输出仍然一致（值由 Excel 现算）。"""
    book = make_book(tmp_path, engine="excel")
    config = make_config(tmp_path, engine="excel")
    write_results(book, config, render_all(config, book).results)

    set_member(book, "WBT6", l_tank=100)
    result = runner.invoke(app, ["check", "-c", str(tmp_path / "g.yaml"), "-x", str(book)])
    assert result.exit_code == 0
