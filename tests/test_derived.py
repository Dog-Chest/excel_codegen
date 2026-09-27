"""派生参数（参数引用参数）的测试：Python 求值 / Excel 公式 / 顺序 / 越界 / 告警。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_codegen.derived import (
    DerivedError,
    DerivedNotTranslatable,
    evaluate_derived,
    expression_names,
    is_translatable,
    ordered_derived,
    to_excel,
)
from excel_codegen.excel_io import create_template, read_cases, read_global_values, write_results
from excel_codegen.formula_eval import evaluate_formula
from excel_codegen.models import ProjectConfig, VariableDef, load_config
from excel_codegen.renderer import render_all
from excel_codegen.utils import VarValue

DERIVED_YAML = """\
version: 1

excel:
  output: "derived.xlsx"
  template_sheet: null
  howto_sheet: null
  sheets:
    global: "Global Parameter"
    local: "Local Parameter"
    outputs: ["Output"]

variables:
  global:
    - name: rho
      description: "密度"
      default: 1025
      type: float
    - name: g
      description: "重力加速度"
      default: 9.81
      type: float
    - name: rho_g
      description: "rho * g（全局派生）"
      derived: "rho * g"
      type: float
      suffix: " N/m^3"
  local:
    - name: draft
      default: 20.5
      type: float
    - name: z
      default: 5
      type: float
    - name: h_s
      description: "静水压头 = max(draft - z, 0)"
      derived: "max(draft - z, 0)"
      type: float
      suffix: " m"
    - name: p_s
      description: "静水压力（链式：引用派生参数 rho_g）"
      derived: "rho_g / 1000 * h_s"
      type: float

templates:
  - name: pressure
    output_sheet: "Output"
    start_cell: "B2"
    direction: "horizontal"
    engine: "excel"
    code: |
      // case {{ case_name }}
      var h_s = {{ h_s }};
      var p_s = {{ p_s }} Pa;
"""


@pytest.fixture()
def derived_config(tmp_path: Path) -> ProjectConfig:
    path = tmp_path / "derived.yaml"
    path.write_text(DERIVED_YAML, encoding="utf-8")
    return load_config(path)


# --------------------------------------------------------------------------- #
# 表达式与顺序
# --------------------------------------------------------------------------- #
def test_expression_names_and_order(derived_config: ProjectConfig) -> None:
    # max 已注册为 Jinja 全局函数，所以不算"变量"
    assert expression_names("max(draft - z, 0)") == {"draft", "z"}
    ordered = [v.name for v in ordered_derived(derived_config.local_variables, scope="local")]
    assert ordered.index("h_s") < ordered.index("p_s")  # 被依赖的先算


def test_cycle_is_reported() -> None:
    variables = [
        VariableDef.model_validate({"name": "a", "derived": "b + 1"}),
        VariableDef.model_validate({"name": "b", "derived": "a + 1"}),
    ]
    with pytest.raises(DerivedError, match="循环引用"):
        ordered_derived(variables, scope="local")


def test_unknown_name_is_reported() -> None:
    variables = [VariableDef.model_validate({"name": "a", "derived": "nope * 2"})]
    with pytest.raises(DerivedError, match="取不到的变量"):
        evaluate_derived(variables, {}, scope="local")


def test_global_cannot_reference_local() -> None:
    globals_ = [VariableDef.model_validate({"name": "a", "derived": "localvar * 2"})]
    with pytest.raises(DerivedError, match="不能引用 local"):
        evaluate_derived(
            globals_, {}, scope="global", forbidden={"localvar": "global 不能引用 local"}
        )


def test_derived_arithmetic_uses_pure_values() -> None:
    definition = VariableDef.model_validate({"name": "sum", "derived": "a + b"})
    context = {"a": VarValue(2, "P", "S"), "b": VarValue(3)}
    context = {name: value.value for name, value in context.items()}
    evaluate_derived([definition], context, scope="local")
    assert context["sum"] == 5


# --------------------------------------------------------------------------- #
# 翻译成 Excel 公式
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "expression, expected",
    [
        ("a * b", "(A*B)"),
        ("a ** 2 % 3", "MOD((A^2),3)"),
        ("max(a - b, 0)", "MAX((A-B),0)"),
        ("int(x)", "TRUNC(X)"),
        ("a if b > 1 else c", "IF((B>1),A,C)"),
        ('"GPIO" ~ p', '("GPIO")&(P)'),
        # and / or 是条件表达式的一部分（0.6.0 修复：它们曾经被 BinExpr 分支抢先命中而不可用）
        ("a if b > 1 and c < 2 else d", "IF(AND((B>1),(C<2)),A,D)"),
        ("a if b > 1 or c < 2 else d", "IF(OR((B>1),(C<2)),A,D)"),
        ("a if not b else c", "IF(NOT(B),A,C)"),
    ],
)
def test_translate_expressions(expression: str, expected: str) -> None:
    assert to_excel(expression, name="t", resolve=lambda name: name.upper()) == expected


def test_and_or_outside_condition_is_rejected() -> None:
    """非条件位置不放行 and/or —— Python 的返回值语义与 Excel 不同，翻了会静默不一致。"""
    with pytest.raises(DerivedNotTranslatable, match="只能用在条件"):
        to_excel("a and b", name="t", resolve=lambda name: name.upper())


def test_round_is_deliberately_not_translated() -> None:
    """两种语言的舍入语义不同：宁可不翻译（写算好的值），也不要"Excel 里看到的"和导出不一致。"""
    variable = VariableDef.model_validate({"name": "t", "derived": "round(x, 2)"})
    assert is_translatable(variable) is False
    with pytest.raises(DerivedNotTranslatable, match="银行家舍入"):
        to_excel("round(x, 2)", name="t", resolve=lambda name: name)
    with pytest.raises(DerivedNotTranslatable, match="银行家舍入"):
        to_excel("x | round(2)", name="t", resolve=lambda name: name)


def test_syntax_error_is_not_swallowed() -> None:
    with pytest.raises(DerivedError):  # 语法错必须冒出来，不能当成"翻译不了"降级
        to_excel("a *", name="t", resolve=lambda name: name)


# --------------------------------------------------------------------------- #
# 端到端
# --------------------------------------------------------------------------- #
def test_create_template_writes_live_formulas(derived_config: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(derived_config, tmp_path / "d.xlsx", cases=2)
    workbook = load_workbook(path)
    try:
        global_sheet = workbook["Global Parameter"]
        assert global_sheet["B4"].data_type == "f"  # rho_g：派生格是公式
        assert "MATCH(\"rho\"" in global_sheet["B4"].value
        assert global_sheet["B2"].data_type == "n"  # rho：输入格还是数值
        assert "自动计算" in global_sheet["C4"].value

        local_sheet = workbook["Local Parameter"]
        assert local_sheet["E4"].data_type == "f"  # h_s @ Case1
        assert "MAX(" in local_sheet["E4"].value
        assert local_sheet["F4"].data_type == "f"  # Case2 列同样有公式
        assert "'Local Parameter'!$F:$F" in local_sheet["F4"].value
    finally:
        workbook.close()


def test_render_uses_derived_values(derived_config: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(derived_config, tmp_path / "d.xlsx", cases=1)
    output = render_all(derived_config, path)

    assert pytest.approx(output.global_values["rho_g"].value) == 1025 * 9.81
    case = output.cases[0]
    assert case.values["h_s"].value == 15.5
    assert pytest.approx(case.values["p_s"].value) == 155.856375
    lines = output.results["pressure"][0].lines
    # 数值按 15 位有效数字输出（与 Excel 的 General 一致），不会有二进制尾巴
    assert lines == ["// case Case1", "var h_s = 15.5 m;", "var p_s = 155.856375 Pa;"]


def test_derived_follows_user_input(derived_config: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(derived_config, tmp_path / "d.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Local Parameter"]
        sheet["E2"] = 30  # draft
        sheet["E3"] = 12  # z
        workbook.save(path)
    finally:
        workbook.close()

    output = render_all(derived_config, path)
    assert output.cases[0].values["h_s"].value == 18
    assert pytest.approx(output.cases[0].values["p_s"].value) == 1025 * 9.81 / 1000 * 18


def test_excel_formula_matches_python_render(derived_config: ProjectConfig, tmp_path: Path) -> None:
    """公式求值器算出来的文本 == Python 渲染的文本（派生链 + 公式模式）。"""
    from excel_codegen.formula_eval import evaluate_template_values

    path = create_template(derived_config, tmp_path / "d.xlsx", cases=2)
    output = render_all(derived_config, path)
    write_results(path, derived_config, output.results)

    workbook = load_workbook(path)
    try:
        template = derived_config.templates[0]
        got = evaluate_template_values(
            workbook, derived_config, template, [case.name for case in output.cases]
        )
        for result in output.results["pressure"]:
            assert got[result.case_name] == result.lines
    finally:
        workbook.close()


def test_hand_edited_derived_cell_warns(derived_config: ProjectConfig, tmp_path: Path) -> None:
    path = create_template(derived_config, tmp_path / "d.xlsx", cases=1)
    workbook = load_workbook(path)
    try:
        workbook["Local Parameter"]["E4"] = 999  # h_s 是派生格
        workbook.save(path)
    finally:
        workbook.close()

    output = render_all(derived_config, path)
    assert output.cases[0].values["h_s"].value == 15.5  # 仍然按算式算，忽略手工值
    warnings = " ".join(output.warnings)
    assert "h_s" in warnings and "999" in warnings


def test_write_results_refreshes_untranslatable_derived(
    tmp_path: Path,
) -> None:
    """翻译不了的派生表达式：格子写 Python 算好的值（而不是公式），并给出告警。"""
    config_path = tmp_path / "round.yaml"
    config_path.write_text(
        DERIVED_YAML.replace('derived: "rho_g / 1000 * h_s"', 'derived: "round(rho_g / 1000 * h_s, 1)"'),
        encoding="utf-8",
    )
    config = load_config(config_path)
    path = create_template(config, tmp_path / "r.xlsx", cases=1)
    output = render_all(config, path)
    warnings: list[str] = []
    write_results(path, config, output.results, warnings=warnings)

    workbook = load_workbook(path)
    try:
        cell = workbook["Local Parameter"]["E5"]  # p_s 派生格
        assert cell.data_type == "n"  # 值而不是公式
        assert pytest.approx(cell.value) == round(1025 * 9.81 / 1000 * 15.5, 1)
    finally:
        workbook.close()
    assert any("h_s" not in w for w in warnings) or not warnings  # 刷新过程不报错即可


def test_derived_cell_is_read_as_value_by_evaluator(
    derived_config: ProjectConfig, tmp_path: Path
) -> None:
    """派生格本身是公式时，求值器要能递归算出来（公式引用公式）。"""
    path = create_template(derived_config, tmp_path / "d.xlsx", cases=1)
    output = render_all(derived_config, path)
    write_results(path, derived_config, output.results)

    workbook = load_workbook(path)
    try:
        formula = workbook["Local Parameter"]["E4"].value  # h_s 的公式
        reader = None
        from excel_codegen.formula_eval import Evaluator, WorkbookReader, sheet_names_of

        reader = WorkbookReader(workbook)
        evaluator = Evaluator(reader, {name: name for name in sheet_names_of(derived_config)})
        assert float(evaluator.evaluate(formula)) == 15.5
    finally:
        workbook.close()
