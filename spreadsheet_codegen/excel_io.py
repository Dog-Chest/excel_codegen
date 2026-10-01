"""Excel 侧实现：生成模板、读取用户填写内容、把渲染结果写回工作表。

工作表约定
----------
Global Parameter
    ``A=Variable``、``B=Value``（用户填写）、``C=Description``、``D=Prefix``、``E=Suffix``

Local Parameter
    ``A=Variable``、``B=Description``、``C=Prefix``、``D=Suffix``、从 ``E`` 列开始每个 Case 一列
    （``E1=Case1``、``F1=Case2`` …），用户可右拉增加 Case。

Output / …
    配置中 ``excel.sheets.outputs`` 声明的输出表，渲染结果写在这里。

Template（可选，隐藏）
    模板原文 + 机器可读的渲染元信息（时间 / 参数指纹 / 输出指纹），方便对照与 `check`。

HOWTO（可选，第一张）
    写进工作簿本身的"下一步跑什么"说明；每次写回结果时刷新。
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from contextlib import suppress
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .derived import (
    DerivedError,
    DerivedNotTranslatable,
    evaluate_derived,
    is_translatable,
    to_excel,
)
from .derived import validate_config as derived_validate_config
from .formula import (
    LONG_FORMULA_WARN,
    compile_formulas,
    excel_literal,
    guard_default,
    guarded_lookup,
    local_cell,
    long_formula_warning,
    longest_formula,
)
from .models import (
    FIRST_CASE_COLUMN,
    CaseData,
    ProjectConfig,
    RenderResult,
    VariableDef,
)
from .utils import (
    CodeGenError,
    DerivedValue,
    ExcelError,
    InputError,
    VarValue,
    WorkbookIOError,
    column_index_to_letter,
    fingerprint,
    parse_cell,
    to_text,
)

__all__ = [
    "GLOBAL_HEADERS",
    "GROUP_HEADER",
    "LOCAL_CASE_HEADER",
    "LOCAL_HEADERS",
    "META_MARKER",
    "case_anchor_map",
    "check_required_sheets",
    "create_template",
    "get_sheet",
    "input_fingerprint",
    "load_workbook_file",
    "output_fingerprint",
    "read_cases",
    "read_global_values",
    "read_group_members",
    "read_metadata",
    "template_source",
    "write_results",
    "write_run_scripts",
]

GLOBAL_HEADERS: tuple[str, ...] = ("Variable", "Value", "Description", "Prefix", "Suffix")
LOCAL_HEADERS: tuple[str, ...] = ("Variable", "Description", "Prefix", "Suffix")
#: 纵向 Local 表 A1 的表头（那一列写 Case 名）。
LOCAL_CASE_HEADER = "Case"

#: 隐藏 Template 表里"机器可读元信息块"的起始标记。
META_MARKER = "## spreadsheet-codegen-meta"

#: 改名前的标记（`excel_codegen` 时代写进工作簿的）。只为**向后兼容**保留：
#: 0.9.x 生成的工作簿里存的是它，读的时候要认得，写回时顺手换成新标记。
LEGACY_META_MARKERS = ("## excel-codegen-meta",)
_META_MARKERS = (META_MARKER, *LEGACY_META_MARKERS)

#: 元信息键名（Template 表与 HOWTO 表共用同一套键，读取方不必区分来源）。
META_TIME = "时间"
META_INPUT = "参数指纹"
META_OUTPUT = "输出指纹"

#: 清理旧结果时，向右/向下最多扫多少列/行（防止误伤同表里的无关内容）。
_SCAN_LIMIT = 256

#: Global 表列号
_GLOBAL_COL = {"name": 1, "value": 2, "description": 3, "prefix": 4, "suffix": 5}
#: 成员表（第三层作用域）列号：A=成员名，B 起一个变量一列
_GROUP_FIRST_VAR_COLUMN = 2
#: 成员表 A 列表头
GROUP_HEADER = "成员"
#: Local 表列号
_LOCAL_COL = {"name": 1, "description": 2, "prefix": 3, "suffix": 4}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_CASE_HEADER_FILL = PatternFill("solid", fgColor="EDEDED")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
#: 派生参数（自动计算）的格：淡绿，和"要你填"的黄色区分开
_DERIVED_FILL = PatternFill("solid", fgColor="E2EFDA")
_TOP_ALIGN = Alignment(vertical="top", wrap_text=False)


def _cell_or(cell_value: Any, default: Any) -> Any:
    """只有"单元格真的为空"才回落 ``default``；非空值原样返回。

    与 ``.strip()`` 的关键区别：``" m"``（带前导空格的单位）是**有内容**的值，
    必须原样保留 —— 否则 ``suffix: " m"`` 会被读成 ``"m"``，生成 ``340m``。
    """
    if cell_value is None:
        return default
    if isinstance(cell_value, str) and cell_value.strip() == "":
        return default
    return cell_value


def _text_or(cell_value: Any, default: Any) -> str:
    """:func:`_cell_or` 的文本版本（前缀/后缀用）。"""
    return to_text(_cell_or(cell_value, default))


def _prefill_value(variable: VariableDef) -> Any:
    """``init`` 往新格子里写什么：``prefill: false`` 时**什么都不写**（留空）。

    与"回落"分开是有意的：作者写 ``default`` 的本意常常只是"给新表一个提示值"，
    而运行时"空 = 用默认值"会让"某类型才有的字段"被默认值污染。
    两个开关各管一头，默认都开着（老行为不变）。
    """
    return variable.default if variable.prefill else None


# --------------------------------------------------------------------------- #
# 生成 Excel 模板
# --------------------------------------------------------------------------- #
def _fresh_value(variable: VariableDef) -> VarValue:
    """``init`` 刚建好、还没人动过时，这个变量在表里长什么样。

    ``prefill: false`` 的格子是**空的**（``init`` 根本不写），读回来就是回落值
    （``fallback`` 默认跟着 ``prefill`` 走，所以那就是空串）。基线必须照这个来 ——
    否则"刚生成的工作簿"会被误判成"人填过东西"，``init --force`` 莫名其妙被拦。
    """
    return VarValue(
        variable.default if variable.prefill else variable.effective_default,
        variable.prefix,
        variable.suffix,
    )


def default_input_fingerprint(config: ProjectConfig, *, cases: Sequence[str] | None = None) -> str:
    """「所有参数都还是 YAML 默认值」时的参数指纹。

    用来判断一本已有工作簿里"人有没有填过东西"：与 :func:`input_fingerprint` 同一套算法，
    只是参数全部取自配置（``default`` + ``prefix`` / ``suffix``），并且**按 ``prefill``
    的口径**（不预填的变量就是空的，见 :func:`_fresh_value`）。

    派生参数不在里面：它不是"人填的东西"，而是输入的函数 —— 比输入就够了，
    见 :func:`_input_only_fingerprint`。
    """
    global_values = {
        variable.name: _fresh_value(variable) for variable in config.global_variables if not variable.is_derived
    }
    case_values = {
        variable.name: _fresh_value(variable) for variable in config.local_variables if not variable.is_derived
    }
    names = list(cases) if cases is not None else ["Case1"]
    case_data = [
        CaseData(name=name, column=FIRST_CASE_COLUMN + offset, values=dict(case_values))
        for offset, name in enumerate(names)
    ]
    return input_fingerprint(global_values, case_data)


def _input_only_fingerprint(
    config: ProjectConfig,
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
) -> str:
    """只对**用户能填的**参数取指纹（派生参数剔掉）。

    为什么要剔：派生值是表达式算出来的。读工作簿时是"现算"（读表 → 求值），
    而基线是"直接取 default + prefix/suffix"，两套口径稍微差一点（哪怕只是
    ``prefill: false`` 的空格子），"刚生成的工作簿"就会被判成"人填过东西"。

    比输入既**充分**（派生值是输入的函数，输入没变派生值就不会变）又**稳**。
    """
    derived = {item.name for item in (*config.global_variables, *config.local_variables) if item.is_derived}
    only_globals = {name: value for name, value in global_values.items() if name not in derived}
    only_cases = [
        CaseData(
            name=case.name,
            column=case.column,
            values={name: value for name, value in case.values.items() if name not in derived},
        )
        for case in cases
    ]
    return input_fingerprint(only_globals, only_cases)


def workbook_deviates_from_defaults(path: str | Path, config: ProjectConfig) -> str | None:
    """这本工作簿里的参数是不是**已经偏离** YAML 默认值（即"人填过东西"）？

    :returns: 一句给人看的说明；"看起来还是默认值"或读不出来时返回 ``None``。

    判据是**参数指纹**：等于"全部取 default"的指纹时，这本表就是刚生成的、没人动过
    —— 覆盖它没有代价（老行为不变）。偏离了才提示，避免每次 ``--force`` 都来问。

    只看**输入**参数（派生值剔掉），也**不看渲染记录**：记录里存的是"上次渲染时的参数"，
    它等于当前参数只能说明"渲染过"，**完全不能说明参数还是默认值** —— 早先拿它当短路条件，
    于是"改了参数 → 渲染 → init --force"这条路会绕过防呆，把用户填的值静默换成默认值。
    """
    try:
        workbook = load_workbook_file(path)
    except CodeGenError:
        return None  # 读不出来（不是 Excel / 已损坏）：交给 init 自己的覆盖与报错逻辑
    try:
        try:
            values = read_global_values(workbook, config)
            cases = read_cases(workbook, config, global_values=values)
        except CodeGenError:
            return None  # 结构对不上（老版本生成的？）：保守起见不拦
        names = [case.name for case in cases]
        current = _input_only_fingerprint(config, values, cases)
        baseline = default_input_fingerprint(config, cases=names)
    finally:
        workbook.close()

    if not names:
        return None
    if current == baseline:
        return None  # 表里就是 YAML 的默认值 —— 刚生成、没人动过
    return (
        f"这本工作簿里 {len(names)} 个 Case 的参数已经不是 YAML 默认值"
        f"（当前指纹 {current}，默认值应为 {baseline}）—— 看起来人填过东西"
    )


def create_template(
    config: ProjectConfig,
    path: str | Path,
    *,
    cases: int | Sequence[str] = 2,
    overwrite: bool = False,
    include_template_sheet: bool | None = None,
    include_howto_sheet: bool | None = None,
    include_comments: bool = True,
    include_scripts: bool = True,
    base_dir: str | Path | None = None,
) -> Path:
    """按配置生成 Excel 模板文件。

    :param cases: 初始 Case 列数量（``int``）或直接给出 Case 名称列表。
    :param overwrite: 目标文件已存在时是否覆盖。
    :param include_template_sheet: ``None`` 时遵循配置；``True/False`` 强制生成/不生成隐藏 Template 表。
    :param include_howto_sheet: ``None`` 时遵循配置；``True/False`` 强制生成/不生成 HOWTO 说明表。
    :param include_comments: 是否给"变量名"那一格加批注（描述 / 单位 / 约束 / 前缀后缀 / 派生表达式）。
    :param include_scripts: 是否在工作簿旁边生成 ``*_render.bat`` / ``*_render.sh`` 一键刷新脚本。
    :param base_dir: 解析 ``template_file`` 相对路径的基准目录，默认使用 ``config.source_dir``。
    """
    target = Path(path)
    if target.exists() and not overwrite:
        raise ExcelError(f"Excel 模板已存在: {target}（需要覆盖请加 --force）")
    if target.exists() and target.is_dir():
        raise ExcelError(f"目标路径是目录: {target}")

    # 派生参数先过一遍配置期检查（语法 / 引用范围 / 循环），别等到渲染才炸
    derived_validate_config(config)
    # 模板也先过一遍：语法 / 片段 / **公式模式能不能表达**
    _preflight_templates(config, base_dir=base_dir or config.source_dir)

    case_names = _normalise_case_names(cases)
    config.scripts_enabled = include_scripts

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    global_name = config.excel.sheets.global_
    local_name = config.excel.sheets.local
    _build_global_sheet(workbook.create_sheet(global_name), config, comments=include_comments)
    _build_local_sheet(workbook.create_sheet(local_name), config, case_names, comments=include_comments)
    if config.group is not None:
        _build_group_sheet(
            workbook.create_sheet(config.group.sheet), config, config.group.members, comments=include_comments
        )

    for name in config.excel.sheets.outputs:
        workbook.create_sheet(name)

    sheet_name = config.excel.template_sheet
    if include_template_sheet is False:
        sheet_name = None
    elif include_template_sheet is True and not sheet_name:
        sheet_name = "Template"
    if sheet_name:
        if sheet_name in workbook.sheetnames:
            raise ExcelError(f"Template 工作表名 {sheet_name!r} 与已有工作表冲突，请修改 excel.template_sheet")
        template_sheet = workbook.create_sheet(sheet_name)
        _build_template_sheet(template_sheet, config, base_dir=base_dir or config.source_dir)
        template_sheet.sheet_state = "hidden"

    howto_name = config.excel.howto_sheet
    if include_howto_sheet is False:
        howto_name = None
    elif include_howto_sheet is True and not howto_name:
        howto_name = "HOWTO"
    if howto_name:
        if howto_name in workbook.sheetnames:
            raise ExcelError(f"HOWTO 工作表名 {howto_name!r} 与已有工作表冲突，请修改 excel.howto_sheet")
        howto_sheet = workbook.create_sheet(howto_name)
        # 放在第一张：打开工作簿先看到"下一步跑什么"
        workbook.move_sheet(howto_sheet, offset=-len(workbook.sheetnames) + 1)
        write_howto_sheet(howto_sheet, config, metadata=None)

    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(target)
    except OSError as exc:
        raise WorkbookIOError(f"无法写入 Excel 模板 {target}: {exc}（文件被 Excel 占用？）") from exc
    finally:
        workbook.close()

    if include_scripts:
        # 放在工作簿旁边：改完参数双击就能刷新，不用记命令
        write_run_scripts(config, target)
    return target


def _preflight_templates(config: ProjectConfig, *, base_dir: str | Path | None) -> None:
    """建表之前先把每个模板过一遍：语法、``{% include %}`` 片段、公式模式能否表达。

    为什么要在这里做：**公式模式是默认引擎**，而它只支持"占位符 + 单行 ``{% if %}``"。
    如果等到 ``render`` 才发现模板用了过滤器，那时工作簿已经生成、参数也填了一半 ——
    而这个工具的正常用法是"建一次工作簿，之后就在 Excel 里干活"，越早报错越好。
    报错里会明确说"请把该模板改回 engine: snapshot"。
    """
    from .formula import compile_formulas  # 本模块已导入，这里只是让依赖显式
    from .renderer import validate_template  # 延迟导入：renderer 依赖本模块

    for template in config.templates:
        validate_template(template, base_dir=base_dir)
        if template.engine != "excel":
            continue
        # 轴取哪个都行 —— 这里只关心"这一行能不能编译成公式"（与 validate 命令同一套检查）
        compile_formulas(
            template,
            config,
            case_axes=[FIRST_CASE_COLUMN],
            source=_template_source(template, base_dir),
        )


def _normalise_case_names(cases: int | Sequence[str]) -> list[str]:
    if isinstance(cases, int):
        if cases < 1:
            raise ExcelError("Case 列数量至少为 1")
        return [f"Case{index}" for index in range(1, cases + 1)]
    names = [to_text(item).strip() for item in cases]
    if not names or any(not name for name in names):
        raise ExcelError("Case 名称不能为空")
    if len(set(names)) != len(names):
        raise ExcelError(f"Case 名称重复: {', '.join(names)}")
    return names


def _build_global_sheet(worksheet: Worksheet, config: ProjectConfig, *, comments: bool = True) -> None:
    _write_headers(worksheet, GLOBAL_HEADERS)
    for row, variable in enumerate(config.global_variables, start=2):
        name_cell = worksheet.cell(row=row, column=_GLOBAL_COL["name"], value=variable.name)
        if comments:
            _attach_comment(name_cell, variable, where="Global 表 B 列（所有 Case 共用）")
        if variable.is_derived:
            _write_derived_cell(
                worksheet.cell(row=row, column=_GLOBAL_COL["value"]),
                variable,
                resolve=_global_resolver(config),
            )
        else:
            value_cell = worksheet.cell(row=row, column=_GLOBAL_COL["value"], value=_prefill_value(variable))
            value_cell.fill = _INPUT_FILL
            value_cell.alignment = _TOP_ALIGN
            _add_value_validation(worksheet, variable, [value_cell.coordinate])
        worksheet.cell(row=row, column=_GLOBAL_COL["description"], value=_described(variable))
        worksheet.cell(row=row, column=_GLOBAL_COL["prefix"], value=variable.prefix)
        worksheet.cell(row=row, column=_GLOBAL_COL["suffix"], value=variable.suffix)
    for index, width in enumerate((26, 30, 46, 16, 16), start=1):
        worksheet.column_dimensions[column_index_to_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def _build_local_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    case_names: Sequence[str],
    *,
    comments: bool = True,
) -> None:
    """搭 Local Parameter 表：布局由 ``excel.local_direction`` 决定。

    * ``horizontal``（默认）：变量做行，一个工况一列（E1 起写 Case 名）；
    * ``vertical``：变量做列，一个工况一行（A2 起写 Case 名）。

    两种布局都**没有**给 Prefix / Suffix 留列吗？不是 —— 横向布局保留 C/D 两列
    （老行为，可在表里覆盖 YAML 的前后缀）；纵向布局第 1 行整行都是变量名，
    所以前后缀只来自 YAML，表头格的批注里写着它们是什么。
    """
    if config.excel.local_direction == "vertical":
        _build_local_sheet_vertical(worksheet, config, case_names, comments=comments)
    else:
        _build_local_sheet_horizontal(worksheet, config, case_names, comments=comments)


def _build_local_sheet_horizontal(
    worksheet: Worksheet,
    config: ProjectConfig,
    case_names: Sequence[str],
    *,
    comments: bool,
) -> None:
    _write_headers(worksheet, list(LOCAL_HEADERS) + list(case_names))
    for row, variable in enumerate(config.local_variables, start=2):
        name_cell = worksheet.cell(row=row, column=_LOCAL_COL["name"], value=variable.name)
        if comments:
            _attach_comment(name_cell, variable, where="Local 表 E 列起（一个 Case 一列）")
        worksheet.cell(row=row, column=_LOCAL_COL["description"], value=_described(variable))
        worksheet.cell(row=row, column=_LOCAL_COL["prefix"], value=variable.prefix)
        worksheet.cell(row=row, column=_LOCAL_COL["suffix"], value=variable.suffix)
        input_cells: list[str] = []
        for offset in range(len(case_names)):
            column = FIRST_CASE_COLUMN + offset
            cell = worksheet.cell(row=row, column=column)
            if variable.is_derived:
                _write_derived_cell(cell, variable, resolve=_local_resolver(config, CaseData(name="", column=column)))
            else:
                cell.value = _prefill_value(variable)
                cell.fill = _INPUT_FILL
                cell.alignment = _TOP_ALIGN
                input_cells.append(cell.coordinate)
        _add_value_validation(worksheet, variable, input_cells)
    for index, width in enumerate((24, 40, 16, 16), start=1):
        worksheet.column_dimensions[column_index_to_letter(index)].width = width
    for offset, name in enumerate(case_names):
        column = FIRST_CASE_COLUMN + offset
        worksheet.column_dimensions[column_index_to_letter(column)].width = max(16, min(36, len(name) + 14))
    worksheet.freeze_panes = "E2"


def _build_local_sheet_vertical(
    worksheet: Worksheet,
    config: ProjectConfig,
    case_names: Sequence[str],
    *,
    comments: bool,
) -> None:
    """纵向 Local 表：一行一个工况（第 1 行是变量名），方便整块粘贴参数。"""
    worksheet.cell(row=1, column=_LOCAL_COL["name"], value=LOCAL_CASE_HEADER).font = _HEADER_FONT
    worksheet.cell(row=1, column=_LOCAL_COL["name"]).fill = _CASE_HEADER_FILL
    for offset, variable in enumerate(config.local_variables):
        column = _GROUP_FIRST_VAR_COLUMN + offset
        header = worksheet.cell(row=1, column=column, value=variable.name)
        if comments:
            _attach_comment(header, variable, where=f"Local 表 {column_index_to_letter(column)} 列（一个 Case 一行）")
        worksheet.column_dimensions[column_index_to_letter(column)].width = max(16, min(36, len(variable.name) + 14))
    for offset in range(len(case_names)):
        row = 2 + offset
        case_cell = worksheet.cell(row=row, column=_LOCAL_COL["name"], value=case_names[offset])
        case_cell.fill = _CASE_HEADER_FILL
        case_cell.font = _HEADER_FONT
    last_row = 1 + len(case_names)

    for offset, variable in enumerate(config.local_variables):
        column = _GROUP_FIRST_VAR_COLUMN + offset
        input_cells: list[str] = []
        for index in range(len(case_names)):
            row = 2 + index
            cell = worksheet.cell(row=row, column=column)
            if variable.is_derived:
                _write_derived_cell(cell, variable, resolve=_local_resolver(config, CaseData(name="", row=row)))
            else:
                cell.value = _prefill_value(variable)
                cell.fill = _INPUT_FILL
                cell.alignment = _TOP_ALIGN
                input_cells.append(cell.coordinate)
        if input_cells:
            letter = column_index_to_letter(column)
            _add_value_validation(worksheet, variable, [f"{letter}2:{letter}{last_row}"])
    worksheet.column_dimensions[column_index_to_letter(_LOCAL_COL["name"])].width = 22
    worksheet.freeze_panes = "B2"


def _variable_comment(variable: VariableDef, *, where: str) -> str:
    """变量名那格的批注正文：把"这一格填什么"一次说清。"""
    lines: list[str] = []
    if variable.description:
        lines.append(variable.description)
        lines.append("")
    if variable.is_derived:
        lines.append(f"自动计算：{variable.derived}")
        lines.append("不用手填；改了它引用的输入后会自动重算。")
    else:
        lines.append(f"填写位置：{where}")
    if variable.unit:
        lines.append(f"单位：{variable.unit}")
    lines.append(f"类型：{variable.type}")
    if variable.has_constraints:
        lines.append(f"约束：{variable.constraint_text}")
    if variable.prefix or variable.suffix:
        lines.append(f"前缀 / 后缀：{variable.prefix!r} / {variable.suffix!r}")
    if variable.prefill:
        if to_text(variable.default) != "":
            lines.append(f"默认值：{to_text(variable.default)}")
    else:
        lines.append("默认值：（本变量不预填 —— 新表格里这一格是空的）")
    # "清空这一格"是什么意思，必须写在格子上：默认语义下它会回落到默认值
    if not variable.fallback:
        lines.append("★ 留空 = 这个参数没有值（本变量不回落到默认值）")
    elif to_text(variable.default) != "":
        lines.append("留空 = 使用上面的默认值")
    lines.append(f"模板里引用：{{{{ {variable.name} }}}}")
    return "\n".join(lines)


def _attach_comment(cell, variable: VariableDef, *, where: str) -> None:
    """给"变量名"那一格加批注。

    只加在名字格（A 列）而不是每个取值格：取值格已经有数据有效性的输入提示，
    而 A 列是冻结的、永远可见 —— 鼠标一放就知道这是什么、该填什么、有没有约束。
    """
    comment = Comment(_variable_comment(variable, where=where), "spreadsheet_codegen")
    comment.width = 340
    comment.height = 190
    cell.comment = comment


def _relative_to(path: Path, base: Path, *, windows: bool) -> str:
    """把 ``path`` 表示成相对 ``base`` 的路径，并按目标平台换算分隔符。"""
    import os

    try:
        text = os.path.relpath(path, base)
    except ValueError:  # 跨盘符（Windows）时 relpath 会失败，退回绝对路径
        text = str(path)
    if windows:
        return text.replace("/", "\\")
    return text.replace("\\", "/")


def _render_args(config: ProjectConfig, target: Path, *, windows: bool) -> str:
    """生成 render 的**参数**部分（不含 ``render`` 这个词，也不含怎么调用）。

    拆出来是因为"进程内的模块调用"与"命令行调用"拼法不同：
    ``python -m spreadsheet_codegen render -c …`` 里有 ``-m`` 这一段。
    """
    x_flag = target.name if target.parent else str(target)
    if config.config_path is not None:
        c_flag = _relative_to(config.config_path, target.parent or Path("."), windows=windows)
    else:  # 直接调库、没经过 load_config 时拿不到 YAML 路径：交给用户自己改
        c_flag = "<你的配置>.yaml"
    return f'render -c "{c_flag}" -x "{x_flag}" --write-excel'


def _render_command(config: ProjectConfig, target: Path, *, windows: bool) -> str:
    """兼容旧签名：命令行形态的 render 命令（不含程序名）。

    .. deprecated:: 0.11.0
       生成脚本改用 :func:`_render_args` + 完整回退链；本函数只给调用方与测试留着。
    """
    return _render_args(config, target, windows=windows)


def _bat_text(config: ProjectConfig, target: Path) -> str:
    """Windows 批处理：按"最可能可用"的顺序试遍三种调用方式。

    为什么要有回退链：只会在 PATH 上找 ``spreadsheet-codegen`` 的脚本，在"用源码 /
    ``python -m`` 跑"的机器上**两条路都不通**，而失败提示还指向"依赖没装"这个
    不存在的原因。包名是 ``spreadsheet_codegen``（下划线）、命令才是
    ``spreadsheet-codegen``（连字符）—— ``python -m`` 那条路最稳，必须覆盖。

    ``chcp 65001`` 是必需的：本文件里有中文，而中文 Windows 的控制台默认是 GBK，
    会把 UTF-8 字节读成乱码甚至当成命令去执行。
    """
    args = _render_args(config, target, windows=True)
    return f"""@echo off
REM ===========================================================================
REM  由 spreadsheet_codegen 生成 —— 改完参数双击本文件即可把结果写回 Output 表。
REM  重新生成工作簿（init）时会一并覆盖本文件。
REM ===========================================================================
REM  本文件是 UTF-8 编码，先切到 65001 代码页，否则中文提示在 GBK 控制台是乱码。
chcp 65001 >nul
cd /d "%~dp0"

REM 依次尝试：uv + 安装好的命令 / uv + 模块 / 本机 python + 模块。
REM SPREADSHEET_CODEGEN_PY 可以指定一个"能跑 -m spreadsheet_codegen"的解释器。
set "SPREADSHEET_CODEGEN_OK="

where uv >nul 2>nul
if %errorlevel%==0 (
  uv run spreadsheet-codegen {args} && set "SPREADSHEET_CODEGEN_OK=1"
  if defined SPREADSHEET_CODEGEN_OK goto :done
  uv run python -m spreadsheet_codegen {args} && set "SPREADSHEET_CODEGEN_OK=1"
  if defined SPREADSHEET_CODEGEN_OK goto :done
)

if defined SPREADSHEET_CODEGEN_PY (
  "%SPREADSHEET_CODEGEN_PY%" -m spreadsheet_codegen {args} && set "SPREADSHEET_CODEGEN_OK=1"
  if defined SPREADSHEET_CODEGEN_OK goto :done
)

spreadsheet-codegen {args} && set "SPREADSHEET_CODEGEN_OK=1"
if defined SPREADSHEET_CODEGEN_OK goto :done
python -m spreadsheet_codegen {args} && set "SPREADSHEET_CODEGEN_OK=1"
if defined SPREADSHEET_CODEGEN_OK goto :done

:done
echo.
if defined SPREADSHEET_CODEGEN_OK (
  echo [OK] Done. Open the Output sheet in Excel to see the result.
) else (
  echo [FAILED] None of the ways to run spreadsheet_codegen worked.
  echo           Tried: uv run spreadsheet-codegen / uv run python -m spreadsheet_codegen
  echo                  %%SPREADSHEET_CODEGEN_PY%% -m spreadsheet_codegen / spreadsheet-codegen / python -m spreadsheet_codegen
  echo           Please make sure the tool is installed: pip install spreadsheet-codegen
  echo           or set SPREADSHEET_CODEGEN_PY to a Python that can run -m spreadsheet_codegen.
  echo           Also check that Excel is not holding this file open.
)
REM pause 放在结论之后：双击运行时窗口留得住、看得见结论。
REM 退出码必须**在 pause 之后显式给** —— echo / pause 这些内建命令会把 ERRORLEVEL
REM 归零，不写这两行的话"五条路全失败"照样以 0 结束，批处理与 CI 都看不出失败。
pause
if defined SPREADSHEET_CODEGEN_OK (exit /b 0) else (exit /b 1)
"""


def _sh_text(config: ProjectConfig, target: Path) -> str:
    """POSIX shell：与 ``.bat`` 同一套回退链（顺序一致，日志才可对照）。

    这里有两处**很容易写错、而且错了不报错**的地方（都真踩过）：

    * ``status`` 初始不能当成"已经试过并成功"，要用独立的 ``solved`` 标记。
      早先的写法是 ``status=0`` + 每条路都套一个 ``[ "$status" -ne 0 ]`` ——
      在**没装 uv** 的机器上第一条路根本不执行，``status`` 永远是 0，于是后面每条路
      都被自己的守卫挡掉：脚本一件正事没干，却打印"[完成]"并以 0 退出。
      这正是本版要消灭的"静默假成功"。
    * 只有**命令不存在**时才该退到下一条路。``spreadsheet-codegen`` 存在却报错，
      那是工具真的失败（工作簿被 Excel 占着、取值不合法……），再退到别的解释器
      会把真错误盖成"没装"，也会让 ``check`` 的退出码在 CI 里失真。
    """
    args = _render_args(config, target, windows=False)
    return f"""#!/usr/bin/env bash
# ===========================================================================
#  由 spreadsheet_codegen 生成 —— 改完参数跑一次本文件即可把结果写回 Output 表。
#  重新生成工作簿（init）时会一并覆盖本文件。
# ===========================================================================
set -uo pipefail
cd "$(dirname "$0")"

# 依次尝试：uv + 安装好的命令 / uv + 模块 / 本机 python + 模块。
# SPREADSHEET_CODEGEN_PY 可以指定一个"能跑 -m spreadsheet_codegen"的解释器。
status=0
solved=0
# 跑一条路；成功（退出码 0）就定案，后面的路不再执行。
# 函数本身恒定返回 0：调用方只用全局的 status / solved，不靠返回值判断。
run_route() {{
  if [ "$solved" -ne 0 ]; then return 0; fi
  "$@"
  status=$?
  if [ "$status" -eq 0 ]; then solved=1; fi
  return 0
}}

if command -v uv >/dev/null 2>&1; then
  run_route uv run spreadsheet-codegen {args}
  run_route uv run python -m spreadsheet_codegen {args}
fi

if [ -n "${{SPREADSHEET_CODEGEN_PY:-}}" ]; then
  run_route "$SPREADSHEET_CODEGEN_PY" -m spreadsheet_codegen {args}
fi

# 最后两条：命令与模块。命令**存在**时只试它 —— 见函数文档：它若报错，那是工具
# 真的失败，退出码要原样带出去（CI 靠 1/2/3/4 分流），不该被"没装"的假象盖住。
if command -v spreadsheet-codegen >/dev/null 2>&1; then
  run_route spreadsheet-codegen {args}
else
  run_route python3 -m spreadsheet_codegen {args}
fi

echo
if [ "$solved" -ne 0 ]; then
  echo "[完成] 回到 Excel 打开「Output」表看结果。"
else
  echo "[失败] 这些调用方式都没成功：uv run spreadsheet-codegen / uv run python -m spreadsheet_codegen /"
  echo "       spreadsheet-codegen / python3 -m spreadsheet_codegen。"
  echo "       请确认已安装（pip install spreadsheet-codegen），"
  echo "       或设置 SPREADSHEET_CODEGEN_PY 指向能跑 -m spreadsheet_codegen 的 Python。"
  echo "       也请确认 Excel / WPS 没有占着这个文件。"
fi
exit "$status"
"""


def write_run_scripts(config: ProjectConfig, target: Path) -> list[Path]:
    """在**工作簿旁边**生成 ``<工作簿名>_render.bat`` 与 ``<工作簿名>_render.sh``。

    为什么要它：目标用户是工程师，不是终端爱好者。HOWTO 表里写了命令，但还得自己开终端敲；
    双击脚本就能"改完参数 → 刷新 Output 表"，而把 uv / venv / 源码运行方式的差异封在脚本里。

    两个平台都生成（不是只生成当前的）：一本工作簿常常在 Windows 和 Linux 之间传来传去。
    """
    directory = target.parent or Path(".")
    stem = target.stem
    written: list[Path] = []
    for name, text in (
        (f"{stem}_render.bat", _bat_text(config, target)),
        (f"{stem}_render.sh", _sh_text(config, target)),
    ):
        path = directory / name
        path.write_text(text, encoding="utf-8", newline="\r\n" if name.endswith(".bat") else "\n")
        if name.endswith(".sh"):
            path.chmod(path.stat().st_mode | 0o111)  # 让 Linux/macOS 上可以直接 ./ 跑
        written.append(path)
    return written


def _build_group_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    members: Sequence[str],
    *,
    comments: bool = True,
) -> None:
    """建"成员表"：**一行一个成员，B 列起一个变量一列**（表头是变量名）。

    为什么不像 Local 那样"变量做行"：工程师写舱容表就是一行一个舱；而且这样公式模式
    能用与 global / local 同一形态的 ``INDEX/MATCH`` 定位（一次一维查找）。
    """
    group = config.group
    if group is None:
        return
    _write_headers(worksheet, [GROUP_HEADER, *group.names])
    for offset, variable in enumerate(group.variables):
        cell = worksheet.cell(row=1, column=_GROUP_FIRST_VAR_COLUMN + offset)
        if comments:
            _attach_comment(cell, variable, where=f"成员表（{group.sheet}）：一行一个成员")

    for row, member in enumerate(members, start=2):
        worksheet.cell(row=row, column=1, value=member)
        for offset, variable in enumerate(group.variables):
            cell = worksheet.cell(row=row, column=_GROUP_FIRST_VAR_COLUMN + offset, value=_prefill_value(variable))
            cell.fill = _INPUT_FILL
            cell.alignment = _TOP_ALIGN

    for offset, variable in enumerate(group.variables):
        column = _GROUP_FIRST_VAR_COLUMN + offset
        cells = [worksheet.cell(row=row, column=column).coordinate for row in range(2, len(members) + 2)]
        _add_value_validation(worksheet, variable, cells)

    worksheet.column_dimensions["A"].width = 24
    for offset in range(len(group.variables)):
        letter = column_index_to_letter(_GROUP_FIRST_VAR_COLUMN + offset)
        worksheet.column_dimensions[letter].width = 20
    worksheet.freeze_panes = "B2"


def _add_value_validation(worksheet: Worksheet, variable: VariableDef, cells: Sequence[str]) -> None:
    """把变量声明的取值约束写成 Excel 的**数据有效性**（下拉列表 / 数值范围）。

    这是"挡在输入口"的第一道闸，方便人填；**判据仍然是** :func:`check_value_constraints` ——
    读回来的取值一律再查一遍。原因：数据有效性挡不住粘贴、脚本写入和别人发来的老文件，
    而且 Excel 的下拉列表总允许留空（``allow_blank=True``），而工具侧默认认为
    "声明了约束就不许为空"（想放行要写 ``allow_blank: true``）。

    约束本身表达不了时（下拉列表的选项里带逗号、或拼起来超过 Excel 的 255 字符上限）
    直接报错 —— 与其写一个悄悄失效的校验，不如让人知道。
    """
    if not cells or not variable.has_constraints:
        return
    if not variable.choices and variable.min is None and variable.max is None:
        # 只有 pattern：Excel 的数据有效性没有正则，写不出有意义的校验。
        # 不写总比写一个乱报错的强 —— pattern 由 check_value_constraints 在 render/validate 时检查。
        return

    if variable.choices:
        # 空串只是"允许留空"的一种表达（``choices: ["", "储蓄卡"]``）；它是校验层的概念，
        # 不该出现在下拉列表里 —— 那会让第一项看起来像"允许空"的空选项。
        allowed = [to_text(item) for item in variable.choices if to_text(item) != ""]
        if not allowed:
            raise ExcelError(
                f"变量 {variable.name!r} 的 choices 里只有空串 —— 那不是约束。"
                "想表达'可以留空'请写 allow_blank: true，否则删掉 choices"
            )
        if any("," in item for item in allowed):
            raise ExcelError(
                f"变量 {variable.name!r} 的 choices 里有取值含逗号（{', '.join(allowed)}）—— "
                "Excel 下拉列表用逗号分隔，表达不了；请改写取值或去掉 choices"
            )
        joined = ",".join(allowed)
        if len(joined) > 255:
            raise ExcelError(
                f"变量 {variable.name!r} 的 choices 合计超过 Excel 下拉列表的 255 字符上限"
                f"（当前 {len(joined)}）；请减少选项，或去掉 choices（取值仍会在 render / validate 时检查）"
            )
        validation = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
    else:
        operator: str
        low: float | None
        high: float | None
        if variable.min is not None and variable.max is not None:
            operator, low, high = "between", variable.min, variable.max
        elif variable.min is not None:
            operator, low, high = "greaterThanOrEqual", variable.min, None
        else:
            operator, low, high = "lessThanOrEqual", variable.max, None
        validation = DataValidation(
            type="whole" if variable.type == "int" else "decimal",
            operator=operator,
            formula1=to_text(low),
            formula2=None if high is None else to_text(high),
            allow_blank=True,
        )

    prompt = variable.constraint_text
    validation.promptTitle = variable.name
    validation.prompt = prompt
    validation.showInputMessage = True
    validation.errorTitle = "取值不合规"
    validation.error = f"{variable.name}：{prompt}"
    validation.showErrorMessage = True

    worksheet.add_data_validation(validation)
    for coordinate in cells:
        validation.add(coordinate)


# --------------------------------------------------------------------------- #
# 派生参数（derived:）：参数表里那一格写成公式（能翻译时）或算好的值
# --------------------------------------------------------------------------- #
def _described(variable: VariableDef) -> str:
    if not variable.is_derived:
        return variable.description
    expression = variable.derived or ""
    note = f"（自动计算：{expression}）"
    return f"{variable.description} {note}".strip() if variable.description else note


#: ``derived.to_excel`` 会调用的属性名 -> "取哪一列"。``value`` / ``text`` 都是取值列。
_DERIVED_VALUE_ALIASES = frozenset({"value", "text"})


def _attribute_column(attribute: str | None, definition: VariableDef) -> tuple[str, int]:
    """把表达式里的属性名翻成"取哪一列 + 默认值取哪个字段"。

    ``.value`` / ``.text`` 与裸变量都指向**取值**（默认值字段叫 ``default``；
    前缀 / 后缀的默认值字段与属性同名）。

    报错而不是静默降级：**属性拼错时最坏的结果是悄悄写进一个算好的值**
    （公式模式的核心承诺就没了），所以这里必须抛 :class:`DerivedError`。
    """
    if attribute is None or attribute in _DERIVED_VALUE_ALIASES:
        return "default", _GLOBAL_COL["value"]
    if attribute in ("prefix", "suffix"):
        return attribute, _GLOBAL_COL[attribute]
    raise DerivedError(
        f"派生表达式里的 .{attribute}（变量 {definition.name!r}）不支持："
        "只有 .value / .text / .prefix / .suffix 四个属性"
    )


def _global_resolver(config: ProjectConfig):
    """派生表达式里的变量名 -> Excel 引用（Global 表的取值列 / 前后缀列）。

    签名是 ``resolve(name, attribute=None)`` —— ``derived.to_excel`` 翻译 ``x.value``
    这类属性访问时会带上第二个参数。（此前这个回调只收一个参数，于是**任何** ``.value``
    都会被当成"上下文不支持属性访问"而降级成"往格子里写算好的值"：公式模式下参数改了
    它不再自动重算，而且没有任何提示。见 0.11.0 的 CHANGELOG。）
    """
    sheet = config.excel.sheets.global_
    globals_by_name = {item.name: item for item in config.global_variables}

    def resolve(name: str, attribute: str | None = None) -> str:
        definition = globals_by_name.get(name)
        if definition is None:
            raise DerivedError(
                f"派生表达式引用了 {name!r}：global 的派生参数只能引用 global 变量（不能引用 local，也不存在别的表）"
            )
        which, column = _attribute_column(attribute, definition)
        default = getattr(definition, which)
        if which == "default":
            default = definition.effective_default
        return guarded_lookup(sheet, name, column, default, absolute=True)

    return resolve


def _local_resolver(config: ProjectConfig, case: CaseData):
    """派生表达式里的变量名 -> Excel 引用（本 Case 的那一格 / Global 取值列 / 前后缀）。

    同样的 ``resolve(name, attribute=None)`` 签名要求，见 :func:`_global_resolver`。
    """
    local_sheet = config.excel.sheets.local
    global_sheet = config.excel.sheets.global_
    locals_by_name = {item.name: item for item in config.local_variables}
    globals_by_name = {item.name: item for item in config.global_variables}
    horizontal = config.excel.local_direction == "horizontal"

    def resolve(name: str, attribute: str | None = None) -> str:
        if name in locals_by_name:
            definition = locals_by_name[name]
            which, _ = _attribute_column(attribute, definition)
            default = definition.effective_default if which == "default" else getattr(definition, which)
            # 派生格里没有"拖动"语义：坐标锁死（参数表由工具维护）
            if which == "prefix":
                column = _LOCAL_COL["prefix"]
            elif which == "suffix":
                column = _LOCAL_COL["suffix"]
            else:
                column = None
            if horizontal:
                assert case.column is not None
                expr = local_cell(local_sheet, name, column=column or case.column, absolute=True)
            else:
                assert case.row is not None
                if column is not None:
                    # 纵向布局没有前后缀列，它们只来自 YAML —— 写常量，与模板侧一致
                    return excel_literal(default)
                expr = local_cell(local_sheet, name, row=case.row, absolute=True)
            return guard_default(expr, default)
        if name in globals_by_name:
            definition = globals_by_name[name]
            which, column = _attribute_column(attribute, definition)
            default = definition.effective_default if which == "default" else getattr(definition, which)
            return guarded_lookup(global_sheet, name, column, default, absolute=True)
        raise DerivedError(f"派生表达式引用了未定义的变量 {name!r}")

    return resolve


def _write_derived_cell(cell, variable: VariableDef, *, resolve, value: Any = None) -> None:
    """写派生参数格：能翻译就写公式，否则写算好的值（没有值就先留空）。"""
    cell.fill = _DERIVED_FILL
    cell.alignment = _TOP_ALIGN
    try:
        formula = to_excel(variable.derived or "", name=variable.name, resolve=resolve)
    except DerivedNotTranslatable:
        formula = None
    if formula is not None:
        cell.value = "=" + formula
        return
    # 翻译不了：只能写入 Python 算好的值（没有就留空，等 --write-excel）
    cell.value = value if value is not None else None


def _build_template_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    *,
    base_dir: str | Path | None = None,
) -> None:
    hint = worksheet.cell(
        row=1,
        column=1,
        value="!! 只读参考：模板真源是 YAML / template_file；在这里改模板不会影响渲染结果",
    )
    hint.font = Font(bold=True, color="C00000")
    row = 2
    for template in config.templates:
        title = worksheet.cell(row=row, column=1, value=f"### {template.name}")
        title.font = _HEADER_FONT
        row += 1
        metadata = (
            ("output_sheet", template.output_sheet),
            ("start_cell", template.start_cell),
            ("direction", template.direction),
            ("filename", template.filename or ""),
            ("extension", template.extension),
            ("description", template.description),
        )
        for key, value in metadata:
            worksheet.cell(row=row, column=1, value=key)
            worksheet.cell(row=row, column=2, value=to_text(value))
            row += 1

        source = template.source_code
        if template.template_file:
            file_path = Path(template.template_file)
            if base_dir and not file_path.is_absolute():
                file_path = Path(base_dir) / file_path
            try:
                source = file_path.read_text(encoding="utf-8")
            except OSError as exc:
                source = f"<<无法读取模板文件 {file_path}: {exc}>>"
        worksheet.cell(row=row, column=1, value="code ↓").font = _HEADER_FONT
        row += 1
        for line_number, line in enumerate(source.splitlines(), start=1):
            worksheet.cell(row=row, column=1, value=line_number)
            worksheet.cell(row=row, column=2, value=line)
            row += 1
        row += 1

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 110


def _write_headers(worksheet: Worksheet, headers: Sequence[str]) -> None:
    for index, text in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=text)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.fill = _CASE_HEADER_FILL if index >= FIRST_CASE_COLUMN else _HEADER_FILL


# --------------------------------------------------------------------------- #
# HOWTO 表：把"下一步跑什么"写进工作簿本身
# --------------------------------------------------------------------------- #
_HOWTO_TITLE = Font(bold=True, size=14, color="1F3864")
_HOWTO_HEAD = Font(bold=True, size=11, color="1F3864")
_HOWTO_MONO = Font(name="Consolas", size=10)
_HOWTO_NOTE = Font(size=9, color="606060")
_HOWTO_WARN = Font(bold=True, size=10, color="C00000")
_HOWTO_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_howto_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    *,
    metadata: Mapping[str, str] | None = None,
    command: str | None = None,
) -> None:
    """重写 HOWTO 表：三步说明 + 快照提醒 + 本次生成记录。"""
    for row in worksheet.iter_rows(min_row=1, max_row=max(worksheet.max_row, 1), max_col=2):
        for cell in row:
            cell.value = None

    lines: list[tuple[str, Font, PatternFill | None]] = []
    add = lambda text, font, fill=None: lines.append((text, font, fill))  # noqa: E731

    add(f"{config.excel.output} —— 参数填写与代码生成说明", _HOWTO_TITLE, None)
    add("", _HOWTO_NOTE, None)
    add("这个工作簿怎么用", _HOWTO_HEAD, None)
    add(f"  1.  {config.excel.sheets.global_} 的 B 列：全局变量取值（所有 Case 共用）。", _HOWTO_NOTE, None)
    if config.excel.local_direction == "horizontal":
        add(f"  2.  {config.excel.sheets.local} 从 E 列起：每个 Case 一列，右拉复制即可增加。", _HOWTO_NOTE, None)
    else:
        add(f"  2.  {config.excel.sheets.local} 从第 2 行起：每个 Case 一行，下拉复制即可增加。", _HOWTO_NOTE, None)
    all_formula = bool(config.templates) and all(item.engine == "excel" for item in config.templates)
    if all_formula:
        # 全公式模式（默认）：改完参数 Output 表自己就重算了，不需要跑命令 —— 别再让人白跑一趟
        add("  3.  改完参数直接看输出表：里面是公式，Excel 会自己重算，**不用跑任何命令**。", _HOWTO_NOTE, None)
        add("      只有『要把代码导成文件』时才回到命令行：", _HOWTO_NOTE, None)
        add(
            f"          spreadsheet-codegen render -c <配置>.yaml -x {config.excel.output} --outdir <目录>",
            _HOWTO_MONO,
            _HOWTO_FILL,
        )
    else:
        add("  3.  改完参数后回到命令行执行：", _HOWTO_NOTE, None)
        add(
            f"          spreadsheet-codegen render -c <配置>.yaml -x {config.excel.output} --write-excel",
            _HOWTO_MONO,
            _HOWTO_FILL,
        )
        add("      只想导出代码文件就换成 --outdir <目录>；只预览不写回则什么参数都不加。", _HOWTO_NOTE, None)
    if config.scripts_enabled:
        stem = Path(config.excel.output).stem
        add(
            f"      ★ 懒得开终端就双击本文件旁边的 {stem}_render.bat（Windows）"
            f"或跑 {stem}_render.sh（Linux / macOS）。",
            _HOWTO_NOTE,
            None,
        )
    add("      想确认表里的代码是不是已经过期：spreadsheet-codegen check -c <配置>.yaml", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)
    formula_templates = [t for t in config.templates if t.engine == "excel"]
    snapshot_templates = [t for t in config.templates if t.engine != "excel"]
    if formula_templates:
        add("★  部分输出表是公式（engine: excel）：改参数后 Excel 打开即重算，不用跑脚本", _HOWTO_HEAD, None)
        add(
            "   公式支持占位符替换与行内分支（{% if %}）；要生成代码文件（--outdir）或做 CI 检查（check）仍需命令行。",
            _HOWTO_NOTE,
            None,
        )
    if snapshot_templates:
        add("⚠  部分输出表是「快照」，不是活公式", _HOWTO_WARN, None)
        add("   在 Excel 里改了参数、但没跑上面那条命令，输出表里的代码还是上一次的。", _HOWTO_NOTE, None)
    add("   不带 --write-excel 的 render 只预览，不会改变本文件。", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)
    add("输出位置", _HOWTO_HEAD, None)
    for template in config.templates:
        where = f"{template.output_sheet} @ {template.start_cell} ({template.direction})"
        engine = "公式·自动重算" if template.engine == "excel" else "快照·需重跑"
        extra = f"，case_filter: {template.case_filter}" if template.case_filter else ""
        add(f"  {template.name}: {where}［{engine}］{extra}", _HOWTO_MONO, None)
    if config.excel.template_sheet:
        add(f"  {config.excel.template_sheet}（隐藏）：模板原文，只读参考，改它不影响渲染结果。", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)

    add("本次生成", _HOWTO_HEAD, None)
    if metadata:
        for key, value in metadata.items():
            add(f"  {key:<10}{value}", _HOWTO_NOTE, None)
        if command:
            add(f"  {'命令':<8}{command}", _HOWTO_MONO, None)
    else:
        add("  （尚未渲染：本文件由 spreadsheet-codegen init 生成，还没有写回结果）", _HOWTO_WARN, None)
    add("", _HOWTO_NOTE, None)
    add("工况一览（来自上一次渲染）", _HOWTO_HEAD, None)
    cases = (metadata or {}).get("cases") or "（未知）"
    add(f"  {cases}", _HOWTO_MONO, None)

    for offset, (text, font, fill) in enumerate(lines, start=1):
        cell = worksheet.cell(row=offset, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(vertical="top", wrap_text=False)
        if fill is not None:
            cell.fill = fill
    worksheet.column_dimensions["A"].width = 110


def _meta_rows(metadata: Mapping[str, str]) -> list[tuple[str, str]]:
    return [(key, to_text(value)) for key, value in metadata.items()]


def _write_meta_block(worksheet: Worksheet, metadata: Mapping[str, str]) -> None:
    """在隐藏 Template 表末尾写入/刷新机器可读元信息块。"""
    marker_row = None
    for row in range(1, worksheet.max_row + 1):
        if to_text(worksheet.cell(row=row, column=1).value).strip() in _META_MARKERS:
            marker_row = row
            break
    start = marker_row if marker_row else worksheet.max_row + 2
    if marker_row:
        for row in range(marker_row, worksheet.max_row + 1):
            for column in (1, 2):
                worksheet.cell(row=row, column=column).value = None

    cell = worksheet.cell(row=start, column=1, value=META_MARKER)
    cell.font = _HEADER_FONT
    for offset, (key, value) in enumerate(_meta_rows(metadata), start=1):
        worksheet.cell(row=start + offset, column=1, value=key)
        worksheet.cell(row=start + offset, column=2, value=value)


def read_metadata(workbook: Workbook, config: ProjectConfig) -> dict[str, str]:
    """读回渲染元信息（时间 / 指纹 / 命令）。没写过则返回空字典。"""
    sheet_name = config.excel.template_sheet
    if sheet_name and sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        collecting = False
        found: dict[str, str] = {}
        for row in range(1, worksheet.max_row + 1):
            key = to_text(worksheet.cell(row=row, column=1).value).strip()
            value = to_text(worksheet.cell(row=row, column=2).value)
            if key in _META_MARKERS:
                collecting = True
                continue
            if collecting:
                if key == "" and value == "":
                    break
                if key:
                    found[key] = value
        if found:
            return found

    # 退化路径：只有 HOWTO 表（template_sheet 被关掉）时，从人读的那几行里抠出来。
    # 键名必须和 Template 表里的元信息块**完全一致**，否则调用方按 "参数指纹" 取值会取空。
    howto_name = config.excel.howto_sheet
    if howto_name and howto_name in workbook.sheetnames:
        worksheet = workbook[howto_name]
        found = {}
        for row in range(1, worksheet.max_row + 1):
            text = to_text(worksheet.cell(row=row, column=1).value).strip()
            for label in (META_TIME, META_INPUT, META_OUTPUT):
                if text.startswith(label):
                    found[label] = text[len(label) :].strip()
        if found:
            return found
    return {}


# --------------------------------------------------------------------------- #
# 读取工作簿
# --------------------------------------------------------------------------- #
def load_workbook_file(path: str | Path) -> Workbook:
    """打开 Excel 文件，失败时给出可读的提示。"""
    target = Path(path)
    if not target.exists():
        raise ExcelError(f"Excel 文件不存在: {target}（请先运行 `spreadsheet-codegen init` 生成模板）")
    if target.is_dir():
        raise ExcelError(f"路径是目录而不是 Excel 文件: {target}")
    try:
        return load_workbook(target)
    except Exception as exc:  # openpyxl 会抛各种异常类型
        raise WorkbookIOError(f"无法读取 Excel 文件 {target}: {exc}") from exc


def get_sheet(workbook: Workbook, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise ExcelError(f"缺少工作表 {name!r}。当前工作表: {', '.join(workbook.sheetnames)}")
    return workbook[name]


def check_required_sheets(workbook: Workbook, config: ProjectConfig) -> None:
    """校验 Global / Local / Output 工作表是否存在，一次报出所有缺失项。"""
    required = [
        config.excel.sheets.global_,
        config.excel.sheets.local,
        *config.excel.sheets.outputs,
    ]
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise ExcelError(
            "Excel 缺少工作表: "
            + ", ".join(repr(name) for name in missing)
            + f"。当前工作表: {', '.join(workbook.sheetnames)}"
            + "（可重新运行 `spreadsheet-codegen init --force` 生成模板）"
        )
    _check_header(workbook[config.excel.sheets.global_], GLOBAL_HEADERS, config.excel.sheets.global_)
    local_sheet = workbook[config.excel.sheets.local]
    if config.excel.local_direction == "horizontal":
        _check_header(local_sheet, LOCAL_HEADERS, config.excel.sheets.local, strict=False)
    else:
        # 纵向布局第 1 行整行都是变量名，只有 A1 是固定的
        _check_header(local_sheet, (LOCAL_CASE_HEADER,), config.excel.sheets.local, strict=False)


def _check_header(
    worksheet: Worksheet,
    expected: Sequence[str],
    sheet_name: str,
    *,
    strict: bool = True,
) -> None:
    for index, title in enumerate(expected, start=1):
        actual = to_text(worksheet.cell(row=1, column=index).value).strip()
        if actual and actual.lower() != title.lower():
            raise ExcelError(f"工作表 {sheet_name!r} 第 1 行第 {index} 列表头应为 {title!r}，实际是 {actual!r}")
        if not actual and strict:
            raise ExcelError(f"工作表 {sheet_name!r} 第 1 行第 {index} 列表头为空，应为 {title!r}")


def read_global_values(
    workbook: Workbook,
    config: ProjectConfig,
    *,
    warnings: list[str] | None = None,
) -> dict[str, VarValue]:
    """读取 Global Parameter 表：B 列值 + D/E 列前缀后缀；**派生参数在这里算出来**。

    回落规则：**只有单元格真的为空才回落 YAML**；非空值原样使用（首尾空格有意义），
    所以 ``suffix: " m"`` 会得到 ``" m"`` 而不是 ``"m"``。

    派生参数（``derived:``）的格子由工具写成公式 / 算好的值，这里**不读它**，
    而是用表达式现算（同 Case 的其他参数 + 全局参数）。
    """
    worksheet = get_sheet(workbook, config.excel.sheets.global_)
    definitions = {item.name: item for item in config.global_variables}
    values: dict[str, VarValue] = {}
    raw_values: dict[str, Any] = {}
    seen: set[str] = set()

    for row in range(2, worksheet.max_row + 1):
        name = to_text(worksheet.cell(row=row, column=_GLOBAL_COL["name"]).value).strip()
        if not name:
            continue
        if name in seen:
            raise ExcelError(f"工作表 {config.excel.sheets.global_!r} 第 {row} 行变量名 {name!r} 重复")
        seen.add(name)

        definition = definitions.get(name)
        prefix = _text_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["prefix"]).value,
            definition.prefix if definition else "",
        )
        suffix = _text_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["suffix"]).value,
            definition.suffix if definition else "",
        )

        if definition is not None and definition.is_derived:
            # 派生格：留个占位，等输入读完之后统一算
            values[name] = VarValue("", prefix, suffix)
            continue

        raw_value = _cell_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["value"]).value,
            definition.effective_default if definition else "",
        )
        kind = definition.type if definition else "auto"
        coerced = _coerce(raw_value, kind, name)
        raw_values[name] = coerced
        values[name] = VarValue(coerced, prefix, suffix)

    # YAML 中定义但表里没写的变量，用默认值补齐，保证模板引用不报错。
    for definition in config.global_variables:
        if definition.name in values:
            continue
        values[definition.name] = VarValue(definition.default, definition.prefix, definition.suffix)
        if not definition.is_derived:
            raw_values[definition.name] = _coerce(definition.default, definition.type, definition.name)

    _resolve_derived_globals(worksheet, config, values, raw_values, warnings)
    return values


def _derived_context(
    definitions: Sequence[VariableDef],
    values: Mapping[str, VarValue],
    raw_values: Mapping[str, Any],
) -> dict[str, VarValue]:
    """给派生表达式组装上下文：``x`` / ``x.value`` / ``x.prefix`` / ``x.suffix`` 都能写。

    为什么要包一层：派生表达式此前拿到的是**裸值**（一个 str / float），于是
    ``derived: "len(secret.value)"`` 会在求值期报 ``'str' object has no attribute
    'value'`` —— 配置期还查不出来。

    包的是 :class:`~spreadsheet_codegen.utils.DerivedValue` 而**不是**模板侧的
    ``VarValue``：派生表达式里裸名字必须代表**纯值**（指南 §15.1）—— 与 Excel 侧
    ``resolve(name)`` 指向取值列完全一致。用 ``VarValue`` 会让 ``x`` 变成带前后缀的
    组合串，于是 ``a ~ a`` Python 给 ``"XAZXAZ"``、Excel 给 ``"AA"``；
    ``len(a)`` Python 给 3、Excel 给 1。两边对同一个名字的理解必须一样。

    前缀 / 后缀一律取**表里读到的**（``values``）而不是 YAML 里的 —— Local 表允许在
    C / D 列覆盖前后缀，公式模式读的是表里的那一格；用 YAML 的值会让"参数表里看到的"
    与"导出的文件"悄悄不一致（真踩过：表里写 ``GPIOX``、派生表达式却按 ``GPIO`` 算长度）。
    """
    definitions_by_name = {item.name: item for item in definitions}
    context: dict[str, VarValue] = {}
    for name, raw in raw_values.items():
        definition = definitions_by_name.get(name)
        loaded = values.get(name)
        if loaded is not None:
            context[name] = DerivedValue(raw, loaded.prefix, loaded.suffix)
        elif definition is None:
            context[name] = DerivedValue(raw)
        else:
            context[name] = DerivedValue(raw, definition.prefix, definition.suffix)
    return context


def _unwrap_derived(context: dict[str, Any], definitions: Sequence[VariableDef]) -> None:
    """把 ``_derived_context`` 里算出来的**派生结果**写回裸值形态。

    其他调用方（指纹 / 模板上下文）要的是算好的值本身，不是 ``VarValue`` 包装；
    非派生变量本来就该是裸值，原样放回。

    注意算出来的结果**已经是裸值**（表达式求值返回的就是数值 / 文本），
    所以这里只在"确实被包成 VarValue"时才拆 —— 不能假定一定是包装对象。
    """
    derived_names = {item.name for item in definitions if item.is_derived}
    for name in [name for name in context if name in derived_names]:
        value = context[name]
        if isinstance(value, VarValue):
            context[name] = value.value


def _resolve_derived_globals(
    worksheet: Worksheet,
    config: ProjectConfig,
    values: dict[str, VarValue],
    raw_values: dict[str, Any],
    warnings: list[str] | None,
) -> None:
    """把 global 的派生参数算出来（按依赖顺序），并检查格子里有没有被手工改过。"""
    derived = [item for item in config.global_variables if item.is_derived]
    if not derived:
        return
    context = _derived_context(config.global_variables, values, raw_values)
    evaluate_derived(
        config.global_variables,
        context,
        scope="global 的派生参数",
        forbidden={
            item.name: "global 的派生参数不能引用 local 变量（那时还没有当前 Case）" for item in config.local_variables
        },
    )
    _unwrap_derived(context, config.global_variables)
    for definition in derived:
        raw_values[definition.name] = context[definition.name]
    _warn_hand_edited_derived_global(worksheet, derived, _GLOBAL_COL["value"], raw_values, warnings, label="Global 表")
    for definition in derived:
        current = values.get(definition.name, VarValue(""))
        computed = _coerce(raw_values[definition.name], definition.type, definition.name)
        values[definition.name] = VarValue(computed, current.prefix, current.suffix)


def _warn_hand_edited_derived_global(
    worksheet: Worksheet,
    derived: Sequence[VariableDef],
    value_column: int,
    raw_values: Mapping[str, Any],
    warnings: list[str] | None,
    *,
    label: str,
) -> None:
    """Global 表专用：变量永远在行上，派生格固定在 ``value_column`` 列。"""
    if warnings is None:
        return
    rows = {
        to_text(worksheet.cell(row=row, column=_GLOBAL_COL["name"]).value).strip(): row
        for row in range(2, worksheet.max_row + 1)
    }
    for definition in derived:
        row = rows.get(definition.name)
        if row is None:
            continue
        raw = worksheet.cell(row=row, column=value_column).value
        if isinstance(raw, str) and raw.startswith("="):
            continue  # 正常的公式格
        if is_translatable(definition):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 的格子被手工改成了 {raw!r}，"
                "工具会忽略它（该格由表达式算出来；重跑 --write-excel 会把它改回公式）"
            )
        elif definition.name in raw_values and to_text(raw_values[definition.name]) != to_text(raw):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 表里是 {raw!r}，当前算式应为 "
                f"{to_text(raw_values[definition.name])!r} —— 请重跑 --write-excel 刷新"
            )


def _warn_hand_edited_derived(
    worksheet: Worksheet,
    config: ProjectConfig,
    derived: Sequence[VariableDef],
    case: CaseData,
    raw_values: Mapping[str, Any],
    warnings: list[str] | None,
    *,
    label: str,
) -> None:
    """派生格是公式（或工具写的值）—— 如果用户手工改成了别的值，提醒他会被忽略。"""
    if warnings is None:
        return
    slots = {slot.name: slot for slot in _local_slots(worksheet, config)}
    for definition in derived:
        slot = slots.get(definition.name)
        if slot is None:
            continue
        raw = _case_cell(worksheet, config, slot, case).value
        if isinstance(raw, str) and raw.startswith("="):
            continue  # 正常的公式格
        if is_translatable(definition):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 的格子被手工改成了 {raw!r}，"
                "工具会忽略它（该格由表达式算出来；重跑 --write-excel 会把它改回公式）"
            )
        elif definition.name in raw_values and to_text(raw_values[definition.name]) != to_text(raw):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 表里是 {raw!r}，当前算式应为 "
                f"{to_text(raw_values[definition.name])!r} —— 请重跑 --write-excel 刷新"
            )


def read_cases(
    workbook: Workbook,
    config: ProjectConfig,
    *,
    global_values: Mapping[str, VarValue] | None = None,
    warnings: list[str] | None = None,
) -> list[CaseData]:
    """读取 Local Parameter 表里的工况；**派生参数按 Case 算出来**。

    两种布局由 ``excel.local_direction`` 决定：

    * ``horizontal``（默认）：一个工况**一列**（E1 起写 Case 名），变量在行上；
    * ``vertical``：一个工况**一行**（A2 起写 Case 名），变量在列上。

    两种布局都是"遇到空表头就停"，所以右拉 / 下拉就能加工况。
    """
    worksheet = get_sheet(workbook, config.excel.sheets.local)
    sheet_name = config.excel.sheets.local
    if global_values is None:
        global_values = read_global_values(workbook, config, warnings=warnings)
    global_raw = {name: value.value for name, value in global_values.items()}

    horizontal = config.excel.local_direction == "horizontal"
    anchors = _case_anchors(worksheet, config)
    if not anchors:
        hint = (
            "请在 E1 填写 Case1、F1 填写 Case2 …（可右拉复制列）"
            if horizontal
            else "请在 A2 填写 Case1、A3 填写 Case2 …（可下拉复制行）"
        )
        raise ExcelError(f"工作表 {sheet_name!r} 里没有工况。{hint}")

    slots = _local_slots(worksheet, config)
    if not slots:
        raise ExcelError(
            f"工作表 {sheet_name!r} 没有定义任何局部变量 —— "
            "本工具用「局部变量 × 工况」定位计算，所以 variables.local 至少要有一个变量"
            "（哪怕只是个标注用的 kind）"
        )

    cases: list[CaseData] = []
    for anchor in anchors:
        values: dict[str, VarValue] = {}
        raw_values: dict[str, Any] = dict(global_raw)  # 局部派生可以引用全局
        explicit = False
        for slot in slots:
            if slot.definition is not None and slot.definition.is_derived:
                values[slot.name] = VarValue("", slot.prefix, slot.suffix)
                continue
            cell = _case_cell(worksheet, config, slot, anchor)
            if cell.value is not None and not (isinstance(cell.value, str) and cell.value.strip() == ""):
                explicit = True
            raw_value = _cell_or(cell.value, slot.definition.effective_default if slot.definition else "")
            kind = slot.definition.type if slot.definition else "auto"
            coerced = _coerce(raw_value, kind, f"{anchor.name}.{slot.name}")
            raw_values[slot.name] = coerced
            values[slot.name] = VarValue(coerced, slot.prefix, slot.suffix)

        # YAML 中定义但表里缺少的局部变量，用默认值补齐（派生参数稍后算）
        for definition in config.local_variables:
            if definition.name in values:
                continue
            values[definition.name] = VarValue(definition.default, definition.prefix, definition.suffix)
            if not definition.is_derived:
                raw_values[definition.name] = _coerce(definition.default, definition.type, definition.name)

        _resolve_derived_locals(worksheet, config, anchor, values, raw_values, warnings, global_values)
        cases.append(
            CaseData(
                name=anchor.name,
                column=anchor.column,
                row=anchor.row,
                values=values,
                explicit_values=explicit,
            )
        )
    return cases


def _resolve_derived_locals(
    worksheet: Worksheet,
    config: ProjectConfig,
    case: CaseData,
    values: dict[str, VarValue],
    raw_values: dict[str, Any],
    warnings: list[str] | None,
    global_values: Mapping[str, VarValue] | None = None,
) -> None:
    derived = [item for item in config.local_variables if item.is_derived]
    if not derived:
        return
    # 局部派生可以引用全局参数：两边都进上下文。
    # 前后缀取**各自表里读到的**：local 的在 values 里、global 的在 global_values 里 ——
    # 只传 values 会让 global 名字回落到 YAML 的前后缀，而公式侧读的是 Global 表的
    # D/E 列，于是表里改过前缀就出现"Python 算一个数、Excel 算另一个数"。
    lookup = {**(global_values or {}), **values}
    definitions = [*config.global_variables, *config.local_variables]
    context = _derived_context(definitions, lookup, raw_values)
    evaluate_derived(
        config.local_variables,
        context,
        scope=f"Case {case.name!r} 的 local 派生参数",
    )
    _unwrap_derived(context, config.local_variables)
    for definition in derived:
        raw_values[definition.name] = context[definition.name]
    _warn_hand_edited_derived(
        worksheet,
        config,
        derived,
        case,
        raw_values,
        warnings,
        label=f"Local 表 Case {case.name!r}",
    )
    for definition in derived:
        current = values.get(definition.name, VarValue(""))
        computed = _coerce(raw_values[definition.name], definition.type, definition.name)
        values[definition.name] = VarValue(computed, current.prefix, current.suffix)


def read_group_members(
    workbook: Workbook,
    config: ProjectConfig,
    *,
    warnings: list[str] | None = None,
) -> dict[str, dict[str, VarValue]]:
    """读成员表（第三层作用域）：**一行一个成员，B 列起一个变量一列**。

    :returns: ``{成员名: {变量名: VarValue}}``；没配 ``variables.group`` 时返回空字典。
    """
    group = config.group
    if group is None:
        return {}
    sheet = group.sheet
    worksheet = get_sheet(workbook, sheet)

    # 表头：B 列起是变量名（遇到空表头就停）
    columns: dict[str, int] = {}
    for column in range(_GROUP_FIRST_VAR_COLUMN, worksheet.max_column + 1):
        name = to_text(worksheet.cell(row=1, column=column).value).strip()
        if not name:
            break
        if name in columns:
            raise ExcelError(f"成员表 {sheet!r} 第 1 行表头 {name!r} 重复")
        columns[name] = column
    defined = {variable.name: variable for variable in group.variables}
    missing = [name for name in defined if name not in columns]
    if missing:
        raise ExcelError(
            f"成员表 {sheet!r} 缺少这些变量的列: {', '.join(missing)}"
            "（表头要写变量名；改了 YAML 的 group.variables 之后要重跑 init / 手工补列）"
        )
    unknown = [name for name in columns if name not in defined]
    if unknown and warnings is not None:
        warnings.append(f"成员表 {sheet!r} 里有 YAML 未定义的列: {', '.join(unknown)}（会被读进上下文，按 auto 类型）")

    members: dict[str, dict[str, VarValue]] = {}
    for row in range(2, worksheet.max_row + 1):
        name = to_text(worksheet.cell(row=row, column=1).value).strip()
        if not name:
            continue
        if name in members:
            raise ExcelError(f"成员表 {sheet!r} 第 {row} 行成员名 {name!r} 重复")
        values: dict[str, VarValue] = {}
        for column_name, column in columns.items():
            definition = defined.get(column_name)
            raw = worksheet.cell(row=row, column=column).value
            value = _cell_or(raw, definition.effective_default if definition else "")
            kind = definition.type if definition else "auto"
            values[column_name] = VarValue(
                _coerce(value, kind, f"{name}.{column_name}"),
                definition.prefix if definition else "",
                definition.suffix if definition else "",
            )
        members[name] = values
    return members


@dataclass(frozen=True)
class _VarSlot:
    """Local 表里一个局部变量的位置。

    ``axis`` 的含义随布局而变：横向布局是**行号**（名字在 A 列），纵向布局是**列号**（名字在第 1 行）。
    """

    name: str
    definition: VariableDef | None
    prefix: str
    suffix: str
    axis: int


def _local_slots(worksheet: Worksheet, config: ProjectConfig) -> list[_VarSlot]:
    """扫出 Local 表里的变量槽。

    横向布局：A 列从第 2 行起写变量名，C/D 列可覆盖 Prefix/Suffix（与 YAML 一致的老行为）。
    纵向布局：第 1 行从 B 列起写变量名（A1 是 "Case"），**没有** Prefix/Suffix 列 ——
    它们只来自 YAML（表头格有批注写着）。
    """
    sheet_name = config.excel.sheets.local
    definitions = {item.name: item for item in config.local_variables}
    slots: list[_VarSlot] = []
    seen: set[str] = set()

    if config.excel.local_direction == "horizontal":
        candidates = [(row, _LOCAL_COL["name"]) for row in range(2, worksheet.max_row + 1)]
    else:
        candidates = [(1, column) for column in range(_GROUP_FIRST_VAR_COLUMN, worksheet.max_column + 1)]

    for row, column in candidates:
        name = to_text(worksheet.cell(row=row, column=column).value).strip()
        if not name:
            continue
        if name in seen:
            where = f"第 {row} 行" if config.excel.local_direction == "horizontal" else f"第 {column} 列"
            raise ExcelError(f"工作表 {sheet_name!r} {where}变量名 {name!r} 重复")
        seen.add(name)
        definition = definitions.get(name)
        if config.excel.local_direction == "horizontal":
            prefix = _text_or(
                worksheet.cell(row=row, column=_LOCAL_COL["prefix"]).value,
                definition.prefix if definition else "",
            )
            suffix = _text_or(
                worksheet.cell(row=row, column=_LOCAL_COL["suffix"]).value,
                definition.suffix if definition else "",
            )
            axis = row
        else:
            prefix = definition.prefix if definition else ""
            suffix = definition.suffix if definition else ""
            axis = column
        slots.append(_VarSlot(name, definition, prefix, suffix, axis))
    return slots


def _case_cell(worksheet: Worksheet, config: ProjectConfig, slot: _VarSlot, case: CaseData):
    """某个变量在某个 Case 上的那格（两种布局各取一个坐标）。"""
    if config.excel.local_direction == "horizontal":
        assert case.column is not None
        return worksheet.cell(row=slot.axis, column=case.column)
    assert case.row is not None
    return worksheet.cell(row=case.row, column=slot.axis)


def _case_anchors(worksheet: Worksheet, config: ProjectConfig) -> list[CaseData]:
    """Local 表里的工况位置（只读名字与位置，不读取值）。

    * ``horizontal``（默认）：一个工况一列，名字在**第 1 行**从 E 列起；
    * ``vertical``：一个工况一行，名字在 **A 列**从第 2 行起。

    两种布局都是"遇到空表头就停"，所以右拉 / 下拉加一列 / 一行即可增工况。
    """
    sheet_name = config.excel.sheets.local
    anchors: list[CaseData] = []
    seen: set[str] = set()

    def take(name: str, *, column: int | None, row: int | None, position: str) -> None:
        if name in seen:
            raise ExcelError(f"工作表 {sheet_name!r} 的 Case 名重复: {name!r}")
        seen.add(name)
        anchors.append(CaseData(name=name, column=column, row=row, values={}))
        del position

    if config.excel.local_direction == "horizontal":
        upper = max(worksheet.max_column, FIRST_CASE_COLUMN)
        for column in range(FIRST_CASE_COLUMN, upper + 1):
            name = to_text(worksheet.cell(row=1, column=column).value).strip()
            if not name:
                if anchors:  # 遇到空列说明 Case 列已经结束
                    break
                continue
            take(name, column=column, row=None, position=f"第 {column} 列")
    else:
        for row in range(2, worksheet.max_row + 1):
            name = to_text(worksheet.cell(row=row, column=1).value).strip()
            if not name:
                if anchors:  # 遇到空行说明 Case 行已经结束
                    break
                continue
            take(name, column=None, row=row, position=f"第 {row} 行")
    return anchors


def _coerce(value: Any, kind: str, label: str) -> Any:
    if value is None or kind in ("auto", "raw"):
        return value
    if kind == "string":
        return to_text(value)
    text = value.strip() if isinstance(value, str) else value
    if kind == "int":
        try:
            return int(float(text))
        except (TypeError, ValueError):
            raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 int") from None
    if kind == "float":
        try:
            return float(text)
        except (TypeError, ValueError):
            raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 float") from None
    if kind == "bool":
        if isinstance(value, bool):
            return value
        text_value = to_text(value).strip().lower()
        if text_value in {"1", "true", "yes", "y", "on"}:
            return True
        if text_value in {"0", "false", "no", "n", "off", ""}:
            return False
        raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 bool")
    return value


def _warn_choices_with_blank(variables: Sequence[VariableDef], warnings: list[str] | None) -> None:
    """``choices`` 里塞空串是"允许留空"的**老绕法**，现在有 ``allow_blank`` 了。

    只提示、不改行为：老配置照常能跑；但把"为什么这么写、更好的写法是什么"说出来，
    免得这个反直觉的写法继续被当成唯一出路（它没有任何文档提示）。
    """
    if warnings is None:
        return
    for variable in variables:
        if not variable.choices:
            continue
        if any(to_text(item) == "" for item in variable.choices) and not variable.allow_blank:
            warnings.append(
                f"变量 {variable.name!r} 的 choices 里有空串（{variable.choices!r}）—— "
                "那是「允许留空」的老写法。建议改成 choices 只写真实取值 + allow_blank: true："
                "下拉列表里就不会再出现一个空选项，「可以留空」也写得明明白白"
            )


def check_value_constraints(
    config: ProjectConfig,
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
    members: Mapping[str, Mapping[str, VarValue]] | None = None,
    *,
    warnings: list[str] | None = None,
) -> None:
    """校验表里填的取值是否满足变量声明的 ``min`` / ``max`` / ``choices`` / ``pattern``。

    为什么要它：工具此前只查"变量有没有定义、类型对不对"，**完全不看值** —— 把
    ``20.559`` 手滑打成 ``205.59`` 会一路渲染成错误代码，而 ``check`` 只会说"与参数一致"。
    声明了约束就一定查，**空值默认也算不合格**（``choices`` 意味着"必须给一个合法取值"）；
    想让某个变量能留空就写 ``allow_blank: true``。

    不满足时抛 :class:`ExcelError`，一次列全部问题并指出是哪张表、哪一列、哪个变量。

    :param warnings: 传一个列表进来，会把"``choices`` 里塞了空串"这类"能跑但该改"的写法写进去。
    """
    problems: list[str] = []
    global_sheet = config.excel.sheets.global_

    _warn_choices_with_blank([*config.global_variables, *config.local_variables, *config.group_variables], warnings)

    for row, variable in enumerate(config.global_variables, start=2):
        if variable.is_derived or not variable.has_constraints:
            continue
        value = global_values.get(variable.name)
        if value is None:
            continue
        problem = variable.value_problem(value.text)
        if problem:
            problems.append(f"  {global_sheet} 第 {row} 行 '{variable.name}'：{problem}")

    local_sheet = config.excel.sheets.local
    for case in cases:
        for variable in config.local_variables:
            if variable.is_derived or not variable.has_constraints:
                continue
            value = case.values.get(variable.name)
            if value is None:
                continue
            problem = variable.value_problem(value.text)
            if problem:
                problems.append(f"  {local_sheet} {case.where} '{case.name}' 的 '{variable.name}'：{problem}")

    # 成员表（第三层作用域）：一个成员一行
    if members:
        for member, values in members.items():
            for variable in config.group_variables:
                if not variable.has_constraints:
                    continue
                value = values.get(variable.name)
                if value is None:
                    continue
                problem = variable.value_problem(value.text)
                if problem:
                    problems.append(f"  {config.group.sheet} 成员 '{member}' 的 '{variable.name}'：{problem}")

    if problems:
        raise InputError(
            f"参数取值不满足变量声明的约束，共 {len(problems)} 处：\n"
            + "\n".join(problems)
            + "\n  → 改 Excel 里的取值，或放宽 YAML 里的 min / max / choices / pattern"
        )


# --------------------------------------------------------------------------- #
# 写回渲染结果
# --------------------------------------------------------------------------- #
def input_fingerprint(
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
) -> str:
    """参数指纹：Global 取值 + 每个 Case 的取值（前缀/后缀也算参数）。"""
    parts: list[str] = []
    for name, value in global_values.items():
        parts.append(f"G|{name}|{value.prefix}|{value.text}|{value.suffix}")
    for case in cases:
        for name, value in case.values.items():
            parts.append(f"L|{case.name}|{name}|{value.prefix}|{value.text}|{value.suffix}")
    return fingerprint(*parts)


def output_fingerprint(rendered: Mapping[str, Sequence[RenderResult]]) -> str:
    """输出指纹：所有模板 × Case 的渲染文本。"""
    parts: list[str] = []
    for template_name, per_case in rendered.items():
        for result in per_case:
            parts.append(f"{template_name}/{result.case_name}\n{result.text}")
    return fingerprint(*parts)


@dataclass(frozen=True)
class _Block:
    """写进输出表的一块内容：一个 Case 名 + 若干"行"。"""

    case_name: str
    lines: list[str]

    @property
    def line_count(self) -> int:
        return len(self.lines)


def _template_source(template, base_dir: str | Path | None) -> str:
    """读取模板源码（公式模式需要它；``template_file`` 相对**声明它的那个文件**所在目录解析）。"""
    if not template.template_file:
        return template.source_code
    path = Path(template.template_file)
    base = template.source_dir or base_dir
    if base and not path.is_absolute():
        path = Path(base) / path
    if not path.exists():
        raise ExcelError(f"模板 {template.name!r} 引用的模板文件不存在: {path}")
    try:
        return path.read_text(encoding="utf-8")
    except OSError as exc:
        raise ExcelError(f"无法读取模板文件 {path}: {exc}") from exc


#: 公开别名（CLI 的 check/validate 也要用同一份解析逻辑）
template_source = _template_source


def case_anchor_map(workbook: Workbook, config: ProjectConfig) -> dict[str, CaseData]:
    """Local 表的 ``Case 名 -> 位置`` 映射（公式模式要知道每个 Case 在哪一列 / 哪一行）。"""
    worksheet = get_sheet(workbook, config.excel.sheets.local)
    return {anchor.name: anchor for anchor in _case_anchors(worksheet, config)}


def _case_axis(anchor: CaseData, config: ProjectConfig) -> int:
    """Case 的"工况轴"：横向布局是列号，纵向布局是行号。"""
    if config.excel.local_direction == "horizontal":
        assert anchor.column is not None
        return anchor.column
    assert anchor.row is not None
    return anchor.row


def case_axis_map(workbook: Workbook, config: ProjectConfig) -> dict[str, int]:
    return {name: _case_axis(anchor, config) for name, anchor in case_anchor_map(workbook, config).items()}


def _formula_blocks(
    workbook: Workbook,
    config: ProjectConfig,
    template,
    results: Sequence[RenderResult],
    case_axes: dict[str, int],
) -> list[_Block]:
    """把模板编译成"每行一个公式"的块。"""
    missing = [result.case_name for result in results if result.case_name not in case_axes]
    if missing:
        raise ExcelError(
            f"模板 {template.name!r} 使用公式模式，但 Local 表里找不到这些 Case 列: "
            f"{', '.join(missing)}（可用: {', '.join(case_axes) or '（无）'}）"
        )
    source = _template_source(template, config.source_dir)
    per_case = compile_formulas(
        template,
        config,
        case_axes=[case_axes[result.case_name] for result in results],
        source=source,
    )
    return [_Block(case_name=result.case_name, lines=lines) for result, lines in zip(results, per_case, strict=False)]


def _row_of(worksheet: Worksheet, variable: str) -> int | None:
    """按变量名在 A 列找行号。"""
    for row in range(2, worksheet.max_row + 1):
        if to_text(worksheet.cell(row=row, column=1).value).strip() == variable:
            return row
    return None


def refresh_derived_cells(
    workbook: Workbook,
    config: ProjectConfig,
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
    *,
    warnings: list[str] | None = None,
) -> bool:
    """刷新参数表里的派生参数格：能写公式就写公式，否则写入 Python 算好的值。

    :returns: 是否写过 Excel 公式（调用方据此决定要不要设 ``fullCalcOnLoad``）
    """
    if not any(item.is_derived for item in [*config.global_variables, *config.local_variables]):
        return False

    wrote_formula = False
    global_sheet = get_sheet(workbook, config.excel.sheets.global_)
    global_resolve = _global_resolver(config)
    for definition in config.global_variables:
        if not definition.is_derived:
            continue
        row = _row_of(global_sheet, definition.name)
        if row is None:
            continue
        value = global_values.get(definition.name)
        _write_derived_cell(
            global_sheet.cell(row=row, column=_GLOBAL_COL["value"]),
            definition,
            resolve=global_resolve,
            value=value.value if value is not None else None,
        )
        wrote_formula = wrote_formula or is_translatable(definition)

    local_sheet = get_sheet(workbook, config.excel.sheets.local)
    slots = {slot.name: slot for slot in _local_slots(local_sheet, config)}
    for case in cases:
        resolve = _local_resolver(config, case)
        for definition in config.local_variables:
            if not definition.is_derived:
                continue
            slot = slots.get(definition.name)
            if slot is None:
                continue
            value = case.values.get(definition.name)
            _write_derived_cell(
                _case_cell(local_sheet, config, slot, case),
                definition,
                resolve=resolve,
                value=value.value if value is not None else None,
            )
            wrote_formula = wrote_formula or is_translatable(definition)
    return wrote_formula


def write_results(
    path: str | Path,
    config: ProjectConfig,
    rendered: Mapping[str, Sequence[RenderResult]],
    *,
    update_howto: bool = True,
    command: str | None = None,
    warnings: list[str] | None = None,
) -> Path:
    """把渲染结果写入各模板对应的 Output 表，并在 HOWTO / Template 表里记录指纹。

    * ``direction: horizontal`` —— 每个 Case 一列（结果行向下展开）
    * ``direction: vertical``   —— 每个 Case 一行（结果行向右展开）
    * ``engine: excel``         —— 写 **Excel 公式**（默认）：改参数后由 Excel 自己重算，不用重跑脚本
    * ``engine: snapshot``      —— 写**文本快照**：改参数必须重跑命令（要用循环 / 过滤器时才选它）

    写入前会清理旧的输出区域（按**表内原有的真实边界**算，而不是按本次行数），
    因此"改短模板 / 减少 Case 之后重渲染"不会残留上一次的内容。

    :param warnings: 传一个列表进来，会把"公式过长"这类不致命的问题写进去。
    """
    target = Path(path)
    workbook = load_workbook_file(target)
    try:
        case_axes: dict[str, int] | None = None
        formula_written = False
        for template in config.templates:
            results = list(rendered.get(template.name, ()))
            if not results:
                continue
            if template.output_sheet not in workbook.sheetnames:
                workbook.create_sheet(template.output_sheet)
            worksheet = workbook[template.output_sheet]
            column, row = parse_cell(template.start_cell)

            if template.engine == "excel":
                if case_axes is None:
                    case_axes = case_axis_map(workbook, config)
                blocks = _formula_blocks(workbook, config, template, results, case_axes)
                formula_written = True
                if warnings is not None:
                    source = _template_source(template, config.source_dir)
                    longest, placeholders = longest_formula([block.lines for block in blocks], source)
                    if longest > LONG_FORMULA_WARN:
                        warnings.append(long_formula_warning(template.name, longest, placeholders=placeholders))
            else:
                blocks = [_Block(result.case_name, result.lines) for result in results]

            if template.direction == "horizontal":
                _write_horizontal(worksheet, template, blocks, column, row)
            else:
                _write_vertical(worksheet, template, blocks, column, row)

        if formula_written:
            # 让 Excel / WPS 打开文件时立刻重算，而不是显示上一次的缓存值。
            # 公式模式下必须重算（openpyxl 默认恰好也是 True，这里显式写死，不依赖上游默认）
            with suppress(AttributeError):  # pragma: no cover - 老版本 openpyxl 兜底
                workbook.calculation.fullCalcOnLoad = True

        metadata = {
            META_TIME: datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            META_OUTPUT: output_fingerprint(rendered),
        }
        values: dict[str, VarValue] | None = None
        cases: list[CaseData] | None = None
        try:
            values = read_global_values(workbook, config, warnings=warnings)
            cases = read_cases(workbook, config, global_values=values, warnings=warnings)
            metadata[META_INPUT] = input_fingerprint(values, cases)
            metadata["case 数"] = str(len(cases))
            metadata["cases"] = ", ".join(case.name for case in cases) or "（无）"
        except CodeGenError:
            metadata[META_INPUT] = "（读取参数失败）"
            metadata["cases"] = ", ".join(_case_names(rendered)) or "（无）"

        # 派生参数的格子也刷新一遍：能写公式就写公式（改输入自动重算）
        if (
            values is not None
            and cases is not None
            and refresh_derived_cells(workbook, config, values, cases, warnings=warnings)
        ):
            formula_written = True

        # 脚本是否还在工作簿旁边要现查：config 可能是刚从 YAML 重新加载的，
        # 那个标记会回落成默认的 False（HOWTO 表据此决定要不要提「双击」）
        config.scripts_enabled = (target.parent / f"{target.stem}_render.sh").exists()
        _record_metadata(workbook, config, metadata, command=command, update_howto=update_howto)
        workbook.save(target)
    except OSError as exc:
        raise WorkbookIOError(f"无法写回 Excel {target}: {exc}（文件被 Excel 占用？）") from exc
    finally:
        workbook.close()
    return target


def _record_metadata(
    workbook: Workbook,
    config: ProjectConfig,
    metadata: Mapping[str, str],
    *,
    command: str | None,
    update_howto: bool,
) -> None:
    sheet_name = config.excel.template_sheet
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(sheet_name)
            _build_template_sheet(worksheet, config)
            worksheet.sheet_state = "hidden"
        _write_meta_block(workbook[sheet_name], metadata)
    if update_howto and config.excel.howto_sheet:
        howto_name = config.excel.howto_sheet
        if howto_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(howto_name)
            workbook.move_sheet(worksheet, offset=-len(workbook.sheetnames) + 1)
        write_howto_sheet(workbook[howto_name], config, metadata=metadata, command=command)


def _case_names(rendered: Mapping[str, Sequence[RenderResult]]) -> list[str]:
    ordered: list[str] = []
    for per_case in rendered.values():
        for result in per_case:
            if result.case_name not in ordered:
                ordered.append(result.case_name)
    return ordered


def _write_horizontal(worksheet, template, blocks: Sequence[_Block], column: int, row: int) -> None:
    max_lines = max(block.line_count for block in blocks)
    header_row = row - 1 if template.write_case_headers and row > 1 else None
    top = header_row if header_row else row
    last_column = _last_used_column(worksheet, top, worksheet.max_row, column)
    last_column = max(last_column, column + len(blocks) - 1)
    # 底边取"表里原有的真实底边"，而不是本次行数 —— 否则改短模板会留下旧行
    last_row = _last_used_row(worksheet, top, column, last_column)
    bottom = max(row + max_lines + 1, last_row + 1)
    _clear_region(worksheet, top, bottom, column, last_column)

    for offset, block in enumerate(blocks):
        target_column = column + offset
        if header_row:
            cell = worksheet.cell(row=header_row, column=target_column, value=block.case_name)
            cell.font = _HEADER_FONT
            cell.fill = _CASE_HEADER_FILL
        for line_offset, line in enumerate(block.lines):
            worksheet.cell(row=row + line_offset, column=target_column, value=line)
        width = max((len(line) for line in block.lines), default=0)
        worksheet.column_dimensions[column_index_to_letter(target_column)].width = max(
            14, min(160, max(width, len(block.case_name)) + 2)
        )


def _write_vertical(worksheet, template, blocks: Sequence[_Block], column: int, row: int) -> None:
    max_columns = max(block.line_count for block in blocks)
    header_column = column - 1 if template.write_case_headers and column > 1 else None
    left = header_column if header_column else column
    right = max(column + max_columns - 1, column)
    last_column = _last_used_column(worksheet, row, worksheet.max_row, left)
    last_column = max(last_column, right)
    last_row = _last_used_row(worksheet, row, left, last_column)
    bottom = max(row + len(blocks) + 1, last_row + 1)
    _clear_region(worksheet, row, bottom, left, last_column)

    for offset, block in enumerate(blocks):
        target_row = row + offset
        if header_column:
            cell = worksheet.cell(row=target_row, column=header_column, value=block.case_name)
            cell.font = _HEADER_FONT
            cell.fill = _CASE_HEADER_FILL
        for line_offset, line in enumerate(block.lines):
            target_column = column + line_offset
            worksheet.cell(row=target_row, column=target_column, value=line)
            width = max(14, min(160, len(line) + 2))
            current = worksheet.column_dimensions[column_index_to_letter(target_column)].width or 0
            if width > current:
                worksheet.column_dimensions[column_index_to_letter(target_column)].width = width


def _last_used_column(worksheet: Worksheet, top: int, bottom: int, left: int) -> int:
    """在给定行区间内找到 left 右侧最后一个有内容的列（用于清理旧结果）。"""
    upper = max(worksheet.max_column, left)
    upper = min(upper, left + _SCAN_LIMIT)
    last = left - 1
    for row in worksheet.iter_rows(min_row=top, max_row=max(bottom, top), min_col=left, max_col=upper):
        for cell in row:
            if cell.value not in (None, ""):
                last = max(last, cell.column)
    return last


def _last_used_row(worksheet: Worksheet, top: int, left: int, right: int) -> int:
    """在给定列区间内找到 top 之下最后一个有内容的行号（用于清理上一次更长的结果）。"""
    lower = max(worksheet.max_row, top)
    lower = min(lower, top + _SCAN_LIMIT)
    last = top - 1
    for row in worksheet.iter_rows(min_row=top, max_row=lower, min_col=left, max_col=max(right, left)):
        for cell in row:
            if cell.value not in (None, ""):
                last = max(last, cell.row)
    return last


def _clear_region(worksheet: Worksheet, top: int, bottom: int, left: int, right: int) -> None:
    if bottom < top or right < left:
        return
    for row in worksheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
        for cell in row:
            if cell.value is not None:
                cell.value = None
