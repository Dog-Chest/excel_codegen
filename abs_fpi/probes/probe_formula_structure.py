"""探针 N：公式模式下改*结构*会怎样。

背景：公式模式下"改**参数**免重跑"是文档写明的。但结构变化是另一回事：
公式是写死在格子里的**文本**（`INDEX/MATCH` + 相对列），所以

* 插/删**变量行** → 应该没事（`MATCH` 按变量名找行）；
* 增/删/移动 **Case 列** → 公式里的列标就不再指着原来的工况了。

本探针用两份副本各做一次，并看两件事：
  1. `verify_excel_engine.py` 的求值比对是否失败（说明"Excel 里算出来的代码"已经和参数不一致）；
  2. `excel-codegen check` 是否能发现（这决定它能不能放进 CI 当护栏）。

    python probe_formula_structure.py
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent                      # abs_fpi/
PROJECT = ROOT / "abs_fpi.yaml"
BOOK = ROOT / "ABS_FPI_load_cases.xlsx"
PY = Path(sys.executable)


def run(args: list[str]) -> tuple[int, str]:
    proc = subprocess.run(args, capture_output=True, text=True, encoding="utf-8",
                          errors="replace", cwd=str(ROOT))
    return proc.returncode, (proc.stdout or "") + (proc.stderr or "")


def verify(book: Path) -> tuple[int, str]:
    return run([str(PY), str(ROOT / "verify_excel_engine.py"),
                str(PROJECT), str(book)])


def check(book: Path) -> tuple[int, str]:
    return run([str(PY), "-m", "excel_codegen", "check",
                "-c", str(PROJECT), "-x", str(book)])


def banner(text: str) -> None:
    print("\n" + "=" * 74)
    print(text)
    print("=" * 74)


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="probe_struct_"))

    # ---- N1：插一行变量 -------------------------------------------------- #
    banner("N1  在 Local 表中间插一行变量（不改任何 Case 列）")
    n1 = tmp / "n1.xlsx"
    shutil.copy(BOOK, n1)
    wb = load_workbook(n1)
    ws = wb["Local Parameter"]
    ws.insert_rows(6)                       # 在变量行之间插一行
    ws.cell(row=6, column=1, value="probe_inserted")
    ws.cell(row=6, column=2, value="探针插入的变量行（无取值 → 回落默认值）")
    wb.save(n1)
    wb.close()
    code, out = verify(n1)
    print(f"  verify_excel_engine.py → exit {code}   "
          f"{'公式仍然按变量名找到了新位置 ✓' if code == 0 else '失败（见下）'}")
    if code:
        print("   " + "\n   ".join(out.strip().splitlines()[-6:]))
    code2, out2 = check(n1)
    print(f"  excel-codegen check    → exit {code2}   "
          f"{'（插变量行不影响公式文本，所以 check 也不该报过期）' if code2 == 0 else '报过期'}")
    if code2:
        print("   " + "\n   ".join(out2.strip().splitlines()[-6:]))

    # ---- N2：改一个 Case 列的位置 ---------------------------------------- #
    banner("N2  在 Local 表里插一个 Case 列（工况整体右移一列）")
    n2 = tmp / "n2.xlsx"
    shutil.copy(BOOK, n2)
    wb = load_workbook(n2)
    ws = wb["Local Parameter"]
    ws.insert_cols(6)                       # 在第一个 Case 列之后插一列
    ws.cell(row=1, column=6, value="probe_inserted_case")
    wb.save(n2)
    wb.close()
    code, out = verify(n2)
    print(f"  verify_excel_engine.py → exit {code}   "
          f"{'（居然还对？）' if code == 0 else '公式已过期 ✓（求值出来的代码与参数不一致）'}")
    for line in out.strip().splitlines()[-8:]:
        print("   " + line)
    code2, out2 = check(n2)
    print(f"  excel-codegen check    → exit {code2}   "
          f"{'没发现' if code2 == 0 else '发现过期 ✓（可放进 CI）'}")
    for line in out2.strip().splitlines():
        if "ERROR" in line or "不同" in line or "参数指纹" in line or "输出指纹" in line:
            print("   " + line.strip())

    shutil.rmtree(tmp, ignore_errors=True)
    print("\n完成。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
