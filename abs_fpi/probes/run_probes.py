"""跑全部探针，把结果打出来。

    python run_probes.py

这些探针最初是给 0.1.0 写的"最小复现"（当时不改 excel_codegen 的任何文件）。
修复会话（0.2.0）之后**原样重跑**，用来对照每一条的行为变化；
探针 J / K 是修复新增的：分别验证 case_filter 与 render/check 的写回判定。
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

from openpyxl import load_workbook  # noqa: E402

from excel_codegen import (  # noqa: E402
    create_template,
    export_files,
    load_config,
    render_all,
    write_results,
)


def banner(text: str) -> None:
    print("\n" + "=" * 78)
    print(text)
    print("=" * 78)


def fresh(yaml_name: str, cases: int = 1):
    """按 YAML 生成一个干净的 Excel 并返回 (config, path)。"""
    cfg = load_config(HERE / yaml_name)
    path = HERE / cfg.excel.output
    if path.exists():
        path.unlink()
    create_template(cfg, path, cases=cases, overwrite=True, include_template_sheet=False)
    return cfg, path


def show_outdir(cfg, path, tag: str) -> None:
    out = HERE / "_out" / tag
    if out.exists():
        shutil.rmtree(out)
    files = export_files(cfg, render_all(cfg, path).results, out)
    for f in files:
        print(f"--- {f.name}")
        print(f.read_text(encoding="utf-8").rstrip("\n"))


# --------------------------------------------------------------------------- #
def probe_suffix() -> None:
    banner("探针 A  prefix / suffix 的首尾空格   (FINDINGS #2 · probe_suffix.yaml)")
    print('YAML 写的是 suffix: " m" / " kg/m^3"（带前导空格）')
    cfg, path = fresh("probe_suffix.yaml")
    print("YAML 里的 suffix :", [(d.name, d.suffix) for d in cfg.global_variables])
    ws = load_workbook(path)["Global Parameter"]
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=1).value
        print(f"  Excel E 列 {name:9s} = {ws.cell(row=row, column=5).value!r}")
    out = render_all(cfg, path)
    for name in ("L", "rho_sea"):
        v = out.global_values[name]
        print(f"  上下文     {name:9s} prefix={v.prefix!r} suffix={v.suffix!r} -> str={str(v)!r}")
    print()
    show_outdir(cfg, path, "suffix")


def probe_value() -> None:
    banner("探针 B  数值规范化：{{ x }} vs {{ x.value }}   (FINDINGS #4 · probe_value.yaml)")
    cfg, path = fresh("probe_value.yaml")
    out = render_all(cfg, path)
    for name in ("L", "g", "Li", "La"):
        v = out.global_values[name]
        print(
            f"  {name:3s} type={[d.type for d in cfg.global_variables if d.name == name][0]:5s} "
            f"value={v.value!r}  str(v)={str(v)!r}  repr(v.value)={v.value!r}"
        )
    v = out.cases[0].values["draft"]
    print(f"  draft type=float value={v.value!r}  str(v)={str(v)!r}")
    print()
    show_outdir(cfg, path, "value")


def probe_twosets() -> None:
    banner("探针 C  一个工作簿两套规则：模板 × Case 是交叉积   (FINDINGS #3b · probe_twosets.yaml)")
    cfg, path = fresh("probe_twosets.yaml", cases=2)
    # Case1 -> EXT, Case2 -> INT
    ws = load_workbook(path)["Local Parameter"]
    row_of = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    ws.cell(row=1, column=5, value="case_ext")
    ws.cell(row=1, column=6, value="case_int")
    ws.cell(row=row_of["kind"], column=5, value="EXT")
    ws.cell(row=row_of["kind"], column=6, value="INT")
    ws.cell(row=row_of["draft"], column=5, value=8)
    ws.cell(row=row_of["draft"], column=6, value=15)
    path_save = path
    load_workbook(path).save(path_save) if False else ws.parent.save(path_save)

    out = render_all(cfg, path)
    print("  期望：ext_code 只出 case_ext，int_code 只出 case_int（各 1 个）")
    print(f"  实际：{out.total()} 个渲染结果")
    for tname, per in out.results.items():
        print(f"    {tname:9s} -> {[r.case_name for r in per]}")
    write_results(path, cfg, out.results)
    wb = load_workbook(path)
    for sheet in ("Code EXT", "Code INT"):
        sh = wb[sheet]
        head = [sh.cell(row=1, column=c).value for c in range(2, 6)]
        print(f"  {sheet:9s} 表头行 = {[h for h in head if h]}")
    wb.close()
    print("  -> 不用 case_filter 时仍是交叉积（4 个结果）；加一行 case_filter 即可只出该模板适用的 Case（见探针 J）。")


def probe_casefilter() -> None:
    banner("探针 J  用 case_filter 把两套规则放进一本工作簿   (FINDINGS #3b 修复验证 · probe_casefilter.yaml)")
    cfg, path = fresh("probe_casefilter.yaml", cases=2)
    ws = load_workbook(path)["Local Parameter"]
    row_of = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    ws.cell(row=1, column=5, value="case_ext")
    ws.cell(row=1, column=6, value="case_int")
    ws.cell(row=row_of["kind"], column=5, value="EXT")
    ws.cell(row=row_of["kind"], column=6, value="INT")
    ws.parent.save(path)

    out = render_all(cfg, path)
    print("  期望：ext_code 只出 case_ext，int_code 只出 case_int（各 1 个）")
    print(f"  实际：{out.total()} 个渲染结果")
    for tname, per in out.results.items():
        print(f"    {tname:9s} -> {[r.case_name for r in per]}")
    print(f"  被 case_filter 跳过：{out.skipped}")
    write_results(path, cfg, out.results)
    wb = load_workbook(path)
    for sheet in ("Code EXT", "Code INT"):
        sh = wb[sheet]
        head = [sh.cell(row=1, column=c).value for c in range(2, 6)]
        print(f"  {sheet:9s} 表头行 = {[h for h in head if h]}")
    wb.close()
    print("  -> Global 表只填一次（L），两套规则各写各的列。")


def probe_formula() -> None:
    banner("探针 L  公式模式：输出表写公式，改参数 Excel 自己重算   (probe_formula.yaml)")
    cfg, path = fresh("probe_formula.yaml", cases=2)
    out = render_all(cfg, path)
    write_results(path, cfg, out.results)

    wb = load_workbook(path)
    print(f"  fullCalcOnLoad = {wb.calculation.fullCalcOnLoad}（Excel/WPS 打开即重算）")
    ws = wb["Output"]
    for row in range(2, 7):
        formula = ws.cell(row=row, column=2).value
        print(f"  Output!B{row}: {str(formula)[:96]}{'…' if len(str(formula)) > 96 else ''}")
    wb.close()
    print("  -> 单元格里是公式，不是文本：openpyxl 不会算它（值只在 Excel 里存在）")

    print("\n  Python 侧渲染（用于导出文件 / 指纹；公式应当产出相同文本）：")
    for result in out.results["probe_formula"]:
        print(f"    [{result.case_name}] {result.lines[1]}")

    print("\n  改参数不需要重跑脚本 —— 直接把 Local 表 E2 改成 12，Excel 打开就是：")
    ws = load_workbook(path)["Local Parameter"]
    row_of = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    ws.cell(row=row_of["draft"], column=5, value=12)
    ws.parent.save(path)
    out2 = render_all(cfg, path)
    print(f"    [Case1] {out2.results['probe_formula'][0].lines[2]}   ← 公式会自动变成这一行")
    print("  （在 Excel 里这一步是「打开文件」就完成的；脚本只在改模板时才需要重跑）")


def probe_stale() -> None:
    banner("探针 D  长渲染之后用短模板重渲染   (FINDINGS #5 症状 A · probe_long/short.yaml)")
    long_cfg, path = fresh("probe_long.yaml")
    render_all(long_cfg, path)
    write_results(path, long_cfg, render_all(long_cfg, path).results)
    ws = load_workbook(path)["Output"]
    print("  长渲染之后 B 列：", [ws.cell(row=r, column=2).value for r in range(1, 11)])

    short_cfg = load_config(HERE / "probe_short.yaml")
    render_all(short_cfg, path)
    write_results(path, short_cfg, render_all(short_cfg, path).results)
    ws = load_workbook(path)["Output"]
    col = [ws.cell(row=r, column=2).value for r in range(1, 11)]
    print("  短渲染之后 B 列：", col)
    leftover = [v for v in col if v and "LONG" in str(v)]
    print(f"  残留的 LONG 行：{leftover if leftover else '（无）'}")


def probe_dupkey() -> None:
    banner("探针 E  YAML 同一个键写两遍   (FINDINGS #6 · probe_dupkey.yaml)")
    try:
        cfg = load_config(HERE / "probe_dupkey.yaml")
        print(
            f"  load_config 没有报错；global 变量数 = {len(cfg.global_variables)}，"
            f"local 变量数 = {len(cfg.local_variables)}"
        )
        print(f"  global 名字 = {cfg.global_names}")
        path = HERE / cfg.excel.output
        if path.exists():
            path.unlink()
        create_template(cfg, path, cases=1, overwrite=True, include_template_sheet=False)
        ws = load_workbook(path)["Global Parameter"]
        rows = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        print(f"  生成的 Global 表数据行 = {rows or '（空）'}")
        try:
            render_all(cfg, path)
            print("  渲染成功（说明 L 从 Excel 表里读到了）")
        except Exception as exc:  # noqa: BLE001
            print(f"  直到渲染才炸：{type(exc).__name__}: {exc}")
    except Exception as exc:  # noqa: BLE001
        print(f"  配置阶段就报错：{type(exc).__name__}: {exc}")


def probe_gbk() -> None:
    banner("探针 F  中文 Windows（GBK 控制台）下 CLI 的收尾打印   (FINDINGS #1)")
    print(f"  当前 sys.stdout.encoding = {sys.stdout.encoding}")
    import subprocess

    env = dict(os.environ, PYTHONIOENCODING="gbk")
    env.pop("PYTHONUTF8", None)
    proc = subprocess.run(
        [sys.executable, "-m", "excel_codegen", "validate", "-c", "probe_suffix.yaml"],
        cwd=str(HERE),
        capture_output=True,
        env=env,
        check=False,
    )
    tail = (proc.stdout or b"").decode("gbk", "replace").strip().splitlines()
    print("  用 PYTHONIOENCODING=gbk 真跑一次 validate：")
    print("    " + "\n    ".join(tail[-2:]))
    print(f"    exit={proc.returncode}   （0.1.0 时：UnicodeEncodeError，exit=1）")


def probe_cli_writeback() -> None:
    banner("探针 K  render 不带 --write-excel / check 的判定   (FINDINGS #8.1 / #8.3 修复验证)")
    import subprocess

    env = dict(os.environ, PYTHONIOENCODING="utf-8")

    def run(*args: str) -> tuple[int, list[str]]:
        proc = subprocess.run(
            [sys.executable, "-m", "excel_codegen", *args],
            cwd=str(HERE),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=env,
            check=False,
        )
        lines = [line for line in (proc.stdout or "").splitlines() if line.strip()]
        return proc.returncode, lines

    code, lines = run("render", "-c", "probe_suffix.yaml", "-x", "probe_suffix.xlsx", "--no-show")
    print("  render（不带 --write-excel）：")
    for line in lines[-6:]:
        print(f"    {line}")
    print(f"    exit={code}")

    code, lines = run("check", "-c", "probe_suffix.yaml", "-x", "probe_suffix.xlsx")
    print("  check（第 1 步没写回；表里是更早一次渲染留下的指纹）：")
    for line in lines[-4:]:
        print(f"    {line}")
    print(f"    exit={code}   （0.1.0 时没有 check 命令）")

    code, lines = run("render", "-c", "probe_suffix.yaml", "-x", "probe_suffix.xlsx", "--write-excel", "--no-show")
    code2, lines2 = run("check", "-c", "probe_suffix.yaml", "-x", "probe_suffix.xlsx")
    print("  写回之后再 check：")
    for line in lines2[-3:]:
        print(f"    {line}")
    print(f"    exit={code2}")


def probe_case_shrink() -> None:
    banner("探针 D2  减少 Case 列之后重渲染（FINDINGS #5 反例）")
    cfg, path = fresh("probe_long.yaml", cases=3)
    out = render_all(cfg, path)
    write_results(path, cfg, out.results)
    ws = load_workbook(path)["Output"]
    print("  3 个 Case 渲染后 D 列（第 3 列）：", [ws.cell(row=r, column=4).value for r in range(1, 11)])

    only = [out.cases[0].name, out.cases[1].name]
    out2 = render_all(cfg, path, only_cases=only)
    write_results(path, cfg, out2.results)
    ws = load_workbook(path)["Output"]
    col_d = [ws.cell(row=r, column=4).value for r in range(1, 11)]
    print("  只渲染 2 个 Case 之后 D 列：", col_d)
    leftover = [v for v in col_d if v]
    print(f"  第 3 列残留：{leftover if leftover else '（无）'}")


def probe_cli_cases() -> None:
    banner("探针 G  init --cases：数量或 Case 名字   (FINDINGS #7)")
    import inspect

    from excel_codegen import cli

    sig = inspect.signature(cli.init_command)
    print(f"  cli.init_command 的 --cases 类型: {sig.parameters['cases'].annotation}")
    sig2 = inspect.signature(create_template)
    print(f"  create_template 的 cases 类型   : {sig2.parameters['cases'].annotation}")
    print("  -> 现在 CLI 的 --cases 也接受逗号分隔的工况名（0.1.0 时只接受整数）。")

    print("  试一下 CLI（--cases 直接给名字）：")
    import subprocess

    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "excel_codegen",
            "init",
            "-c",
            "probe_suffix.yaml",
            "--cases",
            "case_alpha,case_beta",
            "--force",
        ],
        cwd=str(HERE),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    tail = (proc.stderr or proc.stdout).strip().splitlines()
    print("    " + "\n    ".join(tail[-4:]))
    print(f"    exit={proc.returncode}")


def probe_template_sheet() -> None:
    banner("探针 H  --write-excel 之后隐藏的 Template 表还在不在   (原报告「没问题的几件事」)")
    cfg = load_config(HERE / "probe_suffix.yaml")
    path = HERE / cfg.excel.output
    if path.exists():
        path.unlink()
    create_template(cfg, path, cases=1, overwrite=True, include_template_sheet=True)
    for stage in ("init 之后",):
        wb = load_workbook(path)
        print(f"  {stage}: Template 表 state = {wb['Template'].sheet_state!r}")
        wb.close()
    out = render_all(cfg, path)
    write_results(path, cfg, out.results)
    wb = load_workbook(path)
    print(
        f"  write-excel 之后: Template 表 state = {wb['Template'].sheet_state!r}, 表内行数 = {wb['Template'].max_row}"
    )
    wb.close()


def probe_vertical_shrink() -> None:
    banner("探针 I  vertical 布局：Case 数从 8 减到 1   (FINDINGS #5 症状 B)")
    # bottom = row + len(results) + 1，所以只有减少的行数 > 2 才会露出残留。
    cfg, path = fresh("probe_vert.yaml", cases=8)
    out = render_all(cfg, path)
    write_results(path, cfg, out.results)
    ws = load_workbook(path)["Output"]
    print("  8 个 Case 之后 A 列行 1..10：", [ws.cell(row=r, column=1).value for r in range(1, 11)])

    out2 = render_all(cfg, path, only_cases=[out.cases[0].name])
    write_results(path, cfg, out2.results)
    ws = load_workbook(path)["Output"]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, 11)]
    print("  只渲染 1 个 Case 之后 A 列：", col_a)
    leftover = [v for v in col_a if v and v != "Case1"]
    print(f"  Case2..Case8 残留：{leftover if leftover else '（无）'}")


def main() -> int:
    import excel_codegen

    print(f"excel_codegen {excel_codegen.__version__} 探针复测（探针标题里的 FINDINGS 编号指向 0.1.0 的原始报告）")
    probe_suffix()
    probe_value()
    probe_twosets()
    probe_casefilter()
    probe_formula()
    probe_stale()
    probe_case_shrink()
    probe_vertical_shrink()
    probe_dupkey()
    probe_cli_cases()
    probe_template_sheet()
    probe_gbk()
    probe_cli_writeback()
    print("\n完成。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
