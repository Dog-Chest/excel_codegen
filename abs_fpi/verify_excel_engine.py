"""（项目侧工具，**不是** spreadsheet_codegen 的一部分）把 ``engine: excel`` 的公式在 Python 里算一遍。

⚠ 如果你是 spreadsheet_codegen 的**使用者**，你要找的大概率不是这个脚本 ——
工具自带 ``spreadsheet-codegen check --values`` 就能把公式算一遍再与 Python 渲染比对，
而且文法更全（含算术与派生格递归求值）。本脚本是本仓库 ``abs_fpi`` 项目侧的
**独立复核**（另一套实现），用来给"公式引擎算得对"再加一条独立证据。

为什么需要它
------------
公式模式的工作簿里，``Output`` 表的每一格是 Excel 公式。``openpyxl`` **不会算公式**，
所以：

* 0.3.0 时工具自带的 ``check`` 在公式模式下只能比"公式文本是否与当前 YAML 一致"，
  **比不了公式算出来的值** —— 本脚本最早就是为补这个缺口写的；
* 0.4.0 起工具自带的 ``check`` **默认也会**把公式在 Python 里算一遍再比对
  （``--no-values`` 可关），所以本脚本的角色从"补缺口"变成**独立复核**：
  它用的求值器是本仓库自己实现的，与工具内部实现相互独立，两边都过才算真过。

本脚本实现公式子集的一个小求值器（``&`` 拼接 / ``IF`` / ``ISBLANK`` /
``TEXT`` / ``INDEX``+``MATCH`` / 单元格引用），把每格公式算出文本，再与同一份模板的
Python 渲染结果（``render_all``）逐行比对。两边一致 ⇒ 公式模式产出的代码与命令行产出的
代码逐字相同。

它同时是本模板库的**回归工具**：以后新增规范（DNV / BV …）时，
只要模板是纯替换，就能用它验证"Excel 侧"与"Python 侧"没有跑偏。

已知边界：本求值器的文法**不含算术**（见 FINDINGS 0.5.0）。模板一旦用上 ``derived:``
（参数表里会出现 ``+`` / ``*`` 之类），本脚本会报"不能识别的公式片段"，
那时以工具自带的 ``spreadsheet-codegen check --values`` 为准（它含算术、且会递归求值派生格）。

用法::

    python verify_excel_engine.py abs_fpi.yaml workbook.xlsx
    python verify_excel_engine.py abs_fpi.yaml workbook.xlsx --change L=400
"""

from __future__ import annotations

import argparse
import contextlib
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from spreadsheet_codegen import load_config, render_all  # noqa: E402
from spreadsheet_codegen.utils import parse_cell, to_text  # noqa: E402

# --------------------------------------------------------------------------- #
# 词法
# --------------------------------------------------------------------------- #
_TOKEN_RE = re.compile(
    r"""
      (?P<string>"(?:[^"]|"")*")
    | (?P<sheet>'(?:[^']|'')*')
    | (?P<number>\d+(?:\.\d*)?)
    | (?P<ident>[A-Za-z_][A-Za-z0-9_.]*)
    | (?P<punct>[&(),!=:$])
    | (?P<space>\s+)
    """,
    re.VERBOSE,
)


class FormulaError(Exception):
    pass


def tokenize(text: str) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    i = 0
    while i < len(text):
        m = _TOKEN_RE.match(text, i)
        if not m:
            raise FormulaError(f"不能识别的公式片段: {text[i : i + 30]!r}")
        i = m.end()
        kind = m.lastgroup
        if kind == "space":
            continue
        out.append((kind, m.group()))
    return out


# --------------------------------------------------------------------------- #
# 值模型
# --------------------------------------------------------------------------- #
class Cell:
    """一个单元格引用（或 INDEX 的结果）：值 + 是否真空白（ISBLANK 用）。"""

    __slots__ = ("blank", "value", "where")

    def __init__(self, value: Any, blank: bool, where: str = ""):
        self.value = value
        self.blank = blank
        self.where = where

    def text(self) -> str:
        return as_text(self.value)

    def __repr__(self) -> str:  # pragma: no cover
        return f"Cell({self.value!r}, blank={self.blank}, {self.where})"


class ColRef:
    """整列引用，只作为 INDEX / MATCH 的参数出现。"""

    __slots__ = ("column", "sheet")

    def __init__(self, sheet: str, column: str):
        self.sheet = sheet
        self.column = column

    def __repr__(self) -> str:  # pragma: no cover
        return f"ColRef({self.sheet}!{self.column})"


def as_text(value: Any) -> str:
    """Excel 把值拼进字符串时的形态。

    数值走 General 规则 = 最多 15 位有效数字，这正是工具侧 ``to_text`` 的规则
    （0.5.0 起公式里不再用 ``TEXT()``：它会打出 20.559000000000001 这种二进制尾巴）。
    所以这里直接复用工具的实现，两边逐字一致；``TRUE``/``FALSE`` 仍按 Excel 的写法。
    """
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return to_text(value)


# --------------------------------------------------------------------------- #
# 语法（递归下降）
# --------------------------------------------------------------------------- #
class Evaluator:
    def __init__(self, workbook, sheet_names: dict[str, str]):
        self.wb = workbook
        self.sheets = sheet_names  # 小写名 -> 真实表名
        # 变量名 -> {表名: 行号}
        self._name_rows: dict[str, dict[str, int]] = {}

    # -- 引用解析 ---------------------------------------------------------- #
    def sheet_of(self, name: str):
        real = self.sheets.get(name.strip().lower())
        if real is None:
            raise FormulaError(f"公式引用了不存在的工作表 {name!r}")
        return self.wb[real]

    def row_of_name(self, sheet_name: str, variable: str) -> int:
        cache = self._name_rows.setdefault(sheet_name, {})
        if not cache:
            ws = self.wb[sheet_name]
            for r in range(1, ws.max_row + 1):
                v = ws.cell(row=r, column=1).value
                if v is not None and str(v).strip():
                    cache.setdefault(str(v).strip(), r)
        if variable not in cache:
            raise FormulaError(f"{sheet_name} 表里没有变量 {variable!r}")
        return cache[variable]

    def cell(self, sheet_name: str, column: str, row: int) -> Cell:
        ws = self.wb[sheet_name]
        c = 0
        for ch in column:
            c = c * 26 + (ord(ch) - 64)
        value = ws.cell(row=row, column=c).value
        if isinstance(value, str) and value.startswith("="):
            raise FormulaError(f"{sheet_name}!{column}{row} 本身是公式，本求值器只读参数表的值")
        return Cell(value, value is None or value == "", f"{sheet_name}!{column}{row}")

    # -- 解析 -------------------------------------------------------------- #
    def evaluate(self, formula: str, *, relative_column_offset: int = 0) -> str:
        if not formula.startswith("="):
            return str(formula)
        self.tokens = tokenize(formula[1:])
        self.i = 0
        self._offset = relative_column_offset
        value = self.comparison()
        if self.i != len(self.tokens):
            raise FormulaError(f"公式没有解析完: {self.tokens[self.i :]}")
        return value if isinstance(value, str) else as_text(value.value if isinstance(value, Cell) else value)

    # 比较层（工具会生成 `REF=""` 这种条件）
    def comparison(self):
        left = self.concat()
        if self._peek("="):
            self._next()
            right = self.concat()
            return as_text(self._scalar(left)) == as_text(self._scalar(right))
        return left

    # 拼接层
    def concat(self):
        left = self.term()
        while self._peek("&"):
            self._next()
            right = self.term()
            left = as_text(self._scalar(left)) + as_text(self._scalar(right))
        return left

    def term(self):
        kind, text = self._next()
        if kind == "string":
            return text[1:-1].replace('""', '"')
        if kind == "number":
            return float(text) if "." in text else int(text)
        if kind == "sheet" or (kind == "ident" and self._peek("!")):
            return self.reference(kind, text)
        if kind == "ident":
            if not self._peek("("):
                raise FormulaError(f"裸标识符 {text!r} 无法求值")
            self._next()  # (
            name = text.upper()
            if name == "IF":
                cond = self.comparison()
                self._expect(",")
                yes = self.comparison()
                self._expect(",")
                no = self.comparison()
                self._expect(")")
                return yes if self.truthy(cond) else no
            if name == "ISBLANK":
                arg = self.comparison()
                self._expect(")")
                return bool(arg.blank) if isinstance(arg, Cell) else False
            if name == "TEXT":
                value = self.comparison()
                self._expect(",")
                fmt = self.comparison()
                self._expect(")")
                return format_number(self._scalar(value), self._scalar(fmt))
            if name == "INDEX":
                ref = self.comparison()
                self._expect(",")
                row = self.comparison()
                self._expect(")")
                if not isinstance(ref, ColRef):
                    raise FormulaError("INDEX 的第一个参数不是整列引用")
                return self.cell(ref.sheet, ref.column, int(self._scalar(row)))
            if name == "MATCH":
                needle = self.comparison()
                self._expect(",")
                ref = self.comparison()
                self._expect(",")
                self.comparison()  # 0
                self._expect(")")
                if not isinstance(ref, ColRef):
                    raise FormulaError("MATCH 的第二个参数不是整列引用")
                return self.row_of_name(ref.sheet, str(self._scalar(needle)))
            raise FormulaError(f"不支持的函数 {name}()")
        raise FormulaError(f"看不懂的记号 {kind}={text!r}")

    def reference(self, kind: str, text: str):
        sheet = None
        if kind == "sheet":
            sheet = text[1:-1].replace("''", "'")
            self._expect("!")
        # 列
        col = ""
        if self._peek("$"):
            self._next()
        k, t = self._next()
        if k != "ident":
            raise FormulaError(f"引用里缺列标: {t!r}")
        col = t
        # 列区间 或 行号
        if self._peek(":"):
            self._next()
            if self._peek("$"):
                self._next()
            _k2, t2 = self._next()
            if t2.upper() != col.upper():
                raise FormulaError(f"只支持整列引用，不支持列区间 {col}:{t2}")
            return ColRef(sheet, col.upper())
        dollar_row = self._peek("$")
        if dollar_row:
            self._next()
        k3, t3 = self._next()
        if k3 != "number":
            raise FormulaError(f"引用里缺行号: {t3!r}")
        if sheet is None:
            raise FormulaError("公式里出现了没有表名的单元格引用")
        # 相对列：横向拖拽会平移，本求值器按目标列算出实际列
        if not dollar_row and self._offset:
            col = shift_column(col, self._offset)
        return self.cell(sheet, col.upper(), int(float(t3)))

    # -- 小工具 ------------------------------------------------------------ #
    def _next(self):
        if self.i >= len(self.tokens):
            raise FormulaError("公式意外结束")
        tok = self.tokens[self.i]
        self.i += 1
        return tok

    def _peek(self, punct: str) -> bool:
        return self.i < len(self.tokens) and self.tokens[self.i] == ("punct", punct)

    def _expect(self, punct: str) -> None:
        if not self._peek(punct):
            got = self.tokens[self.i] if self.i < len(self.tokens) else "结束"
            raise FormulaError(f"期望 {punct!r}，实际 {got}")
        self.i += 1

    @staticmethod
    def _scalar(value):
        return value.value if isinstance(value, Cell) else value

    def truthy(self, value) -> bool:
        if isinstance(value, Cell):
            return value.blank is False and value.value not in (0, "", None)
        if isinstance(value, str):
            return value.upper() not in ("", "FALSE", "0")
        return bool(value)


def shift_column(letter: str, offset: int) -> str:
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    n += offset
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def format_number(value: Any, fmt: str) -> str:
    """只实现工具用到的两种格式串：``0`` 与 ``0.############``。"""
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if fmt == "0":
        return str(round(number))
    if fmt.startswith("0."):
        decimals = len(fmt) - 2
        text = f"{number:.{decimals}f}".rstrip("0").rstrip(".")
        return text if text not in ("", "-") else "0"
    raise FormulaError(f"不支持的数字格式 {fmt!r}")


# --------------------------------------------------------------------------- #
# 从工作簿里把某个模板的输出区域读出来算一遍
# --------------------------------------------------------------------------- #
def sheet_map(cfg) -> dict[str, str]:
    names = [
        cfg.excel.sheets.global_,
        cfg.excel.sheets.local,
        *cfg.excel.sheets.outputs,
    ]
    if cfg.excel.template_sheet:
        names.append(cfg.excel.template_sheet)
    if cfg.excel.howto_sheet:
        names.append(cfg.excel.howto_sheet)
    return {n.strip().lower(): n for n in names}


def evaluate_block(wb, cfg, template, n_cases: int) -> list[list[str]]:
    """把某个模板在 Output 表里的公式逐格算出来 → ``[case][line]``。"""
    ev = Evaluator(wb, sheet_map(cfg))
    ws = wb[template.output_sheet]
    col, row = parse_cell(template.start_cell)
    blocks: list[list[str]] = []
    for i in range(n_cases):
        lines: list[str] = []
        if template.direction == "horizontal":
            r = row
            while True:
                v = ws.cell(row=r, column=col + i).value
                if v is None:
                    break
                lines.append(ev.evaluate(v))
                r += 1
        else:
            c = col
            target_row = row + i
            while True:
                v = ws.cell(row=target_row, column=c).value
                if v is None:
                    break
                lines.append(ev.evaluate(v))
                c += 1
        blocks.append(lines)
    return blocks


def check_workbook(
    yaml_path: Path, book_path: Path, *, change: str | None = None, verbose: bool = True
) -> tuple[int, int]:
    """返回 ``(检查数, 失败数)``。"""
    cfg = load_config(yaml_path)
    checks = failures = 0

    def cmp(what, got, want):
        nonlocal checks, failures
        checks += 1
        if got == want:
            return True
        failures += 1
        print(f"    FAIL  {what}")
        for k in range(max(len(got), len(want))):
            g = got[k] if k < len(got) else "<缺>"
            w = want[k] if k < len(want) else "<多>"
            if g != w:
                print(f"          第 {k + 1} 行：公式算出 {g!r}")
                print(f"                     Python  {w!r}")
                break
        return False

    if change:
        # 模拟"用户在 Excel 里改了一个参数，Excel 自动重算"：改单元格 → 重新求值公式
        var, value = change.split("=", 1)
        wb = load_workbook(book_path)
        target = None
        for sheet in (cfg.excel.sheets.global_, cfg.excel.sheets.local):
            ws = wb[sheet]
            for r in range(1, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value).strip() == var:
                    target = (sheet, r)
                    break
            if target:
                break
        if not target:
            raise SystemExit(f"工作簿里找不到变量 {var}")
        sheet, r = target
        numeric = None
        with contextlib.suppress(ValueError):
            numeric = float(value)
        wb[sheet].cell(row=r, column=2, value=numeric if numeric is not None else value)
        wb.save(book_path)
        wb.close()
        if verbose:
            print(f"  [改参数] {sheet}!B{r}  {var} = {value}（没有重跑 render）")

    wb = load_workbook(book_path)
    try:
        out = render_all(cfg, book_path)
        for template in cfg.templates:
            if template.engine == "snapshot":
                continue
            results = out.results[template.name]
            if not results:
                continue
            if verbose:
                print(
                    f"  [公式] {template.name}  {len(results)} 个 Case  "
                    f"→ {template.output_sheet} @ {template.start_cell}"
                )
            got_blocks = evaluate_block(wb, cfg, template, len(results))
            for res, got in zip(results, got_blocks, strict=False):
                cmp(f"{template.name} / {res.case_name}（公式 vs Python）", got, res.lines)
    finally:
        wb.close()
    return checks, failures


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="验证 engine: excel 的公式算出来的代码对不对")
    ap.add_argument("yaml", type=Path)
    ap.add_argument("excel", type=Path)
    ap.add_argument("--change", default=None, help="额外做一次'改参数不改工作簿公式'的验证，例如 --change L=400")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    print(f"== {args.yaml.name}  ::  {args.excel.name}")
    checks, failures = check_workbook(args.yaml, args.excel, verbose=not args.quiet)
    print(f"  {checks} 项，{failures} 项失败")

    if args.change:
        import shutil
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / args.excel.name
            shutil.copy(args.excel, copy)
            print(f"\n== 改参数后重算（副本 {args.excel.name}）")
            c2, f2 = check_workbook(args.yaml, copy, change=args.change, verbose=not args.quiet)
            checks += c2
            failures += f2
            print(f"  {c2} 项，{f2} 项失败")

    print(f"\n合计 {checks} 项，{failures} 项失败")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
