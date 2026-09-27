"""Excel 侧实现：生成模板、读取用户填写内容、把渲染结果写回工作表。

工作表约定
----------
Global Parameter
    ``A=Variable``、``B=Value``（用户填写）、``C=Description``、``D=Prefix``、``E=Suffix``

Local Parameter
    ``A=Variable``、``B=Description``、``C=Prefix``、``D=Suffix``、从 ``E`` 列开始每个 Case 一列
    （``E1=Case1``、``F1=Case2`` …），用户可右拉增加 Case。

Output / …
    配置中 ``excel.sheets.outputs`` 声明的输出表，渲染结果写在这里。

Template（可选，隐藏）
    模板原文 + 机器可读的渲染元信息（时间 / 参数指纹 / 输出指纹），方便对照与 `check`。

HOWTO（可选，第一张）
    写进工作簿本身的"下一步跑什么"说明；每次写回结果时刷新。
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from contextlib import suppress
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .derived import (
    DerivedError,
    DerivedNotTranslatable,
    evaluate_derived,
    is_translatable,
    to_excel,
)
from .derived import validate_config as derived_validate_config
from .formula import LONG_FORMULA_WARN, compile_formulas, guarded_lookup
from .models import (
    FIRST_CASE_COLUMN,
    CaseData,
    ProjectConfig,
    RenderResult,
    VariableDef,
)
from .utils import (
    CodeGenError,
    ExcelError,
    VarValue,
    column_index_to_letter,
    fingerprint,
    parse_cell,
    to_text,
)

__all__ = [
    "GLOBAL_HEADERS",
    "LOCAL_HEADERS",
    "META_MARKER",
    "case_column_map",
    "check_required_sheets",
    "create_template",
    "get_sheet",
    "input_fingerprint",
    "load_workbook_file",
    "output_fingerprint",
    "read_cases",
    "read_global_values",
    "read_metadata",
    "template_source",
    "write_results",
]

GLOBAL_HEADERS: tuple[str, ...] = ("Variable", "Value", "Description", "Prefix", "Suffix")
LOCAL_HEADERS: tuple[str, ...] = ("Variable", "Description", "Prefix", "Suffix")

#: 隐藏 Template 表里"机器可读元信息块"的起始标记。
META_MARKER = "## excel-codegen-meta"

#: 元信息键名（Template 表与 HOWTO 表共用同一套键，读取方不必区分来源）。
META_TIME = "时间"
META_INPUT = "参数指纹"
META_OUTPUT = "输出指纹"

#: 清理旧结果时，向右/向下最多扫多少列/行（防止误伤同表里的无关内容）。
_SCAN_LIMIT = 256

#: Global 表列号
_GLOBAL_COL = {"name": 1, "value": 2, "description": 3, "prefix": 4, "suffix": 5}
#: Local 表列号
_LOCAL_COL = {"name": 1, "description": 2, "prefix": 3, "suffix": 4}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_CASE_HEADER_FILL = PatternFill("solid", fgColor="EDEDED")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
#: 派生参数（自动计算）的格：淡绿，和"要你填"的黄色区分开
_DERIVED_FILL = PatternFill("solid", fgColor="E2EFDA")
_TOP_ALIGN = Alignment(vertical="top", wrap_text=False)


def _cell_or(cell_value: Any, default: Any) -> Any:
    """只有"单元格真的为空"才回落 ``default``；非空值原样返回。

    与 ``.strip()`` 的关键区别：``" m"``（带前导空格的单位）是**有内容**的值，
    必须原样保留 —— 否则 ``suffix: " m"`` 会被读成 ``"m"``，生成 ``340m``。
    """
    if cell_value is None:
        return default
    if isinstance(cell_value, str) and cell_value.strip() == "":
        return default
    return cell_value


def _text_or(cell_value: Any, default: Any) -> str:
    """:func:`_cell_or` 的文本版本（前缀/后缀用）。"""
    return to_text(_cell_or(cell_value, default))


# --------------------------------------------------------------------------- #
# 生成 Excel 模板
# --------------------------------------------------------------------------- #
def create_template(
    config: ProjectConfig,
    path: str | Path,
    *,
    cases: int | Sequence[str] = 2,
    overwrite: bool = False,
    include_template_sheet: bool | None = None,
    include_howto_sheet: bool | None = None,
    base_dir: str | Path | None = None,
) -> Path:
    """按配置生成 Excel 模板文件。

    :param cases: 初始 Case 列数量（``int``）或直接给出 Case 名称列表。
    :param overwrite: 目标文件已存在时是否覆盖。
    :param include_template_sheet: ``None`` 时遵循配置；``True/False`` 强制生成/不生成隐藏 Template 表。
    :param include_howto_sheet: ``None`` 时遵循配置；``True/False`` 强制生成/不生成 HOWTO 说明表。
    :param base_dir: 解析 ``template_file`` 相对路径的基准目录，默认使用 ``config.source_dir``。
    """
    target = Path(path)
    if target.exists() and not overwrite:
        raise ExcelError(f"Excel 模板已存在: {target}（需要覆盖请加 --force）")
    if target.exists() and target.is_dir():
        raise ExcelError(f"目标路径是目录: {target}")

    # 派生参数先过一遍配置期检查（语法 / 引用范围 / 循环），别等到渲染才炸
    derived_validate_config(config)

    case_names = _normalise_case_names(cases)

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    global_name = config.excel.sheets.global_
    local_name = config.excel.sheets.local
    _build_global_sheet(workbook.create_sheet(global_name), config)
    _build_local_sheet(workbook.create_sheet(local_name), config, case_names)

    for name in config.excel.sheets.outputs:
        workbook.create_sheet(name)

    sheet_name = config.excel.template_sheet
    if include_template_sheet is False:
        sheet_name = None
    elif include_template_sheet is True and not sheet_name:
        sheet_name = "Template"
    if sheet_name:
        if sheet_name in workbook.sheetnames:
            raise ExcelError(f"Template 工作表名 {sheet_name!r} 与已有工作表冲突，请修改 excel.template_sheet")
        template_sheet = workbook.create_sheet(sheet_name)
        _build_template_sheet(template_sheet, config, base_dir=base_dir or config.source_dir)
        template_sheet.sheet_state = "hidden"

    howto_name = config.excel.howto_sheet
    if include_howto_sheet is False:
        howto_name = None
    elif include_howto_sheet is True and not howto_name:
        howto_name = "HOWTO"
    if howto_name:
        if howto_name in workbook.sheetnames:
            raise ExcelError(f"HOWTO 工作表名 {howto_name!r} 与已有工作表冲突，请修改 excel.howto_sheet")
        howto_sheet = workbook.create_sheet(howto_name)
        # 放在第一张：打开工作簿先看到"下一步跑什么"
        workbook.move_sheet(howto_sheet, offset=-len(workbook.sheetnames) + 1)
        write_howto_sheet(howto_sheet, config, metadata=None)

    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(target)
    except OSError as exc:
        raise ExcelError(f"无法写入 Excel 模板 {target}: {exc}（文件被 Excel 占用？）") from exc
    finally:
        workbook.close()
    return target


def _normalise_case_names(cases: int | Sequence[str]) -> list[str]:
    if isinstance(cases, int):
        if cases < 1:
            raise ExcelError("Case 列数量至少为 1")
        return [f"Case{index}" for index in range(1, cases + 1)]
    names = [to_text(item).strip() for item in cases]
    if not names or any(not name for name in names):
        raise ExcelError("Case 名称不能为空")
    if len(set(names)) != len(names):
        raise ExcelError(f"Case 名称重复: {', '.join(names)}")
    return names


def _build_global_sheet(worksheet: Worksheet, config: ProjectConfig) -> None:
    _write_headers(worksheet, GLOBAL_HEADERS)
    for row, variable in enumerate(config.global_variables, start=2):
        worksheet.cell(row=row, column=_GLOBAL_COL["name"], value=variable.name)
        if variable.is_derived:
            _write_derived_cell(
                worksheet.cell(row=row, column=_GLOBAL_COL["value"]),
                variable,
                resolve=_global_resolver(config),
            )
        else:
            value_cell = worksheet.cell(row=row, column=_GLOBAL_COL["value"], value=variable.default)
            value_cell.fill = _INPUT_FILL
            value_cell.alignment = _TOP_ALIGN
            _add_value_validation(worksheet, variable, [value_cell.coordinate])
        worksheet.cell(row=row, column=_GLOBAL_COL["description"], value=_described(variable))
        worksheet.cell(row=row, column=_GLOBAL_COL["prefix"], value=variable.prefix)
        worksheet.cell(row=row, column=_GLOBAL_COL["suffix"], value=variable.suffix)
    for index, width in enumerate((26, 30, 46, 16, 16), start=1):
        worksheet.column_dimensions[column_index_to_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def _build_local_sheet(worksheet: Worksheet, config: ProjectConfig, case_names: Sequence[str]) -> None:
    _write_headers(worksheet, list(LOCAL_HEADERS) + list(case_names))
    for row, variable in enumerate(config.local_variables, start=2):
        worksheet.cell(row=row, column=_LOCAL_COL["name"], value=variable.name)
        worksheet.cell(row=row, column=_LOCAL_COL["description"], value=_described(variable))
        worksheet.cell(row=row, column=_LOCAL_COL["prefix"], value=variable.prefix)
        worksheet.cell(row=row, column=_LOCAL_COL["suffix"], value=variable.suffix)
        input_cells: list[str] = []
        for offset in range(len(case_names)):
            column = FIRST_CASE_COLUMN + offset
            cell = worksheet.cell(row=row, column=column)
            if variable.is_derived:
                _write_derived_cell(cell, variable, resolve=_local_resolver(config, column))
            else:
                cell.value = variable.default
                cell.fill = _INPUT_FILL
                cell.alignment = _TOP_ALIGN
                input_cells.append(cell.coordinate)
        _add_value_validation(worksheet, variable, input_cells)
    for index, width in enumerate((24, 40, 16, 16), start=1):
        worksheet.column_dimensions[column_index_to_letter(index)].width = width
    for offset, name in enumerate(case_names):
        column = FIRST_CASE_COLUMN + offset
        worksheet.column_dimensions[column_index_to_letter(column)].width = max(16, min(36, len(name) + 14))
    worksheet.freeze_panes = "E2"


def _add_value_validation(worksheet: Worksheet, variable: VariableDef, cells: Sequence[str]) -> None:
    """把变量声明的取值约束写成 Excel 的**数据有效性**（下拉列表 / 数值范围）。

    这是"挡在输入口"的第一道闸，方便人填；**判据仍然是** :func:`check_value_constraints` ——
    读回来的取值一律再查一遍。原因：数据有效性挡不住粘贴、脚本写入和别人发来的老文件，
    而且我们允许空单元格（``allow_blank``），而工具侧认为"声明了约束就不许为空"。

    约束本身表达不了时（下拉列表的选项里带逗号、或拼起来超过 Excel 的 255 字符上限）
    直接报错 —— 与其写一个悄悄失效的校验，不如让人知道。
    """
    if not cells or not variable.has_constraints:
        return
    if not variable.choices and variable.min is None and variable.max is None:
        # 只有 pattern：Excel 的数据有效性没有正则，写不出有意义的校验。
        # 不写总比写一个乱报错的强 —— pattern 由 check_value_constraints 在 render/validate 时检查。
        return

    if variable.choices:
        allowed = [to_text(item) for item in variable.choices]
        if any("," in item for item in allowed):
            raise ExcelError(
                f"变量 {variable.name!r} 的 choices 里有取值含逗号（{', '.join(allowed)}）—— "
                "Excel 下拉列表用逗号分隔，表达不了；请改写取值或去掉 choices"
            )
        joined = ",".join(allowed)
        if len(joined) > 255:
            raise ExcelError(
                f"变量 {variable.name!r} 的 choices 合计超过 Excel 下拉列表的 255 字符上限"
                f"（当前 {len(joined)}）；请减少选项，或去掉 choices（取值仍会在 render / validate 时检查）"
            )
        validation = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
    else:
        operator: str
        low: float | None
        high: float | None
        if variable.min is not None and variable.max is not None:
            operator, low, high = "between", variable.min, variable.max
        elif variable.min is not None:
            operator, low, high = "greaterThanOrEqual", variable.min, None
        else:
            operator, low, high = "lessThanOrEqual", variable.max, None
        validation = DataValidation(
            type="whole" if variable.type == "int" else "decimal",
            operator=operator,
            formula1=to_text(low),
            formula2=None if high is None else to_text(high),
            allow_blank=True,
        )

    prompt = variable.constraint_text
    validation.promptTitle = variable.name
    validation.prompt = prompt
    validation.showInputMessage = True
    validation.errorTitle = "取值不合规"
    validation.error = f"{variable.name}：{prompt}"
    validation.showErrorMessage = True

    worksheet.add_data_validation(validation)
    for coordinate in cells:
        validation.add(coordinate)


# --------------------------------------------------------------------------- #
# 派生参数（derived:）：参数表里那一格写成公式（能翻译时）或算好的值
# --------------------------------------------------------------------------- #
def _described(variable: VariableDef) -> str:
    if not variable.is_derived:
        return variable.description
    expression = variable.derived or ""
    note = f"（自动计算：{expression}）"
    return f"{variable.description} {note}".strip() if variable.description else note


def _global_resolver(config: ProjectConfig):
    """派生表达式里的变量名 -> Excel 引用（Global 表的取值列）。"""
    sheet = config.excel.sheets.global_
    globals_by_name = {item.name: item for item in config.global_variables}

    def resolve(name: str) -> str:
        definition = globals_by_name.get(name)
        if definition is None:
            raise DerivedError(
                f"派生表达式引用了 {name!r}：global 的派生参数只能引用 global 变量（不能引用 local，也不存在别的表）"
            )
        return guarded_lookup(sheet, name, _GLOBAL_COL["value"], definition.default, absolute=True)

    return resolve


def _local_resolver(config: ProjectConfig, case_column: int):
    """派生表达式里的变量名 -> Excel 引用（本 Case 列 / Global 取值列）。"""
    local_sheet = config.excel.sheets.local
    global_sheet = config.excel.sheets.global_
    locals_by_name = {item.name: item for item in config.local_variables}
    globals_by_name = {item.name: item for item in config.global_variables}

    def resolve(name: str) -> str:
        if name in locals_by_name:
            # 派生格里没有"右拉"语义：列标锁死（表格由工具维护）
            return guarded_lookup(local_sheet, name, case_column, locals_by_name[name].default, absolute=True)
        if name in globals_by_name:
            return guarded_lookup(
                global_sheet, name, _GLOBAL_COL["value"], globals_by_name[name].default, absolute=True
            )
        raise DerivedError(f"派生表达式引用了未定义的变量 {name!r}")

    return resolve


def _write_derived_cell(cell, variable: VariableDef, *, resolve, value: Any = None) -> None:
    """写派生参数格：能翻译就写公式，否则写算好的值（没有值就先留空）。"""
    cell.fill = _DERIVED_FILL
    cell.alignment = _TOP_ALIGN
    try:
        formula = to_excel(variable.derived or "", name=variable.name, resolve=resolve)
    except DerivedNotTranslatable:
        formula = None
    if formula is not None:
        cell.value = "=" + formula
        return
    # 翻译不了：只能写入 Python 算好的值（没有就留空，等 --write-excel）
    cell.value = value if value is not None else None


def _build_template_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    *,
    base_dir: str | Path | None = None,
) -> None:
    hint = worksheet.cell(
        row=1,
        column=1,
        value="!! 只读参考：模板真源是 YAML / template_file；在这里改模板不会影响渲染结果",
    )
    hint.font = Font(bold=True, color="C00000")
    row = 2
    for template in config.templates:
        title = worksheet.cell(row=row, column=1, value=f"### {template.name}")
        title.font = _HEADER_FONT
        row += 1
        metadata = (
            ("output_sheet", template.output_sheet),
            ("start_cell", template.start_cell),
            ("direction", template.direction),
            ("filename", template.filename or ""),
            ("extension", template.extension),
            ("description", template.description),
        )
        for key, value in metadata:
            worksheet.cell(row=row, column=1, value=key)
            worksheet.cell(row=row, column=2, value=to_text(value))
            row += 1

        source = template.source_code
        if template.template_file:
            file_path = Path(template.template_file)
            if base_dir and not file_path.is_absolute():
                file_path = Path(base_dir) / file_path
            try:
                source = file_path.read_text(encoding="utf-8")
            except OSError as exc:
                source = f"<<无法读取模板文件 {file_path}: {exc}>>"
        worksheet.cell(row=row, column=1, value="code ↓").font = _HEADER_FONT
        row += 1
        for line_number, line in enumerate(source.splitlines(), start=1):
            worksheet.cell(row=row, column=1, value=line_number)
            worksheet.cell(row=row, column=2, value=line)
            row += 1
        row += 1

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 110


def _write_headers(worksheet: Worksheet, headers: Sequence[str]) -> None:
    for index, text in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=text)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.fill = _CASE_HEADER_FILL if index >= FIRST_CASE_COLUMN else _HEADER_FILL


# --------------------------------------------------------------------------- #
# HOWTO 表：把"下一步跑什么"写进工作簿本身
# --------------------------------------------------------------------------- #
_HOWTO_TITLE = Font(bold=True, size=14, color="1F3864")
_HOWTO_HEAD = Font(bold=True, size=11, color="1F3864")
_HOWTO_MONO = Font(name="Consolas", size=10)
_HOWTO_NOTE = Font(size=9, color="606060")
_HOWTO_WARN = Font(bold=True, size=10, color="C00000")
_HOWTO_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_howto_sheet(
    worksheet: Worksheet,
    config: ProjectConfig,
    *,
    metadata: Mapping[str, str] | None = None,
    command: str | None = None,
) -> None:
    """重写 HOWTO 表：三步说明 + 快照提醒 + 本次生成记录。"""
    for row in worksheet.iter_rows(min_row=1, max_row=max(worksheet.max_row, 1), max_col=2):
        for cell in row:
            cell.value = None

    lines: list[tuple[str, Font, PatternFill | None]] = []
    add = lambda text, font, fill=None: lines.append((text, font, fill))  # noqa: E731

    add(f"{config.excel.output} —— 参数填写与代码生成说明", _HOWTO_TITLE, None)
    add("", _HOWTO_NOTE, None)
    add("这个工作簿怎么用", _HOWTO_HEAD, None)
    add(f"  1.  {config.excel.sheets.global_} 的 B 列：全局变量取值（所有 Case 共用）。", _HOWTO_NOTE, None)
    add(f"  2.  {config.excel.sheets.local} 从 E 列起：每个 Case 一列，右拉复制即可增加。", _HOWTO_NOTE, None)
    add("  3.  改完参数后回到命令行执行：", _HOWTO_NOTE, None)
    add(
        f"          excel-codegen render -c <配置>.yaml -x {config.excel.output} --write-excel",
        _HOWTO_MONO,
        _HOWTO_FILL,
    )
    add("      只想导出代码文件就换成 --outdir <目录>；只预览不写回则什么参数都不加。", _HOWTO_NOTE, None)
    add("      想确认表里的代码是不是已经过期：excel-codegen check -c <配置>.yaml", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)
    formula_templates = [t for t in config.templates if t.engine == "excel"]
    snapshot_templates = [t for t in config.templates if t.engine != "excel"]
    if formula_templates:
        add("★  部分输出表是公式（engine: excel）：改参数后 Excel 打开即重算，不用跑脚本", _HOWTO_HEAD, None)
        add("   公式只做「占位符替换」；要生成代码文件（--outdir）或做 CI 检查（check）仍需命令行。", _HOWTO_NOTE, None)
    if snapshot_templates:
        add("⚠  部分输出表是「快照」，不是活公式", _HOWTO_WARN, None)
        add("   在 Excel 里改了参数、但没跑上面那条命令，输出表里的代码还是上一次的。", _HOWTO_NOTE, None)
    add("   不带 --write-excel 的 render 只预览，不会改变本文件。", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)
    add("输出位置", _HOWTO_HEAD, None)
    for template in config.templates:
        where = f"{template.output_sheet} @ {template.start_cell} ({template.direction})"
        engine = "公式·自动重算" if template.engine == "excel" else "快照·需重跑"
        extra = f"，case_filter: {template.case_filter}" if template.case_filter else ""
        add(f"  {template.name}: {where}［{engine}］{extra}", _HOWTO_MONO, None)
    if config.excel.template_sheet:
        add(f"  {config.excel.template_sheet}（隐藏）：模板原文，只读参考，改它不影响渲染结果。", _HOWTO_NOTE, None)
    add("", _HOWTO_NOTE, None)

    add("本次生成", _HOWTO_HEAD, None)
    if metadata:
        for key, value in metadata.items():
            add(f"  {key:<10}{value}", _HOWTO_NOTE, None)
        if command:
            add(f"  {'命令':<8}{command}", _HOWTO_MONO, None)
    else:
        add("  （尚未渲染：本文件由 excel-codegen init 生成，还没有写回结果）", _HOWTO_WARN, None)
    add("", _HOWTO_NOTE, None)
    add("工况一览（来自上一次渲染）", _HOWTO_HEAD, None)
    cases = (metadata or {}).get("cases") or "（未知）"
    add(f"  {cases}", _HOWTO_MONO, None)

    for offset, (text, font, fill) in enumerate(lines, start=1):
        cell = worksheet.cell(row=offset, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(vertical="top", wrap_text=False)
        if fill is not None:
            cell.fill = fill
    worksheet.column_dimensions["A"].width = 110


def _meta_rows(metadata: Mapping[str, str]) -> list[tuple[str, str]]:
    return [(key, to_text(value)) for key, value in metadata.items()]


def _write_meta_block(worksheet: Worksheet, metadata: Mapping[str, str]) -> None:
    """在隐藏 Template 表末尾写入/刷新机器可读元信息块。"""
    marker_row = None
    for row in range(1, worksheet.max_row + 1):
        if to_text(worksheet.cell(row=row, column=1).value).strip() == META_MARKER:
            marker_row = row
            break
    start = marker_row if marker_row else worksheet.max_row + 2
    if marker_row:
        for row in range(marker_row, worksheet.max_row + 1):
            for column in (1, 2):
                worksheet.cell(row=row, column=column).value = None

    cell = worksheet.cell(row=start, column=1, value=META_MARKER)
    cell.font = _HEADER_FONT
    for offset, (key, value) in enumerate(_meta_rows(metadata), start=1):
        worksheet.cell(row=start + offset, column=1, value=key)
        worksheet.cell(row=start + offset, column=2, value=value)


def read_metadata(workbook: Workbook, config: ProjectConfig) -> dict[str, str]:
    """读回渲染元信息（时间 / 指纹 / 命令）。没写过则返回空字典。"""
    sheet_name = config.excel.template_sheet
    if sheet_name and sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        collecting = False
        found: dict[str, str] = {}
        for row in range(1, worksheet.max_row + 1):
            key = to_text(worksheet.cell(row=row, column=1).value).strip()
            value = to_text(worksheet.cell(row=row, column=2).value)
            if key == META_MARKER:
                collecting = True
                continue
            if collecting:
                if key == "" and value == "":
                    break
                if key:
                    found[key] = value
        if found:
            return found

    # 退化路径：只有 HOWTO 表（template_sheet 被关掉）时，从人读的那几行里抠出来。
    # 键名必须和 Template 表里的元信息块**完全一致**，否则调用方按 "参数指纹" 取值会取空。
    howto_name = config.excel.howto_sheet
    if howto_name and howto_name in workbook.sheetnames:
        worksheet = workbook[howto_name]
        found = {}
        for row in range(1, worksheet.max_row + 1):
            text = to_text(worksheet.cell(row=row, column=1).value).strip()
            for label in (META_TIME, META_INPUT, META_OUTPUT):
                if text.startswith(label):
                    found[label] = text[len(label) :].strip()
        if found:
            return found
    return {}


# --------------------------------------------------------------------------- #
# 读取工作簿
# --------------------------------------------------------------------------- #
def load_workbook_file(path: str | Path) -> Workbook:
    """打开 Excel 文件，失败时给出可读的提示。"""
    target = Path(path)
    if not target.exists():
        raise ExcelError(f"Excel 文件不存在: {target}（请先运行 `excel-codegen init` 生成模板）")
    if target.is_dir():
        raise ExcelError(f"路径是目录而不是 Excel 文件: {target}")
    try:
        return load_workbook(target)
    except Exception as exc:  # openpyxl 会抛各种异常类型
        raise ExcelError(f"无法读取 Excel 文件 {target}: {exc}") from exc


def get_sheet(workbook: Workbook, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise ExcelError(f"缺少工作表 {name!r}。当前工作表: {', '.join(workbook.sheetnames)}")
    return workbook[name]


def check_required_sheets(workbook: Workbook, config: ProjectConfig) -> None:
    """校验 Global / Local / Output 工作表是否存在，一次报出所有缺失项。"""
    required = [
        config.excel.sheets.global_,
        config.excel.sheets.local,
        *config.excel.sheets.outputs,
    ]
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise ExcelError(
            "Excel 缺少工作表: "
            + ", ".join(repr(name) for name in missing)
            + f"。当前工作表: {', '.join(workbook.sheetnames)}"
            + "（可重新运行 `excel-codegen init --force` 生成模板）"
        )
    _check_header(workbook[config.excel.sheets.global_], GLOBAL_HEADERS, config.excel.sheets.global_)
    _check_header(workbook[config.excel.sheets.local], LOCAL_HEADERS, config.excel.sheets.local, strict=False)


def _check_header(
    worksheet: Worksheet,
    expected: Sequence[str],
    sheet_name: str,
    *,
    strict: bool = True,
) -> None:
    for index, title in enumerate(expected, start=1):
        actual = to_text(worksheet.cell(row=1, column=index).value).strip()
        if actual and actual.lower() != title.lower():
            raise ExcelError(f"工作表 {sheet_name!r} 第 1 行第 {index} 列表头应为 {title!r}，实际是 {actual!r}")
        if not actual and strict:
            raise ExcelError(f"工作表 {sheet_name!r} 第 1 行第 {index} 列表头为空，应为 {title!r}")


def read_global_values(
    workbook: Workbook,
    config: ProjectConfig,
    *,
    warnings: list[str] | None = None,
) -> dict[str, VarValue]:
    """读取 Global Parameter 表：B 列值 + D/E 列前缀后缀；**派生参数在这里算出来**。

    回落规则：**只有单元格真的为空才回落 YAML**；非空值原样使用（首尾空格有意义），
    所以 ``suffix: " m"`` 会得到 ``" m"`` 而不是 ``"m"``。

    派生参数（``derived:``）的格子由工具写成公式 / 算好的值，这里**不读它**，
    而是用表达式现算（同 Case 的其他参数 + 全局参数）。
    """
    worksheet = get_sheet(workbook, config.excel.sheets.global_)
    definitions = {item.name: item for item in config.global_variables}
    values: dict[str, VarValue] = {}
    raw_values: dict[str, Any] = {}
    seen: set[str] = set()

    for row in range(2, worksheet.max_row + 1):
        name = to_text(worksheet.cell(row=row, column=_GLOBAL_COL["name"]).value).strip()
        if not name:
            continue
        if name in seen:
            raise ExcelError(f"工作表 {config.excel.sheets.global_!r} 第 {row} 行变量名 {name!r} 重复")
        seen.add(name)

        definition = definitions.get(name)
        prefix = _text_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["prefix"]).value,
            definition.prefix if definition else "",
        )
        suffix = _text_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["suffix"]).value,
            definition.suffix if definition else "",
        )

        if definition is not None and definition.is_derived:
            # 派生格：留个占位，等输入读完之后统一算
            values[name] = VarValue("", prefix, suffix)
            continue

        raw_value = _cell_or(
            worksheet.cell(row=row, column=_GLOBAL_COL["value"]).value,
            definition.default if definition else "",
        )
        kind = definition.type if definition else "auto"
        coerced = _coerce(raw_value, kind, name)
        raw_values[name] = coerced
        values[name] = VarValue(coerced, prefix, suffix)

    # YAML 中定义但表里没写的变量，用默认值补齐，保证模板引用不报错。
    for definition in config.global_variables:
        if definition.name in values:
            continue
        values[definition.name] = VarValue(definition.default, definition.prefix, definition.suffix)
        if not definition.is_derived:
            raw_values[definition.name] = _coerce(definition.default, definition.type, definition.name)

    _resolve_derived_globals(worksheet, config, values, raw_values, warnings)
    return values


def _resolve_derived_globals(
    worksheet: Worksheet,
    config: ProjectConfig,
    values: dict[str, VarValue],
    raw_values: dict[str, Any],
    warnings: list[str] | None,
) -> None:
    """把 global 的派生参数算出来（按依赖顺序），并检查格子里有没有被手工改过。"""
    derived = [item for item in config.global_variables if item.is_derived]
    if not derived:
        return
    evaluate_derived(
        config.global_variables,
        raw_values,
        scope="global 的派生参数",
        forbidden={
            item.name: "global 的派生参数不能引用 local 变量（那时还没有当前 Case）" for item in config.local_variables
        },
    )
    _warn_hand_edited_derived(worksheet, derived, _GLOBAL_COL["value"], raw_values, warnings, label="Global 表")
    for definition in derived:
        current = values.get(definition.name, VarValue(""))
        computed = _coerce(raw_values[definition.name], definition.type, definition.name)
        values[definition.name] = VarValue(computed, current.prefix, current.suffix)


def _warn_hand_edited_derived(
    worksheet: Worksheet,
    derived: Sequence[VariableDef],
    value_column: int,
    raw_values: Mapping[str, Any],
    warnings: list[str] | None,
    *,
    label: str,
) -> None:
    """派生格是公式（或工具写的值）—— 如果用户手工改成了别的值，提醒他会被忽略。"""
    if warnings is None:
        return
    rows = {to_text(worksheet.cell(row=row, column=1).value).strip(): row for row in range(2, worksheet.max_row + 1)}
    for definition in derived:
        row = rows.get(definition.name)
        if row is None:
            continue
        raw = worksheet.cell(row=row, column=value_column).value
        if isinstance(raw, str) and raw.startswith("="):
            continue  # 正常的公式格
        if is_translatable(definition):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 的格子被手工改成了 {raw!r}，"
                "工具会忽略它（该格由表达式算出来；重跑 --write-excel 会把它改回公式）"
            )
        elif definition.name in raw_values and to_text(raw_values[definition.name]) != to_text(raw):
            warnings.append(
                f"{label} 的派生参数 {definition.name!r} 表里是 {raw!r}，当前算式应为 "
                f"{to_text(raw_values[definition.name])!r} —— 请重跑 --write-excel 刷新"
            )


def read_cases(
    workbook: Workbook,
    config: ProjectConfig,
    *,
    global_values: Mapping[str, VarValue] | None = None,
    warnings: list[str] | None = None,
) -> list[CaseData]:
    """读取 Local Parameter 表：从 E 列开始每个 Case 一列；**派生参数按 Case 算出来**。

    :param global_values: 已读好的全局参数（派生表达式要用）；不给就自己读一遍。
    """
    worksheet = get_sheet(workbook, config.excel.sheets.local)
    definitions = {item.name: item for item in config.local_variables}
    if global_values is None:
        global_values = read_global_values(workbook, config, warnings=warnings)
    global_raw = {name: value.value for name, value in global_values.items()}

    columns = _case_columns(worksheet, config.excel.sheets.local)
    if not columns:
        raise ExcelError(
            f"工作表 {config.excel.sheets.local!r} 从 E1 开始没有 Case 列。"
            "请在 E1 填写 Case1、F1 填写 Case2 …（可右拉复制列）"
        )

    # 行：变量名 / 前缀 / 后缀
    rows: list[tuple[int, str, str, str, VariableDef | None]] = []
    seen: set[str] = set()
    for row in range(2, worksheet.max_row + 1):
        name = to_text(worksheet.cell(row=row, column=_LOCAL_COL["name"]).value).strip()
        if not name:
            continue
        if name in seen:
            raise ExcelError(f"工作表 {config.excel.sheets.local!r} 第 {row} 行变量名 {name!r} 重复")
        seen.add(name)
        definition = definitions.get(name)
        prefix = _text_or(
            worksheet.cell(row=row, column=_LOCAL_COL["prefix"]).value,
            definition.prefix if definition else "",
        )
        suffix = _text_or(
            worksheet.cell(row=row, column=_LOCAL_COL["suffix"]).value,
            definition.suffix if definition else "",
        )
        rows.append((row, name, prefix, suffix, definition))

    if not rows:
        raise ExcelError(f"工作表 {config.excel.sheets.local!r} 没有定义任何局部变量")

    cases: list[CaseData] = []
    for column, case_name in columns:
        values: dict[str, VarValue] = {}
        raw_values: dict[str, Any] = dict(global_raw)  # 局部派生可以引用全局
        explicit = False
        for row, name, prefix, suffix, definition in rows:
            if definition is not None and definition.is_derived:
                values[name] = VarValue("", prefix, suffix)
                continue
            raw_cell = worksheet.cell(row=row, column=column).value
            if raw_cell is not None and not (isinstance(raw_cell, str) and raw_cell.strip() == ""):
                explicit = True
            raw_value = _cell_or(raw_cell, definition.default if definition else "")
            kind = definition.type if definition else "auto"
            coerced = _coerce(raw_value, kind, f"{case_name}.{name}")
            raw_values[name] = coerced
            values[name] = VarValue(coerced, prefix, suffix)

        # YAML 中定义但表里缺少的局部变量，用默认值补齐（派生参数稍后算）
        for definition in config.local_variables:
            if definition.name in values:
                continue
            values[definition.name] = VarValue(definition.default, definition.prefix, definition.suffix)
            if not definition.is_derived:
                raw_values[definition.name] = _coerce(definition.default, definition.type, definition.name)

        _resolve_derived_locals(worksheet, config, column, case_name, values, raw_values, warnings)
        cases.append(CaseData(name=case_name, column=column, values=values, explicit_values=explicit))
    return cases


def _resolve_derived_locals(
    worksheet: Worksheet,
    config: ProjectConfig,
    column: int,
    case_name: str,
    values: dict[str, VarValue],
    raw_values: dict[str, Any],
    warnings: list[str] | None,
) -> None:
    derived = [item for item in config.local_variables if item.is_derived]
    if not derived:
        return
    evaluate_derived(
        config.local_variables,
        raw_values,
        scope=f"Case {case_name!r} 的 local 派生参数",
    )
    _warn_hand_edited_derived(
        worksheet,
        derived,
        column,
        raw_values,
        warnings,
        label=f"Local 表 Case {case_name!r}",
    )
    for definition in derived:
        current = values.get(definition.name, VarValue(""))
        computed = _coerce(raw_values[definition.name], definition.type, definition.name)
        values[definition.name] = VarValue(computed, current.prefix, current.suffix)


def _case_columns(worksheet: Worksheet, sheet_name: str) -> list[tuple[int, str]]:
    columns: list[tuple[int, str]] = []
    seen: set[str] = set()
    upper = max(worksheet.max_column, FIRST_CASE_COLUMN)
    for column in range(FIRST_CASE_COLUMN, upper + 1):
        name = to_text(worksheet.cell(row=1, column=column).value).strip()
        if not name:
            if columns:  # 遇到空列说明 Case 列已经结束
                break
            continue
        if name in seen:
            raise ExcelError(f"工作表 {sheet_name!r} 的 Case 列名重复: {name!r}")
        seen.add(name)
        columns.append((column, name))
    return columns


def _coerce(value: Any, kind: str, label: str) -> Any:
    if value is None or kind in ("auto", "raw"):
        return value
    if kind == "string":
        return to_text(value)
    text = value.strip() if isinstance(value, str) else value
    if kind == "int":
        try:
            return int(float(text))
        except (TypeError, ValueError):
            raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 int") from None
    if kind == "float":
        try:
            return float(text)
        except (TypeError, ValueError):
            raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 float") from None
    if kind == "bool":
        if isinstance(value, bool):
            return value
        text_value = to_text(value).strip().lower()
        if text_value in {"1", "true", "yes", "y", "on"}:
            return True
        if text_value in {"0", "false", "no", "n", "off", ""}:
            return False
        raise ExcelError(f"变量 {label} 的值 {value!r} 无法转换为 bool")
    return value


def check_value_constraints(
    config: ProjectConfig,
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
) -> None:
    """校验表里填的取值是否满足变量声明的 ``min`` / ``max`` / ``choices`` / ``pattern``。

    为什么要它：工具此前只查"变量有没有定义、类型对不对"，**完全不看值** —— 把
    ``20.559`` 手滑打成 ``205.59`` 会一路渲染成错误代码，而 ``check`` 只会说"与参数一致"。
    声明了约束就一定查，**空值也算不合格**（``choices`` 意味着"必须给一个合法取值"）。

    不满足时抛 :class:`ExcelError`，一次列全部问题并指出是哪张表、哪一列、哪个变量。
    """
    problems: list[str] = []
    global_sheet = config.excel.sheets.global_

    for row, variable in enumerate(config.global_variables, start=2):
        if variable.is_derived or not variable.has_constraints:
            continue
        value = global_values.get(variable.name)
        if value is None:
            continue
        problem = variable.value_problem(value.text)
        if problem:
            problems.append(f"  {global_sheet} 第 {row} 行 '{variable.name}'：{problem}")

    local_sheet = config.excel.sheets.local
    for case in cases:
        for row, variable in enumerate(config.local_variables, start=2):
            if variable.is_derived or not variable.has_constraints:
                continue
            value = case.values.get(variable.name)
            if value is None:
                continue
            problem = variable.value_problem(value.text)
            if problem:
                column = column_index_to_letter(case.column)
                problems.append(
                    f"  {local_sheet} 第 {column} 列 '{case.name}' 第 {row} 行 '{variable.name}'：{problem}"
                )

    if problems:
        raise ExcelError(
            f"参数取值不满足变量声明的约束，共 {len(problems)} 处：\n"
            + "\n".join(problems)
            + "\n  → 改 Excel 里的取值，或放宽 YAML 里的 min / max / choices / pattern"
        )


# --------------------------------------------------------------------------- #
# 写回渲染结果
# --------------------------------------------------------------------------- #
def input_fingerprint(
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
) -> str:
    """参数指纹：Global 取值 + 每个 Case 的取值（前缀/后缀也算参数）。"""
    parts: list[str] = []
    for name, value in global_values.items():
        parts.append(f"G|{name}|{value.prefix}|{value.text}|{value.suffix}")
    for case in cases:
        for name, value in case.values.items():
            parts.append(f"L|{case.name}|{name}|{value.prefix}|{value.text}|{value.suffix}")
    return fingerprint(*parts)


def output_fingerprint(rendered: Mapping[str, Sequence[RenderResult]]) -> str:
    """输出指纹：所有模板 × Case 的渲染文本。"""
    parts: list[str] = []
    for template_name, per_case in rendered.items():
        for result in per_case:
            parts.append(f"{template_name}/{result.case_name}\n{result.text}")
    return fingerprint(*parts)


@dataclass(frozen=True)
class _Block:
    """写进输出表的一块内容：一个 Case 名 + 若干"行"。"""

    case_name: str
    lines: list[str]

    @property
    def line_count(self) -> int:
        return len(self.lines)


def _template_source(template, base_dir: str | Path | None) -> str:
    """读取模板源码（公式模式需要它；``template_file`` 相对**声明它的那个文件**所在目录解析）。"""
    if not template.template_file:
        return template.source_code
    path = Path(template.template_file)
    base = template.source_dir or base_dir
    if base and not path.is_absolute():
        path = Path(base) / path
    if not path.exists():
        raise ExcelError(f"模板 {template.name!r} 引用的模板文件不存在: {path}")
    try:
        return path.read_text(encoding="utf-8")
    except OSError as exc:
        raise ExcelError(f"无法读取模板文件 {path}: {exc}") from exc


#: 公开别名（CLI 的 check/validate 也要用同一份解析逻辑）
template_source = _template_source


def case_column_map(workbook: Workbook, config: ProjectConfig) -> dict[str, int]:
    """Local 表的 ``Case 名 -> 列号`` 映射（公式模式要知道每个 Case 在哪一列）。"""
    worksheet = get_sheet(workbook, config.excel.sheets.local)
    return {name: column for column, name in _case_columns(worksheet, config.excel.sheets.local)}


def _formula_blocks(
    workbook: Workbook,
    config: ProjectConfig,
    template,
    results: Sequence[RenderResult],
    case_columns: dict[str, int],
) -> list[_Block]:
    """把模板编译成"每行一个公式"的块。"""
    missing = [result.case_name for result in results if result.case_name not in case_columns]
    if missing:
        raise ExcelError(
            f"模板 {template.name!r} 使用公式模式，但 Local 表里找不到这些 Case 列: "
            f"{', '.join(missing)}（可用: {', '.join(case_columns) or '（无）'}）"
        )
    source = _template_source(template, config.source_dir)
    per_case = compile_formulas(
        template,
        config,
        case_columns=[case_columns[result.case_name] for result in results],
        source=source,
    )
    return [_Block(case_name=result.case_name, lines=lines) for result, lines in zip(results, per_case, strict=False)]


def _row_of(worksheet: Worksheet, variable: str) -> int | None:
    """按变量名在 A 列找行号。"""
    for row in range(2, worksheet.max_row + 1):
        if to_text(worksheet.cell(row=row, column=1).value).strip() == variable:
            return row
    return None


def refresh_derived_cells(
    workbook: Workbook,
    config: ProjectConfig,
    global_values: Mapping[str, VarValue],
    cases: Sequence[CaseData],
    *,
    warnings: list[str] | None = None,
) -> bool:
    """刷新参数表里的派生参数格：能写公式就写公式，否则写入 Python 算好的值。

    :returns: 是否写过 Excel 公式（调用方据此决定要不要设 ``fullCalcOnLoad``）
    """
    if not any(item.is_derived for item in [*config.global_variables, *config.local_variables]):
        return False

    wrote_formula = False
    global_sheet = get_sheet(workbook, config.excel.sheets.global_)
    global_resolve = _global_resolver(config)
    for definition in config.global_variables:
        if not definition.is_derived:
            continue
        row = _row_of(global_sheet, definition.name)
        if row is None:
            continue
        value = global_values.get(definition.name)
        _write_derived_cell(
            global_sheet.cell(row=row, column=_GLOBAL_COL["value"]),
            definition,
            resolve=global_resolve,
            value=value.value if value is not None else None,
        )
        wrote_formula = wrote_formula or is_translatable(definition)

    local_sheet = get_sheet(workbook, config.excel.sheets.local)
    for case in cases:
        resolve = _local_resolver(config, case.column)
        for definition in config.local_variables:
            if not definition.is_derived:
                continue
            row = _row_of(local_sheet, definition.name)
            if row is None:
                continue
            value = case.values.get(definition.name)
            _write_derived_cell(
                local_sheet.cell(row=row, column=case.column),
                definition,
                resolve=resolve,
                value=value.value if value is not None else None,
            )
            wrote_formula = wrote_formula or is_translatable(definition)
    return wrote_formula


def write_results(
    path: str | Path,
    config: ProjectConfig,
    rendered: Mapping[str, Sequence[RenderResult]],
    *,
    update_howto: bool = True,
    command: str | None = None,
    warnings: list[str] | None = None,
) -> Path:
    """把渲染结果写入各模板对应的 Output 表，并在 HOWTO / Template 表里记录指纹。

    * ``direction: horizontal`` —— 每个 Case 一列（结果行向下展开）
    * ``direction: vertical``   —— 每个 Case 一行（结果行向右展开）
    * ``engine: snapshot``      —— 写**文本快照**（默认）
    * ``engine: excel``         —— 写 **Excel 公式**：改参数后由 Excel 自己重算，不用重跑脚本

    写入前会清理旧的输出区域（按**表内原有的真实边界**算，而不是按本次行数），
    因此"改短模板 / 减少 Case 之后重渲染"不会残留上一次的内容。

    :param warnings: 传一个列表进来，会把"公式过长"这类不致命的问题写进去。
    """
    target = Path(path)
    workbook = load_workbook_file(target)
    try:
        case_columns: dict[str, int] | None = None
        formula_written = False
        for template in config.templates:
            results = list(rendered.get(template.name, ()))
            if not results:
                continue
            if template.output_sheet not in workbook.sheetnames:
                workbook.create_sheet(template.output_sheet)
            worksheet = workbook[template.output_sheet]
            column, row = parse_cell(template.start_cell)

            if template.engine == "excel":
                if case_columns is None:
                    case_columns = case_column_map(workbook, config)
                blocks = _formula_blocks(workbook, config, template, results, case_columns)
                formula_written = True
                if warnings is not None:
                    longest = max((len(line) for block in blocks for line in block.lines), default=0)
                    if longest > LONG_FORMULA_WARN:
                        warnings.append(
                            f"模板 {template.name!r} 的最长公式 {longest} 字符"
                            f"（警告阈值 {LONG_FORMULA_WARN}）：一个 {{{{ x }}}} 约展开 300–400 字符，"
                            "建议把这一行拆成多行"
                        )
            else:
                blocks = [_Block(result.case_name, result.lines) for result in results]

            if template.direction == "horizontal":
                _write_horizontal(worksheet, template, blocks, column, row)
            else:
                _write_vertical(worksheet, template, blocks, column, row)

        if formula_written:
            # 让 Excel / WPS 打开文件时立刻重算，而不是显示上一次的缓存值。
            # 公式模式下必须重算（openpyxl 默认恰好也是 True，这里显式写死，不依赖上游默认）
            with suppress(AttributeError):  # pragma: no cover - 老版本 openpyxl 兜底
                workbook.calculation.fullCalcOnLoad = True

        metadata = {
            META_TIME: datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            META_OUTPUT: output_fingerprint(rendered),
        }
        values: dict[str, VarValue] | None = None
        cases: list[CaseData] | None = None
        try:
            values = read_global_values(workbook, config, warnings=warnings)
            cases = read_cases(workbook, config, global_values=values, warnings=warnings)
            metadata[META_INPUT] = input_fingerprint(values, cases)
            metadata["case 数"] = str(len(cases))
            metadata["cases"] = ", ".join(case.name for case in cases) or "（无）"
        except CodeGenError:
            metadata[META_INPUT] = "（读取参数失败）"
            metadata["cases"] = ", ".join(_case_names(rendered)) or "（无）"

        # 派生参数的格子也刷新一遍：能写公式就写公式（改输入自动重算）
        if (
            values is not None
            and cases is not None
            and refresh_derived_cells(workbook, config, values, cases, warnings=warnings)
        ):
            formula_written = True

        _record_metadata(workbook, config, metadata, command=command, update_howto=update_howto)
        workbook.save(target)
    except OSError as exc:
        raise ExcelError(f"无法写回 Excel {target}: {exc}（文件被 Excel 占用？）") from exc
    finally:
        workbook.close()
    return target


def _record_metadata(
    workbook: Workbook,
    config: ProjectConfig,
    metadata: Mapping[str, str],
    *,
    command: str | None,
    update_howto: bool,
) -> None:
    sheet_name = config.excel.template_sheet
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(sheet_name)
            _build_template_sheet(worksheet, config)
            worksheet.sheet_state = "hidden"
        _write_meta_block(workbook[sheet_name], metadata)
    if update_howto and config.excel.howto_sheet:
        howto_name = config.excel.howto_sheet
        if howto_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(howto_name)
            workbook.move_sheet(worksheet, offset=-len(workbook.sheetnames) + 1)
        write_howto_sheet(workbook[howto_name], config, metadata=metadata, command=command)


def _case_names(rendered: Mapping[str, Sequence[RenderResult]]) -> list[str]:
    ordered: list[str] = []
    for per_case in rendered.values():
        for result in per_case:
            if result.case_name not in ordered:
                ordered.append(result.case_name)
    return ordered


def _write_horizontal(worksheet, template, blocks: Sequence[_Block], column: int, row: int) -> None:
    max_lines = max(block.line_count for block in blocks)
    header_row = row - 1 if template.write_case_headers and row > 1 else None
    top = header_row if header_row else row
    last_column = _last_used_column(worksheet, top, worksheet.max_row, column)
    last_column = max(last_column, column + len(blocks) - 1)
    # 底边取"表里原有的真实底边"，而不是本次行数 —— 否则改短模板会留下旧行
    last_row = _last_used_row(worksheet, top, column, last_column)
    bottom = max(row + max_lines + 1, last_row + 1)
    _clear_region(worksheet, top, bottom, column, last_column)

    for offset, block in enumerate(blocks):
        target_column = column + offset
        if header_row:
            cell = worksheet.cell(row=header_row, column=target_column, value=block.case_name)
            cell.font = _HEADER_FONT
            cell.fill = _CASE_HEADER_FILL
        for line_offset, line in enumerate(block.lines):
            worksheet.cell(row=row + line_offset, column=target_column, value=line)
        width = max((len(line) for line in block.lines), default=0)
        worksheet.column_dimensions[column_index_to_letter(target_column)].width = max(
            14, min(160, max(width, len(block.case_name)) + 2)
        )


def _write_vertical(worksheet, template, blocks: Sequence[_Block], column: int, row: int) -> None:
    max_columns = max(block.line_count for block in blocks)
    header_column = column - 1 if template.write_case_headers and column > 1 else None
    left = header_column if header_column else column
    right = max(column + max_columns - 1, column)
    last_column = _last_used_column(worksheet, row, worksheet.max_row, left)
    last_column = max(last_column, right)
    last_row = _last_used_row(worksheet, row, left, last_column)
    bottom = max(row + len(blocks) + 1, last_row + 1)
    _clear_region(worksheet, row, bottom, left, last_column)

    for offset, block in enumerate(blocks):
        target_row = row + offset
        if header_column:
            cell = worksheet.cell(row=target_row, column=header_column, value=block.case_name)
            cell.font = _HEADER_FONT
            cell.fill = _CASE_HEADER_FILL
        for line_offset, line in enumerate(block.lines):
            target_column = column + line_offset
            worksheet.cell(row=target_row, column=target_column, value=line)
            width = max(14, min(160, len(line) + 2))
            current = worksheet.column_dimensions[column_index_to_letter(target_column)].width or 0
            if width > current:
                worksheet.column_dimensions[column_index_to_letter(target_column)].width = width


def _last_used_column(worksheet: Worksheet, top: int, bottom: int, left: int) -> int:
    """在给定行区间内找到 left 右侧最后一个有内容的列（用于清理旧结果）。"""
    upper = max(worksheet.max_column, left)
    upper = min(upper, left + _SCAN_LIMIT)
    last = left - 1
    for row in worksheet.iter_rows(min_row=top, max_row=max(bottom, top), min_col=left, max_col=upper):
        for cell in row:
            if cell.value not in (None, ""):
                last = max(last, cell.column)
    return last


def _last_used_row(worksheet: Worksheet, top: int, left: int, right: int) -> int:
    """在给定列区间内找到 top 之下最后一个有内容的行号（用于清理上一次更长的结果）。"""
    lower = max(worksheet.max_row, top)
    lower = min(lower, top + _SCAN_LIMIT)
    last = top - 1
    for row in worksheet.iter_rows(min_row=top, max_row=lower, min_col=left, max_col=max(right, left)):
        for cell in row:
            if cell.value not in (None, ""):
                last = max(last, cell.row)
    return last


def _clear_region(worksheet: Worksheet, top: int, bottom: int, left: int, right: int) -> None:
    if bottom < top or right < left:
        return
    for row in worksheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
        for cell in row:
            if cell.value is not None:
                cell.value = None
